#!/usr/bin/env python3
"""Climate-Shock Descriptive Statistics and Correlations.

Plan: Report natural-unit distributions and exact Pearson dependence for the
five Stage-1 climate exposures on the same complete village-year sample used
by the approved appendix distribution figure.
Framework: AnaSOP Sections 4-7 Stage-1 data audit and workflow step 2. This
table documents exposure support; it does not estimate the NPP response model.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "data/processed/cses_public_village_cropland_npp_absolute_climate_panel_preprocessed.parquet"
FIGURE_EVIDENCE = ROOT / "data/exp/analysis/climate-npp/climate-shock-distributions-and-correlations"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/climate-shock-descriptive-statistics-and-correlations"
OUTPUT = ROOT / "data/results/tables/Table_climate_shock_descriptive_statistics_and_correlations.xlsx"

ID = "National Village Point ID"
YEAR = "Year"
HEAT_DAYS = "Village Buffer Mean Annual Heat Days at or Above 35 C"
HEAT_DD = "Village Buffer Mean Annual Heat Degree-Days Above 35 C"
RX5DAY = "Village Buffer Mean Annual Maximum Consecutive Five-Day Precipitation Rx5day mm"
DRY_DAYS = "Village Buffer Mean Annual Maximum Consecutive Dry Days Below 1 mm"
RAIN_TOTAL = "Village Buffer Mean Annual Precipitation Total mm"

VARIABLES = [HEAT_DAYS, HEAT_DD, RX5DAY, DRY_DAYS, RAIN_TOTAL]
LABELS = {
    HEAT_DAYS: "Heat days ≥35°C",
    HEAT_DD: "Degree-days >35°C",
    RX5DAY: "Rx5day",
    DRY_DAYS: "Maximum dry spell",
    RAIN_TOTAL: "Annual rainfall",
}
UNITS = {
    HEAT_DAYS: "days/year",
    HEAT_DD: "°C-days/year",
    RX5DAY: "mm",
    DRY_DAYS: "days",
    RAIN_TOTAL: "mm/year",
}


def interpolate_color(value: float) -> str:
    """Map correlations to the approved blue-green-white-yellow-red scale."""
    anchors = [
        (-1.0, (33, 102, 172)),
        (-0.5, (26, 152, 80)),
        (0.0, (255, 255, 255)),
        (0.5, (254, 224, 139)),
        (1.0, (215, 48, 39)),
    ]
    clipped = max(-1.0, min(1.0, float(value)))
    for (left_value, left_color), (right_value, right_color) in zip(anchors[:-1], anchors[1:]):
        if left_value <= clipped <= right_value:
            fraction = (clipped - left_value) / (right_value - left_value)
            rgb = tuple(round(left + fraction * (right - left)) for left, right in zip(left_color, right_color))
            return "".join(f"{channel:02X}" for channel in rgb)
    return "FFFFFF"


def build_results() -> tuple[pd.DataFrame, pd.DataFrame, dict[str, int]]:
    frame = pd.read_parquet(INPUT, columns=[ID, YEAR, *VARIABLES]).dropna(subset=VARIABLES).copy()
    descriptive_rows = []
    for variable in VARIABLES:
        values = frame[variable]
        descriptive_rows.append(
            {
                "Variable": LABELS[variable],
                "Unit": UNITS[variable],
                "N": int(values.count()),
                "Mean": float(values.mean()),
                "SD": float(values.std()),
                "Minimum": float(values.min()),
                "P10": float(values.quantile(0.10)),
                "Median": float(values.median()),
                "P90": float(values.quantile(0.90)),
                "Maximum": float(values.max()),
            }
        )
    descriptive = pd.DataFrame(descriptive_rows)
    correlation = frame[VARIABLES].corr(method="pearson")
    correlation.index = [LABELS[variable] for variable in VARIABLES]
    correlation.columns = [LABELS[variable] for variable in VARIABLES]
    support = {
        "First year": int(frame[YEAR].min()),
        "Last year": int(frame[YEAR].max()),
        "Villages": int(frame[ID].nunique()),
        "Village-years": int(len(frame)),
    }
    assert support == {"First year": 2001, "Last year": 2021, "Villages": 2966, "Village-years": 62286}

    old_descriptive = pd.read_csv(
        FIGURE_EVIDENCE / "climate_shock_natural_unit_distributions.csv", index_col=0
    )
    old_correlation = pd.read_csv(
        FIGURE_EVIDENCE / "climate_shock_pearson_correlations.csv", index_col=0
    )
    assert abs(descriptive.loc[0, "Mean"] - old_descriptive.loc[HEAT_DAYS, "mean"]) < 1e-12
    assert abs(correlation.loc["Heat days ≥35°C", "Degree-days >35°C"] - old_correlation.loc[HEAT_DAYS, HEAT_DD]) < 1e-12

    EVIDENCE.mkdir(parents=True, exist_ok=True)
    descriptive.to_csv(EVIDENCE / "climate_shock_descriptive_statistics.csv", index=False)
    correlation.to_csv(EVIDENCE / "climate_shock_pearson_correlations.csv")
    pd.DataFrame([support]).to_csv(EVIDENCE / "climate_shock_common_sample_support.csv", index=False)
    return descriptive, correlation, support


def write_workbook(
    descriptive: pd.DataFrame,
    correlation: pd.DataFrame,
    support: dict[str, int],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Climate Descriptives"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    blue = "5B9BD5"
    white = "FFFFFF"
    light = "F4F6F7"
    thin = Side(style="thin", color="C8D5DE")
    medium = Side(style="medium", color="7F9DB9")

    ws.merge_cells("A1:J1")
    ws["A1"] = "Climate-Shock Descriptive Statistics and Correlations"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 29

    ws.merge_cells("A3:J3")
    ws["A3"] = "A. Natural-unit distributions on the common complete sample"
    ws["A3"].fill = PatternFill("solid", fgColor=navy)
    ws["A3"].font = Font(name="Times New Roman", size=10.2, bold=True, color=white)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws["A3"].border = Border(top=medium, bottom=medium)
    ws.row_dimensions[3].height = 24

    headers = list(descriptive.columns)
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="Times New Roman", size=9.2, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[4].height = 27

    for row_index, record in enumerate(descriptive.itertuples(index=False, name=None), start=5):
        for column, value in enumerate(record, start=1):
            cell = ws.cell(row_index, column, value)
            cell.font = Font(name="Times New Roman", size=9.0, bold=column == 1)
            cell.alignment = Alignment(
                horizontal="right" if column >= 3 else "left",
                vertical="center",
            )
            cell.border = Border(bottom=thin)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            if column == 3:
                cell.number_format = "#,##0"
            elif column >= 4:
                cell.number_format = "0.00"
        ws.row_dimensions[row_index].height = 26

    ws.merge_cells("A11:F11")
    ws["A11"] = "B. Pearson correlation matrix"
    ws["A11"].fill = PatternFill("solid", fgColor=navy)
    ws["A11"].font = Font(name="Times New Roman", size=10.2, bold=True, color=white)
    ws["A11"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws["A11"].border = Border(top=medium, bottom=medium)

    short_headers = ["Variable", "Heat days", "Degree-days", "Rx5day", "Dry spell", "Rainfall"]
    for column, header in enumerate(short_headers, start=1):
        cell = ws.cell(12, column, header)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="Times New Roman", size=8.8, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[12].height = 29

    for row_offset, row_label in enumerate(correlation.index, start=13):
        label_cell = ws.cell(row_offset, 1, row_label)
        label_cell.font = Font(name="Times New Roman", size=8.8, bold=True)
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        label_cell.border = Border(bottom=thin)
        for column_offset, column_label in enumerate(correlation.columns, start=2):
            value = float(correlation.loc[row_label, column_label])
            cell = ws.cell(row_offset, column_offset, value)
            cell.fill = PatternFill("solid", fgColor=interpolate_color(value))
            cell.font = Font(
                name="Times New Roman",
                size=9.0,
                bold=row_label == column_label,
                color=white if abs(value) >= 0.78 else "222222",
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.number_format = "0.000"
            cell.border = Border(bottom=thin)
        ws.row_dimensions[row_offset].height = 26

    ws.merge_cells("H11:J11")
    ws["H11"] = "Common support"
    ws["H11"].fill = PatternFill("solid", fgColor=navy)
    ws["H11"].font = Font(name="Times New Roman", size=10.2, bold=True, color=white)
    ws["H11"].alignment = Alignment(horizontal="center", vertical="center")
    ws["H11"].border = Border(top=medium, bottom=medium)
    support_rows = [
        ("Period", f"{support['First year']}–{support['Last year']}"),
        ("Villages", support["Villages"]),
        ("Village-years", support["Village-years"]),
        ("Buffer", "5 km"),
        ("Correlation", "Pearson"),
    ]
    for row_offset, (label, value) in enumerate(support_rows, start=12):
        ws.merge_cells(start_row=row_offset, start_column=8, end_row=row_offset, end_column=9)
        label_cell = ws.cell(row_offset, 8, label)
        label_cell.font = Font(name="Times New Roman", size=8.8, bold=True)
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        label_cell.border = Border(bottom=thin)
        value_cell = ws.cell(row_offset, 10, value)
        value_cell.font = Font(name="Times New Roman", size=8.8, bold=label in {"Villages", "Village-years"})
        value_cell.alignment = Alignment(horizontal="right", vertical="center")
        value_cell.border = Border(bottom=thin)
        if isinstance(value, int):
            value_cell.number_format = "#,##0"
        if row_offset % 2 == 0:
            for column in range(8, 11):
                ws.cell(row_offset, column).fill = PatternFill("solid", fgColor=light)
        ws.row_dimensions[row_offset].height = max(ws.row_dimensions[row_offset].height or 15, 26)

    note_row = 19
    notes = [
        "Notes: Statistics use all 62,286 complete village-year observations; no observations are trimmed or standardised. P10 and P90 are pooled empirical percentiles.",
        "Heat days and heat degree-days are alternative heat specifications and are not entered together in the main model. Rx5day is a precipitation-extreme measure, not observed flooding; maximum dry-spell length is not observed drought damage.",
    ]
    for note in notes:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
        cell = ws.cell(note_row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.1, italic=True, color="3F4B52")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[note_row].height = 23
        note_row += 1

    widths = [28, 17, 14, 14, 14, 14, 14, 15, 15, 17]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "C5"
    ws.sheet_view.zoomScale = 80
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.22
    ws.page_margins.top = ws.page_margins.bottom = 0.24
    ws.print_area = "A1:J20"
    wb.properties.title = "Climate-Shock Descriptive Statistics and Correlations"
    wb.properties.subject = "Stage-1 natural-unit climate exposure audit"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Climate Descriptives"]
    ws = wb["Climate Descriptives"]
    assert ws.max_row == 20
    assert ws.max_column == 10
    assert not ws.merged_cells.ranges
    assert ws["C5"].value == 62_286
    assert round(float(ws["C13"].value), 3) == 0.930
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    descriptive, correlation, support = build_results()
    write_workbook(descriptive, correlation, support)
    validate_workbook()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(descriptive.to_string(index=False))
    print(correlation.round(4).to_string())
    print(support)


if __name__ == "__main__":
    main()
