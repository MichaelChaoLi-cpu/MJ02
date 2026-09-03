#!/usr/bin/env python3
"""National NPP-to-Consumption Regression Results.

Plan: Present the fixed survey-weighted Stage-2 regression ladder for total
and food consumption in one compact Stargazer-style table.
Framework: AnaSOP Sections 5-7 exact survey-time effects, household sampling
weights, village-clustered inference, household-composition primary controls,
socioeconomic and village-fixed-effect sensitivities, and exact percentage
interpretation for a 0.1 kg C m-2 increase in prior-year cropland NPP.
"""

from __future__ import annotations

from pathlib import Path

from linearmodels.iv import AbsorbingLS
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
HOUSEHOLDS = ROOT / "data/processed/cses_household_cropland_npp_analysis_preprocessed.parquet"
OUTPUT = ROOT / "data/results/tables/Table_national_npp_to_consumption_regression_results.xlsx"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/national-npp-to-consumption-regression-results"

ID = "National Village Point ID"
WEIGHT = "Household Survey Weight"
SURVEY_YEAR = "Interview Calendar Year"
SURVEY_MONTH = "Interview Month"
NPP = "Prior-Year Strict-Cropland NPP"
TOTAL = "Log Real 2021 Annual Total Consumption per Capita"
FOOD = "Log Real 2021 Annual Food Consumption per Capita"
TOTAL_FLAG = "Stage 2 Total Consumption Complete Case"
FOOD_FLAG = "Stage 2 Food Consumption Complete Case"

COMPOSITION = [
    "Household Size",
    "Female Household Member Share",
    "Mean Household Member Age Years",
    "Child Age 0-14 Share",
    "Older Age 65 Plus Share",
    "Household Dependency Ratio",
]
SOCIOECONOMIC_SOURCE = [
    "Urban Rural",
    "Agricultural Participation",
    "Household Head Ever Attended School",
]
URBAN = "Urban location"
AGRICULTURE = "Agricultural participation"
HEAD_SCHOOL = "Household head ever attended school"
SOCIOECONOMIC = [URBAN, AGRICULTURE, HEAD_SCHOOL]
NPP_INCREMENT = 0.1

SPECIFICATIONS = [
    ("Time controls", [], ["Exact Survey Time"]),
    ("+ composition", COMPOSITION, ["Exact Survey Time"]),
    ("+ socioeconomic", [*COMPOSITION, *SOCIOECONOMIC], ["Exact Survey Time"]),
    ("+ village FE", COMPOSITION, ["Exact Survey Time", ID]),
]


def load_data() -> pd.DataFrame:
    columns = [
        ID,
        WEIGHT,
        SURVEY_YEAR,
        SURVEY_MONTH,
        NPP,
        TOTAL,
        FOOD,
        TOTAL_FLAG,
        FOOD_FLAG,
        *COMPOSITION,
        *SOCIOECONOMIC_SOURCE,
    ]
    frame = pd.read_parquet(HOUSEHOLDS, columns=columns)
    numeric = [
        WEIGHT,
        SURVEY_YEAR,
        SURVEY_MONTH,
        NPP,
        TOTAL,
        FOOD,
        *COMPOSITION,
        "Household Head Ever Attended School",
    ]
    for column in numeric:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    frame[URBAN] = frame["Urban Rural"].astype("string").eq("1.0").astype(float)
    frame[AGRICULTURE] = frame["Agricultural Participation"].astype("boolean").astype(float)
    frame[HEAD_SCHOOL] = pd.to_numeric(frame["Household Head Ever Attended School"], errors="coerce")
    frame["Exact Survey Time"] = (
        frame[SURVEY_YEAR].astype("Int64").astype("string")
        + "-"
        + frame[SURVEY_MONTH].astype("Int64").astype("string").str.zfill(2)
    )
    return frame


def fit_model(sample: pd.DataFrame, outcome: str, controls: list[str], absorb_columns: list[str]) -> object:
    exog = pd.DataFrame({NPP: sample[NPP]}, index=sample.index)
    for control in controls:
        exog[control] = sample[control]
    absorb = pd.DataFrame(index=sample.index)
    for column in absorb_columns:
        absorb[column] = sample[column].astype("category")
    clusters = pd.DataFrame(
        {"Village cluster": pd.Categorical(sample[ID]).codes},
        index=sample.index,
    )
    return AbsorbingLS(
        sample[outcome].astype(float),
        exog.astype(float),
        absorb=absorb,
        weights=sample[WEIGHT].astype(float),
        drop_absorbed=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)


def transformed_effect(beta: float) -> float:
    return 100.0 * (np.exp(NPP_INCREMENT * beta) - 1.0)


def estimate_ladder(
    frame: pd.DataFrame,
    outcome: str,
    flag: str,
    outcome_label: str,
    model_start: int,
) -> tuple[list[dict[str, object]], dict[int, object]]:
    eligible = frame.loc[frame[flag].fillna(False).astype(bool)].copy()
    rows: list[dict[str, object]] = []
    models: dict[int, object] = {}
    base_required = [ID, WEIGHT, "Exact Survey Time", NPP, outcome]
    for offset, (specification, controls, absorb_columns) in enumerate(SPECIFICATIONS):
        model_number = model_start + offset
        required = list(dict.fromkeys([*base_required, *controls, *absorb_columns]))
        sample = eligible.dropna(subset=required).copy()
        model = fit_model(sample, outcome, controls, absorb_columns)
        models[model_number] = model
        interval = model.conf_int(level=0.95).loc[NPP]
        beta = float(model.params[NPP])
        standard_error = float(model.std_errors[NPP])
        lower = float(interval["lower"])
        upper = float(interval["upper"])
        rows.append(
            {
                "Model Number": model_number,
                "Outcome": outcome_label,
                "Specification": specification,
                "NPP Coefficient": beta,
                "NPP Coefficient per 0.1": NPP_INCREMENT * beta,
                "Clustered Standard Error": standard_error,
                "Clustered Standard Error per 0.1": NPP_INCREMENT * standard_error,
                "95 Percent CI Lower": lower,
                "95 Percent CI Upper": upper,
                "Probability Value": float(model.pvalues[NPP]),
                "Percent Difference per 0.1 NPP": transformed_effect(beta),
                "Percent Difference 95 Percent CI Lower": transformed_effect(lower),
                "Percent Difference 95 Percent CI Upper": transformed_effect(upper),
                "Household Composition Controls": bool(COMPOSITION[0] in controls),
                "Socioeconomic Controls": bool(SOCIOECONOMIC[0] in controls),
                "Village Fixed Effects": bool(ID in absorb_columns),
                "Exact Survey-Time Fixed Effects": True,
                "Survey Weighted": True,
                "Village-Clustered Inference": True,
                "Observations": int(model.nobs),
                "Villages": int(sample[ID].nunique()),
                "R Squared": float(model.rsquared),
                "Absorbed R Squared": float(model.absorbed_rsquared),
            }
        )
    return rows, models


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def coefficient_text(row: pd.Series) -> str:
    return (
        f"{row['NPP Coefficient per 0.1']:.5f}{stars(float(row['Probability Value']))}\n"
        f"({row['Clustered Standard Error per 0.1']:.5f})"
    )


def percent_text(row: pd.Series) -> str:
    return (
        f"{row['Percent Difference per 0.1 NPP']:.2f}%{stars(float(row['Probability Value']))}\n"
        f"[{row['Percent Difference 95 Percent CI Lower']:.2f}, "
        f"{row['Percent Difference 95 Percent CI Upper']:.2f}]"
    )


def write_workbook(results: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "National Stage 2"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    mid_blue = "5B9BD5"
    total_blue = "2F80A2"
    food_teal = "3A9D8F"
    pale_green = "E2F0D9"
    pale_blue = "DDEBF7"
    light_gray = "F4F6F7"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    medium = Side(style="medium", color="7F9DB9")
    vertical = Side(style="thin", color="D7E0E6")

    ws.merge_cells("A1:I1")
    ws["A1"] = "Table 4. National NPP-to-Consumption Regression Results"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A3"] = "Outcome"
    ws.merge_cells("B3:E3")
    ws["B3"] = "Log real per-capita total consumption"
    ws.merge_cells("F3:I3")
    ws["F3"] = "Log real per-capita food consumption"
    for column in range(1, 10):
        cell = ws.cell(3, column)
        cell.fill = PatternFill("solid", fgColor=navy if column == 1 else total_blue if column <= 5 else food_teal)
        cell.font = Font(name="Times New Roman", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=medium, bottom=medium, left=vertical if column > 1 else None)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 25

    headers = ["Measure"] + [f"({index})" for index in range(1, 9)]
    for column, label in enumerate(headers, start=1):
        cell = ws.cell(4, column, label)
        cell.fill = PatternFill("solid", fgColor=mid_blue)
        cell.font = Font(name="Times New Roman", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=vertical if column > 1 else None)
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    spec_labels = ["Time only", "+ composition", "+ socioeconomic", "+ village FE"] * 2
    ws["A5"] = "Specification"
    ws["A5"].font = Font(name="Times New Roman", size=9.4, bold=True)
    ws["A5"].alignment = Alignment(horizontal="left", vertical="center")
    ws["A5"].border = Border(bottom=thin)
    for column, label in enumerate(spec_labels, start=2):
        cell = ws.cell(5, column, label)
        cell.font = Font(name="Times New Roman", size=9.0, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=vertical, bottom=thin)
        if column == 3:
            cell.fill = PatternFill("solid", fgColor=pale_green)
        elif column == 7:
            cell.fill = PatternFill("solid", fgColor=pale_blue)
    ws.row_dimensions[5].height = 31

    indexed = results.set_index("Model Number")
    coefficient_row = 6
    ws.cell(coefficient_row, 1, "Prior-year strict-cropland NPP (per 0.1 kg C m⁻²)")
    ws.cell(coefficient_row, 1).font = Font(name="Times New Roman", size=9.5, bold=True)
    ws.cell(coefficient_row, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.cell(coefficient_row, 1).border = Border(bottom=thin)
    for column, model_number in enumerate(range(1, 9), start=2):
        cell = ws.cell(coefficient_row, column, coefficient_text(indexed.loc[model_number]))
        cell.font = Font(name="Times New Roman", size=9.4, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=vertical, bottom=thin)
        if column == 3:
            cell.fill = PatternFill("solid", fgColor=pale_green)
        elif column == 7:
            cell.fill = PatternFill("solid", fgColor=pale_blue)
    ws.row_dimensions[coefficient_row].height = 40

    percent_row = 7
    ws.cell(percent_row, 1, "Implied consumption difference per 0.1 NPP (%)")
    ws.cell(percent_row, 1).font = Font(name="Times New Roman", size=9.5, bold=True)
    ws.cell(percent_row, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.cell(percent_row, 1).border = Border(bottom=thin)
    ws.cell(percent_row, 1).fill = PatternFill("solid", fgColor=light_gray)
    for column, model_number in enumerate(range(1, 9), start=2):
        cell = ws.cell(percent_row, column, percent_text(indexed.loc[model_number]))
        cell.font = Font(name="Times New Roman", size=9.2, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=vertical, bottom=thin)
        if column == 3:
            cell.fill = PatternFill("solid", fgColor=pale_green)
        elif column == 7:
            cell.fill = PatternFill("solid", fgColor=pale_blue)
        else:
            cell.fill = PatternFill("solid", fgColor=light_gray)
    ws.row_dimensions[percent_row].height = 40

    row_number = 8
    ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=9)
    ws.cell(row_number, 1, "Model structure and support")
    ws.cell(row_number, 1).fill = PatternFill("solid", fgColor=mid_blue)
    ws.cell(row_number, 1).font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
    ws.cell(row_number, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row_number].height = 20
    row_number += 1

    structure_rows = [
        ("Exact survey-time fixed effects", ["Yes"] * 8, None),
        ("Household composition controls", ["No", "Yes", "Yes", "Yes"] * 2, None),
        ("Socioeconomic controls", ["No", "No", "Yes", "No"] * 2, None),
        ("Village fixed effects", ["No", "No", "No", "Yes"] * 2, None),
        ("Survey weights", ["Yes"] * 8, None),
        ("Standard errors clustered by village", ["Yes"] * 8, None),
        ("Observations", [int(indexed.loc[i, "Observations"]) for i in range(1, 9)], "#,##0"),
        ("Villages", [int(indexed.loc[i, "Villages"]) for i in range(1, 9)], "#,##0"),
        ("R²", [float(indexed.loc[i, "R Squared"]) for i in range(1, 9)], "0.000"),
        ("Absorbed R²", [float(indexed.loc[i, "Absorbed R Squared"]) for i in range(1, 9)], "0.000"),
    ]
    for index, (label, values, number_format) in enumerate(structure_rows):
        label_cell = ws.cell(row_number, 1, label)
        label_cell.font = Font(name="Times New Roman", size=9.0, bold=index < 6)
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        label_cell.border = Border(bottom=thin)
        if index % 2:
            label_cell.fill = PatternFill("solid", fgColor=light_gray)
        for column, value in enumerate(values, start=2):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="Times New Roman", size=9.0)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=vertical, bottom=thin)
            if number_format:
                cell.number_format = number_format
            if column == 3:
                cell.fill = PatternFill("solid", fgColor=pale_green)
            elif column == 7:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
            elif index % 2:
                cell.fill = PatternFill("solid", fgColor=light_gray)
        ws.row_dimensions[row_number].height = 23
        row_number += 1

    row_number += 1
    notes = [
        "Notes: The first result row reports the log-consumption coefficient and village-clustered standard error for a 0.1 kg C m⁻² increase in prior-year 5 km strict-cropland NPP.",
        "The second row reports the exact percentage difference, 100[exp(0.1β)−1], with 95% confidence interval. All models use released household survey weights and exact survey-time fixed effects.",
        "*** p<0.01; ** p<0.05; * p<0.10. Green marks the pre-specified total-consumption primary model (2); blue marks the food-consumption confirmation (6). Associations are not causal mediation estimates.",
    ]
    for note in notes:
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=9)
        cell = ws.cell(row_number, 1, note)
        cell.font = Font(name="Times New Roman", size=8.1, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row_number].height = 22
        row_number += 1

    widths = [43, 18, 18, 18, 18, 18, 18, 18, 18]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B6"
    ws.sheet_view.zoomScale = 70
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.18
    ws.page_margins.top = ws.page_margins.bottom = 0.22
    ws.print_area = f"A1:I{row_number - 1}"
    wb.properties.title = "National NPP-to-Consumption Regression Results"
    wb.properties.subject = "Stage-2 survey-weighted national regression ladder"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["National Stage 2"]
    ws = wb["National Stage 2"]
    assert ws.max_column == 9
    assert not ws.merged_cells.ranges
    assert ws["A1"].value == "Table 4. National NPP-to-Consumption Regression Results"
    assert [ws.cell(4, column).value for column in range(2, 10)] == [f"({index})" for index in range(1, 9)]
    for sheet_row in ws.iter_rows():
        for cell in sheet_row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    frame = load_data()
    total_rows, _ = estimate_ladder(frame, TOTAL, TOTAL_FLAG, "Total consumption", 1)
    food_rows, _ = estimate_ladder(frame, FOOD, FOOD_FLAG, "Food consumption", 5)
    results = pd.DataFrame([*total_rows, *food_rows])
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    results.to_csv(EVIDENCE / "national_npp_to_consumption_regression_ladder.csv", index=False)
    write_workbook(results)
    validate()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(results.to_string(index=False))


if __name__ == "__main__":
    main()
