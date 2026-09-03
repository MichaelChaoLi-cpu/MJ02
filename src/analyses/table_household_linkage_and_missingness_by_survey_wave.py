#!/usr/bin/env python3
"""Household Linkage and Missingness by Survey Wave.

Plan: audit released households, outcome eligibility, exact interview-to-NPP
timing, village linkage, lagged NPP support, missing outcomes, and complete-case
retention by CSES wave.
Framework: AnaSOP Sections 5-7, workflow steps 1 and 8; no model is estimated.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "data/processed/cses_household_cropland_npp_analysis_preprocessed.parquet"
FIGURE_EVIDENCE = (
    ROOT
    / "data/exp/analysis/climate-npp/household-linkage-and-consumption-support-by-wave"
)
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/household-linkage-and-missingness-by-survey-wave"
OUTPUT = ROOT / "data/results/tables/Table_household_linkage_and_missingness_by_survey_wave.xlsx"

WAVES = ["2004", "2007", "2009", "2011-12", "2013", "2014", "2016", "2017", "2019", "2021"]
DISPLAY_WAVES = {
    "2004": "2004",
    "2007": "2007",
    "2009": "2009",
    "2011-12": "2011–12",
    "2013": "2013",
    "2014": "2014",
    "2016": "2016",
    "2017": "2017",
    "2019": "2019",
    "2021": "2021",
}


def year_range(minimum: object, maximum: object) -> str:
    if pd.isna(minimum) or pd.isna(maximum):
        return "Not eligible"
    minimum_int = int(minimum)
    maximum_int = int(maximum)
    if minimum_int == maximum_int:
        return str(minimum_int)
    return f"{minimum_int}–{maximum_int}"


def build_table() -> tuple[pd.DataFrame, pd.DataFrame]:
    columns = [
        "Survey Wave",
        "Household ID",
        "Main Linked Sample",
        "National Village Point ID",
        "Total Consumption Main Outcome Eligible",
        "Real 2021 Annual Total Consumption per Capita Riels",
        "Prior-Year Strict-Cropland NPP",
        "Stage 2 Total Consumption Complete Case",
        "Stage 2 Food Consumption Complete Case",
        "Household Survey Weight",
        "Interview Calendar Year",
        "Prior NPP Calendar Year",
    ]
    frame = pd.read_parquet(INPUT, columns=columns).copy()
    frame["Survey Wave"] = pd.Categorical(frame["Survey Wave"], categories=WAVES, ordered=True)
    frame["Main"] = frame["Main Linked Sample"].fillna(False).astype(bool)
    frame["Eligible"] = frame["Total Consumption Main Outcome Eligible"].fillna(False).astype(bool)
    frame["Total observed"] = frame["Real 2021 Annual Total Consumption per Capita Riels"].notna()
    frame["NPP linked"] = frame["Prior-Year Strict-Cropland NPP"].notna()
    frame["Total complete"] = frame["Stage 2 Total Consumption Complete Case"].fillna(False).astype(bool)
    frame["Food complete"] = frame["Stage 2 Food Consumption Complete Case"].fillna(False).astype(bool)
    frame["Eligible weight"] = frame["Household Survey Weight"].where(frame["Eligible"], 0).fillna(0)
    frame["Total-complete weight"] = frame["Household Survey Weight"].where(
        frame["Total complete"], 0
    ).fillna(0)

    grouped = (
        frame.groupby("Survey Wave", observed=False)
        .agg(
            **{
                "Released households": ("Household ID", "size"),
                "Main-frame households": ("Main", "sum"),
                "Outcome-eligible households": ("Eligible", "sum"),
                "Total outcome observed": ("Total observed", "sum"),
                "Mapped villages": ("National Village Point ID", "nunique"),
                "Prior-year NPP-linked households": ("NPP linked", "sum"),
                "Total-consumption complete cases": ("Total complete", "sum"),
                "Food-consumption complete cases": ("Food complete", "sum"),
                "Eligible survey weight": ("Eligible weight", "sum"),
                "Complete-case survey weight": ("Total-complete weight", "sum"),
                "Interview year minimum": ("Interview Calendar Year", "min"),
                "Interview year maximum": ("Interview Calendar Year", "max"),
                "Prior NPP year minimum": ("Prior NPP Calendar Year", "min"),
                "Prior NPP year maximum": ("Prior NPP Calendar Year", "max"),
            }
        )
        .reset_index()
    )
    grouped["Total outcome missing"] = (
        grouped["Main-frame households"] - grouped["Total outcome observed"]
    )
    grouped["Weighted total-sample retention"] = (
        grouped["Complete-case survey weight"]
        / grouped["Eligible survey weight"].replace(0, np.nan)
    )
    grouped["Interview → prior NPP year"] = grouped.apply(
        lambda row: (
            "Not eligible"
            if pd.isna(row["Interview year minimum"])
            else f"{year_range(row['Interview year minimum'], row['Interview year maximum'])} → "
            f"{year_range(row['Prior NPP year minimum'], row['Prior NPP year maximum'])}"
        ),
        axis=1,
    )
    grouped["Survey wave"] = grouped["Survey Wave"].astype("object").map(DISPLAY_WAVES)

    display_columns = [
        "Survey wave",
        "Interview → prior NPP year",
        "Released households",
        "Outcome-eligible households",
        "Mapped villages",
        "Prior-year NPP-linked households",
        "Total outcome missing",
        "Total-consumption complete cases",
        "Food-consumption complete cases",
        "Weighted total-sample retention",
    ]
    display = grouped[display_columns].copy()

    eligible = frame["Eligible"]
    total_complete = frame["Total complete"]
    total = {
        "Survey wave": "All waves",
        "Interview → prior NPP year": "2007–2021 → 2006–2020",
        "Released households": int(len(frame)),
        "Outcome-eligible households": int(eligible.sum()),
        "Mapped villages": int(frame["National Village Point ID"].nunique()),
        "Prior-year NPP-linked households": int(frame["NPP linked"].sum()),
        "Total outcome missing": int(frame["Main"].sum() - frame["Total observed"].sum()),
        "Total-consumption complete cases": int(total_complete.sum()),
        "Food-consumption complete cases": int(frame["Food complete"].sum()),
        "Weighted total-sample retention": float(
            frame.loc[total_complete, "Household Survey Weight"].sum()
            / frame.loc[eligible, "Household Survey Weight"].sum()
        ),
    }
    display = pd.concat([display, pd.DataFrame([total])], ignore_index=True)

    assert display.loc[0, "Released households"] == 14_984
    assert display.loc[0, "Prior-year NPP-linked households"] == 0
    assert display.loc[8, "Interview → prior NPP year"] == "2019–2020 → 2018–2019"
    assert total["Released households"] == 77_904
    assert total["Outcome-eligible households"] == 62_526
    assert total["Mapped villages"] == 2_966
    assert total["Prior-year NPP-linked households"] == 44_537
    assert total["Total outcome missing"] == 394
    assert total["Total-consumption complete cases"] == 43_120
    assert total["Food-consumption complete cases"] == 43_365

    # Cross-check the already reviewed wave-level figure evidence.
    figure_summary = pd.read_csv(FIGURE_EVIDENCE / "household_linkage_consumption_support_by_wave.csv")
    assert int(figure_summary["Released Households"].sum()) == total["Released households"]
    assert int(figure_summary["Prior-Year NPP Linked Households"].sum()) == total[
        "Prior-year NPP-linked households"
    ]
    assert int(figure_summary["Total-Consumption Complete Cases"].sum()) == total[
        "Total-consumption complete cases"
    ]

    EVIDENCE.mkdir(parents=True, exist_ok=True)
    grouped.to_csv(EVIDENCE / "household_linkage_missingness_full_audit.csv", index=False)
    display.to_csv(EVIDENCE / "household_linkage_missingness_display.csv", index=False)
    return display, grouped


def write_workbook(table: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Household Linkage"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    blue = "3F7CAC"
    teal = "3A9D8F"
    gold = "D6A84B"
    white = "FFFFFF"
    light = "F4F6F7"
    diagnostic = "ECEFF1"
    thin = Side(style="thin", color="C8D5DE")
    medium = Side(style="medium", color="7F9DB9")

    ws.merge_cells("A1:J1")
    ws["A1"] = "Household Linkage and Missingness by Survey Wave"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 29

    groups = [
        ("A3:B3", "Survey timing", navy),
        ("C3:D3", "Released and eligible samples", blue),
        ("E3:F3", "Village and NPP linkage", teal),
        ("G3:J3", "Outcome support and retention", gold),
    ]
    for cell_range, title, color in groups:
        ws.merge_cells(cell_range)
        cell = ws[cell_range.split(":")[0]]
        cell.value = title
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=medium, bottom=medium)
    ws.row_dimensions[3].height = 23

    headers = [
        "Survey wave",
        "Interview → prior NPP year",
        "Released\nhouseholds",
        "Outcome-eligible\nhouseholds",
        "Mapped\nvillages",
        "Prior-year NPP-linked\nhouseholds",
        "Total outcome\nmissing",
        "Total-consumption\ncomplete cases",
        "Food-consumption\ncomplete cases",
        "Weighted total-sample\nretention",
    ]
    header_colors = [navy, navy, blue, blue, teal, teal, gold, gold, gold, gold]
    for column, (header, color) in enumerate(zip(headers, header_colors, strict=True), start=1):
        cell = ws.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(name="Times New Roman", size=8.6, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[4].height = 38

    for row_number, record in enumerate(table.itertuples(index=False, name=None), start=5):
        is_2004 = row_number == 5
        is_total = row_number == 15
        for column, value in enumerate(record, start=1):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(
                name="Times New Roman",
                size=8.8,
                bold=is_total or (is_2004 and column == 1),
                color="FFFFFF" if is_total else "333333",
            )
            cell.alignment = Alignment(
                horizontal="left" if column in {1, 2} else "right",
                vertical="center",
                indent=1 if column in {1, 2} else 0,
            )
            cell.border = Border(bottom=medium if is_total else thin, top=medium if is_total else thin)
            if is_total:
                cell.fill = PatternFill("solid", fgColor=navy)
            elif is_2004:
                cell.fill = PatternFill("solid", fgColor=diagnostic)
            elif row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            if column in range(3, 10):
                cell.number_format = "#,##0"
            elif column == 10:
                cell.number_format = "0.0%"
        ws.row_dimensions[row_number].height = 25

    note_row = 17
    notes = [
        "Notes: 2004 is retained only as a construction diagnostic because it cannot be linked to the national public village-point frame; it is excluded from Stage 2 regressions.",
        "Total outcome missing equals main-frame households minus households with observed real total consumption. NPP linkage requires nonmissing strict-cropland NPP in the exact calendar year before interview.",
        "Weighted retention is the released survey-weight share of outcome-eligible households retained in the primary total-consumption complete-case sample. Missing values are not imputed.",
    ]
    for note in notes:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
        cell = ws.cell(note_row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.0, italic=True, color="3F4B52")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[note_row].height = 23
        note_row += 1

    widths = [13, 24, 15, 17, 13, 20, 15, 19, 19, 18]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "C5"
    ws.sheet_view.zoomScale = 80
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.25
    ws.print_area = "A1:J19"
    wb.properties.title = "Household Linkage and Missingness by Survey Wave"
    wb.properties.subject = "Stage-2 survey-wave linkage, timing, missingness, and retention audit"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Household Linkage"]
    ws = wb["Household Linkage"]
    assert ws.max_row == 19
    assert ws.max_column == 10
    assert not ws.merged_cells.ranges
    assert ws["A5"].value == "2004"
    assert ws["B5"].value == "Not eligible"
    assert ws["A15"].value == "All waves"
    assert ws["C15"].value == 77_904
    assert ws["F15"].value == 44_537
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    table, _ = build_table()
    write_workbook(table)
    validate_workbook()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
