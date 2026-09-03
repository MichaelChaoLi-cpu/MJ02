#!/usr/bin/env python3
"""Decompose the Gate C NPP result across the two 2011 conflict sectors.

The experiment was approved after the pooled NPP result was opened. It estimates
the frozen natural-unit NPP specification separately for Preah Vihear and
Ta Moan/Ta Krabey, re-estimating outcome-blind overlap weights in each sector
sample. It also estimates both sector doses jointly on the pooled common sample
and tests their coefficient difference.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from linearmodels.panel import PanelOLS
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.stats import t as student_t

from estimate_cambodia_thailand_gate_c_npp import (
    BLOCK,
    CELL,
    CLIMATE,
    DOSE,
    DRY_RAIN,
    MAIN_PERIOD,
    NPP_NATURAL,
    PANEL,
    POST,
    TARGET,
    WEIGHT,
    YEAR,
    fit_model,
)
from freeze_cambodia_thailand_gate_b_protocol import (
    load_features,
    weighted_mean_variance,
)
from test_cambodia_thailand_gate_b_placebo_sites import PREDICTORS, make_weights
from test_cambodia_thailand_gate_b_sector_robustness import (
    distance_to_sites,
    event_sector_coordinates,
)


ROOT = Path(__file__).resolve().parents[2]
OUT = (
    ROOT
    / "data/exp/experiments/cambodia-thailand-gate-c-npp/sector-decomposition"
)
TIDY = OUT / "gate_c_npp_sector_decomposition_tidy.csv"
WORKBOOK = OUT / "gate_c_npp_sector_decomposition_summary.xlsx"
METADATA = OUT / "gate_c_npp_sector_decomposition_metadata.json"

PREAH = "Preah Vihear"
TAMOAN = "Ta Moan/Ta Krabey"
STRICT_CONTROL_KM = 60.0


def named_sector_sites() -> dict[str, np.ndarray]:
    """Return the two DBSCAN-frozen 2011 UCDP site clusters with readable names."""
    output: dict[str, np.ndarray] = {}
    for raw_name, sites in event_sector_coordinates():
        if "Preah Vihear" in raw_name or "Kantharalak" in raw_name:
            name = PREAH
        elif "Ta Kwai" in raw_name or "Ta Muen" in raw_name or "Phanom Dong Rak" in raw_name:
            name = TAMOAN
        else:
            raise RuntimeError(f"Unrecognized 2011 conflict sector: {raw_name}")
        output[name] = sites
    if set(output) != {PREAH, TAMOAN}:
        raise RuntimeError(f"Expected exactly the two declared sectors, got {sorted(output)}")
    return output


def target_row(table: pd.DataFrame, specification: str) -> dict[str, object]:
    row = table.loc[table["Term"].eq(TARGET)].iloc[0].to_dict()
    row["Specification"] = specification
    return row


def estimate_sector(
    base: pd.DataFrame,
    panel: pd.DataFrame,
    distance: np.ndarray,
    sector: str,
) -> tuple[pd.DataFrame, dict[str, object], pd.DataFrame]:
    """Estimate one sector against controls clean of every 2008-2011 event."""
    strict_control = (
        base["Candidate All 2008-2011 Nearest Event Distance km"]
        .gt(STRICT_CONTROL_KM)
        .to_numpy()
    )
    treated = distance <= STRICT_CONTROL_KM
    keep = treated | strict_control
    sample = base.loc[keep].copy()
    sample_distance = distance[keep]
    sample_treated = sample_distance <= STRICT_CONTROL_KM
    weights, max_smd, treated_block_ess = make_weights(sample, sample_treated)
    balance_rows = []
    for variable in PREDICTORS:
        treated_mean, treated_variance = weighted_mean_variance(
            sample.loc[sample_treated, variable].to_numpy(float),
            weights[sample_treated],
        )
        control_mean, control_variance = weighted_mean_variance(
            sample.loc[~sample_treated, variable].to_numpy(float),
            weights[~sample_treated],
        )
        denominator = np.sqrt((treated_variance + control_variance) / 2)
        smd = (
            (treated_mean - control_mean) / denominator if denominator > 0 else 0.0
        )
        balance_rows.append(
            {
                "variable": variable,
                "standardized_mean_difference": float(smd),
                "absolute_standardized_mean_difference": float(abs(smd)),
            }
        )
    balance = sorted(
        balance_rows,
        key=lambda row: row["absolute_standardized_mean_difference"],
        reverse=True,
    )
    sample[WEIGHT] = weights
    sample[DOSE] = np.maximum(0.0, 1.0 - sample_distance / STRICT_CONTROL_KM)

    merge_columns = [CELL, BLOCK, WEIGHT, DOSE]
    frame = panel.merge(
        sample[merge_columns], on=CELL, how="inner", validate="many_to_one"
    )
    specification = f"Sector only: {sector}"
    result, model_diagnostics = fit_model(
        frame, NPP_NATURAL, DRY_RAIN, specification
    )
    diagnostics = {
        **model_diagnostics,
        "sector": sector,
        "treated_grid_cells": int(sample_treated.sum()),
        "control_grid_cells": int((~sample_treated).sum()),
        "treated_spatial_blocks": int(
            sample.loc[sample_treated, BLOCK].nunique()
        ),
        "treated_10km_block_ess": float(treated_block_ess),
        "maximum_absolute_weighted_smd": float(max_smd),
        "maximum_imbalance_variable": balance[0]["variable"],
        "weighted_balance_by_predictor": balance,
        "other_sector_exclusion": (
            "Cells within 60 km of the non-target sector are excluded unless they "
            "are also within 60 km of the target sector; controls are over 60 km "
            "from every 2008-2011 event."
        ),
    }
    return result, diagnostics, sample


def estimate_joint(
    panel: pd.DataFrame,
    base: pd.DataFrame,
    sector_distances: dict[str, np.ndarray],
) -> tuple[pd.DataFrame, dict[str, object]]:
    """Estimate both sector-specific target coefficients on the pooled support."""
    cell_doses = base[[CELL]].copy()
    dose_columns: dict[str, str] = {}
    for sector in [PREAH, TAMOAN]:
        column = f"{sector} conflict dose"
        cell_doses[column] = np.maximum(
            0.0, 1.0 - sector_distances[sector] / STRICT_CONTROL_KM
        )
        dose_columns[sector] = column

    sample = panel.merge(cell_doses, on=CELL, how="inner", validate="many_to_one")
    sample = sample.loc[sample[MAIN_PERIOD]].dropna(
        subset=[
            NPP_NATURAL,
            DRY_RAIN,
            CELL,
            YEAR,
            CLIMATE,
            BLOCK,
            WEIGHT,
            POST,
            *dose_columns.values(),
        ]
    ).copy()
    sample["Climate Cell Year"] = (
        sample[CLIMATE].astype("string") + "__" + sample[YEAR].astype("string")
    )

    target_terms: dict[str, str] = {}
    regressors: list[str] = []
    for sector in [PREAH, TAMOAN]:
        dose = dose_columns[sector]
        dose_post = f"{sector}: dose x post"
        dose_drought = f"{sector}: dose x drought"
        dose_post_drought = f"{sector}: dose x post x drought"
        sample[dose_post] = sample[dose] * sample[POST].astype(float)
        sample[dose_drought] = sample[dose] * sample[DRY_RAIN]
        sample[dose_post_drought] = (
            sample[dose] * sample[POST].astype(float) * sample[DRY_RAIN]
        )
        regressors.extend([dose_post, dose_drought, dose_post_drought])
        target_terms[sector] = dose_post_drought

    panel_indexed = sample.set_index([CELL, YEAR]).sort_index()
    other_effects = pd.DataFrame(
        {
            "Climate Cell Year": pd.Categorical(
                panel_indexed["Climate Cell Year"]
            ).codes
        },
        index=panel_indexed.index,
    )
    clusters = pd.DataFrame(
        {"10 km spatial block": pd.Categorical(panel_indexed[BLOCK]).codes},
        index=panel_indexed.index,
    )
    fitted = PanelOLS(
        panel_indexed[NPP_NATURAL].astype(float),
        panel_indexed[regressors].astype(float),
        weights=panel_indexed[WEIGHT].astype(float),
        entity_effects=True,
        other_effects=other_effects,
        drop_absorbed=True,
        check_rank=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)
    confidence = fitted.conf_int(level=0.95)

    rows: list[dict[str, object]] = []
    for term in regressors:
        sector = PREAH if term.startswith(PREAH) else TAMOAN
        rows.append(
            {
                "Specification": f"Joint common sample: {sector}",
                "Outcome": NPP_NATURAL,
                "Hazard": DRY_RAIN,
                "Term": term,
                "Estimate": float(fitted.params[term]),
                "Clustered Standard Error": float(fitted.std_errors[term]),
                "95% CI Lower": float(confidence.loc[term, "lower"]),
                "95% CI Upper": float(confidence.loc[term, "upper"]),
                "p-value": float(fitted.pvalues[term]),
                "Significant at 5%": bool(fitted.pvalues[term] < 0.05),
                "Observations": int(fitted.nobs),
                "Grid Cells": int(sample[CELL].nunique()),
                "Climate Cells": int(sample[CLIMATE].nunique()),
                "Spatial Blocks": int(sample[BLOCK].nunique()),
                "Years": int(sample[YEAR].nunique()),
                "Weighted R-squared Within": float(fitted.rsquared_within),
                "Fixed Effects": "grid cell; climate cell x year",
                "Inference": "10 km spatial-block clustered; debiased",
                "Weights": "frozen pooled outcome-independent overlap weights",
            }
        )

    preah_term = target_terms[PREAH]
    tamoan_term = target_terms[TAMOAN]
    difference = float(fitted.params[preah_term] - fitted.params[tamoan_term])
    covariance = fitted.cov
    variance = float(
        covariance.loc[preah_term, preah_term]
        + covariance.loc[tamoan_term, tamoan_term]
        - 2 * covariance.loc[preah_term, tamoan_term]
    )
    difference_se = float(np.sqrt(max(variance, 0.0)))
    degrees_freedom = float(fitted.df_resid)
    critical = float(student_t.ppf(0.975, degrees_freedom))
    difference_p = float(
        2 * student_t.sf(abs(difference / difference_se), degrees_freedom)
    )
    rows.append(
        {
            "Specification": f"Joint difference: {PREAH} minus {TAMOAN}",
            "Outcome": NPP_NATURAL,
            "Hazard": DRY_RAIN,
            "Term": "Difference in sector-specific dose x post x drought coefficients",
            "Estimate": difference,
            "Clustered Standard Error": difference_se,
            "95% CI Lower": difference - critical * difference_se,
            "95% CI Upper": difference + critical * difference_se,
            "p-value": difference_p,
            "Significant at 5%": bool(difference_p < 0.05),
            "Observations": int(fitted.nobs),
            "Grid Cells": int(sample[CELL].nunique()),
            "Climate Cells": int(sample[CLIMATE].nunique()),
            "Spatial Blocks": int(sample[BLOCK].nunique()),
            "Years": int(sample[YEAR].nunique()),
            "Weighted R-squared Within": float(fitted.rsquared_within),
            "Fixed Effects": "grid cell; climate cell x year",
            "Inference": "10 km spatial-block clustered Wald contrast; debiased",
            "Weights": "frozen pooled outcome-independent overlap weights",
        }
    )
    diagnostics = {
        "observations": int(fitted.nobs),
        "grid_cells": int(sample[CELL].nunique()),
        "spatial_blocks": int(sample[BLOCK].nunique()),
        "degrees_of_freedom": degrees_freedom,
        "preah_target_term": preah_term,
        "tamoan_target_term": tamoan_term,
        "difference_definition": f"{PREAH} minus {TAMOAN}",
        "weights": "frozen pooled outcome-independent overlap weights",
    }
    return pd.DataFrame(rows), diagnostics


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def write_workbook(summary: pd.DataFrame) -> None:
    columns = summary["Specification"].tolist()
    rows: list[dict[str, object]] = []
    for label, field, formatter in [
        ("Target coefficient", "Estimate", lambda value, row: f"{value:.5f}{stars(float(row['p-value']))}"),
        ("Clustered SE", "Clustered Standard Error", lambda value, row: f"({value:.5f})"),
        ("95% CI", None, lambda value, row: f"[{row['95% CI Lower']:.5f}, {row['95% CI Upper']:.5f}]"),
        ("p-value", "p-value", lambda value, row: f"{value:.4f}"),
        ("Observations", "Observations", lambda value, row: int(value)),
        ("Grid cells", "Grid Cells", lambda value, row: int(value)),
        ("Treated grid cells", "Treated Grid Cells", lambda value, row: "—" if pd.isna(value) else int(value)),
        ("Control grid cells", "Control Grid Cells", lambda value, row: "—" if pd.isna(value) else int(value)),
        ("10 km spatial blocks", "Spatial Blocks", lambda value, row: int(value)),
        ("Treated 10 km block ESS", "Treated 10 km Block ESS", lambda value, row: "—" if pd.isna(value) else f"{value:.1f}"),
        ("Maximum absolute weighted SMD", "Maximum Absolute Weighted SMD", lambda value, row: "—" if pd.isna(value) else f"{value:.3f}"),
        ("Grid fixed effects", None, lambda value, row: "Yes"),
        ("Climate-cell x year fixed effects", None, lambda value, row: "Yes"),
    ]:
        output: dict[str, object] = {"Term / statistic": label}
        for _, row in summary.iterrows():
            value = row[field] if field else np.nan
            output[row["Specification"]] = formatter(value, row)
        rows.append(output)

    table = pd.DataFrame(rows, columns=["Term / statistic", *columns])
    note = pd.DataFrame(
        [
            {
                "Term / statistic": "Notes",
                columns[0]: (
                    "Outcome: annual NPP anomaly (kg C/m2). Target: dose x post-2011 "
                    "x rainfall-deficit intensity. Parentheses: 10 km block-clustered "
                    "SEs. * p<0.10, ** p<0.05, *** p<0.01."
                ),
                **{column: "" for column in columns[1:]},
            }
        ]
    )
    table = pd.concat([table, note], ignore_index=True)
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        table.to_excel(writer, index=False, sheet_name="Sector Decomposition")
        sheet = writer.book["Sector Decomposition"]
        sheet.freeze_panes = "B2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions["A"].width = 37
        for column_number in range(2, len(columns) + 2):
            letter = sheet.cell(1, column_number).column_letter
            sheet.column_dimensions[letter].width = 27
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        note_row = sheet.max_row
        sheet.merge_cells(
            start_row=note_row,
            start_column=2,
            end_row=note_row,
            end_column=len(columns) + 1,
        )
        sheet.cell(note_row, 2).font = Font(size=9, italic=True, color="404040")
        sheet.cell(note_row, 2).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[note_row].height = 34
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    base = load_features()
    sites = named_sector_sites()
    sector_distances = {
        sector: distance_to_sites(base, sector_sites)
        for sector, sector_sites in sites.items()
    }

    panel_columns = [
        CELL,
        YEAR,
        CLIMATE,
        BLOCK,
        WEIGHT,
        DOSE,
        POST,
        MAIN_PERIOD,
        NPP_NATURAL,
        DRY_RAIN,
    ]
    pooled_panel = pd.read_parquet(PANEL, columns=panel_columns)
    base_panel = pooled_panel.drop(columns=[BLOCK, WEIGHT, DOSE])

    print("Estimating pooled frozen model", flush=True)
    pooled_table, pooled_diagnostics = fit_model(
        pooled_panel,
        NPP_NATURAL,
        DRY_RAIN,
        "Pooled two-sector baseline",
    )
    pooled_target = target_row(pooled_table, "Pooled two-sector baseline")
    pooled_target.update(
        {
            "Treated Grid Cells": int(base["Candidate Conflict Year 2011 Nearest Event Distance km"].le(60).sum()),
            "Control Grid Cells": int(base["Candidate All 2008-2011 Nearest Event Distance km"].gt(60).sum()),
            "Treated 10 km Block ESS": np.nan,
            "Maximum Absolute Weighted SMD": np.nan,
        }
    )

    all_tables = [pooled_table]
    summary_rows = [pooled_target]
    sector_diagnostics: list[dict[str, object]] = []
    for sector in [PREAH, TAMOAN]:
        print(f"Estimating {sector} sector-only model", flush=True)
        table, diagnostics, _ = estimate_sector(
            base, base_panel, sector_distances[sector], sector
        )
        all_tables.append(table)
        target = target_row(table, f"{sector} only")
        target.update(
            {
                "Treated Grid Cells": diagnostics["treated_grid_cells"],
                "Control Grid Cells": diagnostics["control_grid_cells"],
                "Treated 10 km Block ESS": diagnostics["treated_10km_block_ess"],
                "Maximum Absolute Weighted SMD": diagnostics[
                    "maximum_absolute_weighted_smd"
                ],
            }
        )
        summary_rows.append(target)
        sector_diagnostics.append(diagnostics)

    print("Estimating joint common-sample sector model and coefficient contrast", flush=True)
    joint_table, joint_diagnostics = estimate_joint(
        pooled_panel, base, sector_distances
    )
    all_tables.append(joint_table)
    for sector in [PREAH, TAMOAN]:
        row = joint_table.loc[
            joint_table["Term"].eq(f"{sector}: dose x post x drought")
        ].iloc[0].to_dict()
        row["Specification"] = f"Joint: {sector}"
        row.update(
            {
                "Treated Grid Cells": int((sector_distances[sector] <= 60).sum()),
                "Control Grid Cells": int(
                    base["Candidate All 2008-2011 Nearest Event Distance km"].gt(60).sum()
                ),
                "Treated 10 km Block ESS": np.nan,
                "Maximum Absolute Weighted SMD": np.nan,
            }
        )
        summary_rows.append(row)
    difference = joint_table.loc[
        joint_table["Term"].eq(
            "Difference in sector-specific dose x post x drought coefficients"
        )
    ].iloc[0].to_dict()
    difference["Specification"] = "Joint difference: Preah minus Ta Moan"
    difference.update(
        {
            "Treated Grid Cells": np.nan,
            "Control Grid Cells": np.nan,
            "Treated 10 km Block ESS": np.nan,
            "Maximum Absolute Weighted SMD": np.nan,
        }
    )
    summary_rows.append(difference)

    tidy = pd.concat(all_tables, ignore_index=True)
    tidy.to_csv(TIDY, index=False)
    summary = pd.DataFrame(summary_rows)
    write_workbook(summary)

    metadata = {
        "status": "completed human-approved Gate C NPP sector decomposition",
        "human_approval_record": "MILI-D-20260822-019",
        "outcome": NPP_NATURAL,
        "hazard": DRY_RAIN,
        "target": TARGET,
        "sector_definition": "30 km DBSCAN clusters of unique 2011 UCDP coordinates",
        "sector_only_control_rule": (
            "Target-sector cells within 60 km versus controls more than 60 km from "
            "every 2008-2011 event; other-sector exposure is excluded."
        ),
        "sector_only_weights": (
            "Outcome-independent binary overlap weights re-estimated separately in "
            "each sector sample."
        ),
        "joint_model": (
            "Both sector-specific dose x post, dose x drought, and dose x post x "
            "drought terms on the pooled common sample with frozen pooled weights."
        ),
        "pooled_model": pooled_diagnostics,
        "sector_models": sector_diagnostics,
        "joint_model_diagnostics": joint_diagnostics,
        "created_at_utc": datetime.now(UTC).isoformat(),
    }
    METADATA.write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    print("\nSector decomposition target estimates")
    print(
        summary[
            [
                "Specification",
                "Estimate",
                "Clustered Standard Error",
                "95% CI Lower",
                "95% CI Upper",
                "p-value",
                "Grid Cells",
                "Treated Grid Cells",
                "Treated 10 km Block ESS",
                "Maximum Absolute Weighted SMD",
            ]
        ].to_string(index=False)
    )
    print(f"\nSaved {TIDY.relative_to(ROOT)}")
    print(f"Saved {WORKBOOK.relative_to(ROOT)}")
    print(f"Saved {METADATA.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
