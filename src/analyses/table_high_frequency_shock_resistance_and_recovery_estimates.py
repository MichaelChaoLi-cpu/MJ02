#!/usr/bin/env python3
"""Create the internal one-sheet high-frequency shock resistance and recovery table."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/exp/high-frequency-boundary-recovery/event_loss_and_recovery_estimates.csv"
OUTPUT = ROOT / "data/exp/internal_output_archive/tables/Table_high_frequency_shock_resistance_and_recovery_estimates.xlsx"


def interpretation(row: pd.Series) -> str:
    if row["ci_low"] <= 0 <= row["ci_high"]:
        return "Confidence interval includes no historical-side difference"
    direction = "higher" if row["estimate"] > 0 else "lower"
    return f"Southwest side is {direction} on this estimand"


def main() -> None:
    data = pd.read_csv(SOURCE)
    data["Outcome"] = data["outcome_family"].map({"EVI": "EVI", "NDVI": "NDVI"})
    data["Shock"] = data["shock_family"] + " rainfall extreme"
    data["Estimand"] = data["estimand"]
    data["Specification"] = data["specification"]
    data["Scale"] = data["scale"]
    data["Estimate"] = data["estimate"]
    data["95% CI"] = data.apply(lambda row: f"[{row['ci_low']:.3f}, {row['ci_high']:.3f}]", axis=1)
    data["p-value"] = data["p_value"]
    data["Pretrend p-value"] = data["pretrend_p_value"]
    data["Observations"] = data["observations"].astype(int)
    data["Events"] = data["events"].astype(int)
    data["Villages"] = data["villages"].astype(int)
    data["Inference"] = data["inference"]
    data["Interpretation"] = data.apply(interpretation, axis=1)
    columns = [
        "Outcome", "Shock", "Estimand", "Specification", "Scale", "Estimate",
        "95% CI", "p-value", "Pretrend p-value", "Observations", "Events",
        "Villages", "Inference", "Interpretation",
    ]
    order = {"Immediate loss": 0, "Cumulative negative loss, periods 0-5": 1, "Restricted recovery time": 2, "Recovered by period 8": 3}
    data["_order"] = data["Estimand"].map(order)
    data = data.sort_values(["Outcome", "Shock", "_order", "Specification"])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recovery Estimates"
    sheet.append(columns)
    for row in data[columns].itertuples(index=False, name=None):
        sheet.append(list(row))

    navy = "1F4E78"
    light_blue = "D9EAF7"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="B7B7B7")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)
    for row in range(2, sheet.max_row + 1):
        if row % 2 == 0:
            for cell in sheet[row]:
                cell.fill = PatternFill("solid", fgColor=light_blue)
        for cell in sheet[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D9D9D9"))
        for column in (6, 8, 9):
            sheet.cell(row, column).number_format = "0.000"
    widths = [10, 20, 34, 18, 20, 12, 20, 12, 17, 14, 10, 11, 34, 46]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 38
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.outlinePr.summaryBelow = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows: {len(data):,}; columns: {len(columns)}; sheets: 1")


if __name__ == "__main__":
    main()
