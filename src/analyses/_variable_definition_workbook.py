"""Shared writer for standalone variable-definition workbooks."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.pagebreak import Break
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def write_variable_dictionary(
    table: pd.DataFrame,
    *,
    title: str,
    sheet_name: str,
    output: Path,
    notes: list[str],
    fit_to_height: int = 1,
    page_break_after_data_rows: list[int] | None = None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.showGridLines = False

    navy, light, green, yellow, gray, white = (
        "1F4E78", "F3F3F3", "E2F0D9", "FFF2CC", "E7E6E6", "FFFFFF"
    )
    thin = Side(style="thin", color="B7C9D6")
    medium = Side(style="medium", color="7F9DB9")
    vertical = Side(style="thin", color="D7E0E6")

    ws.merge_cells("A1:J1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for column, label in enumerate(table.columns, 1):
        cell = ws.cell(3, column, label)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=9.0, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=medium,
            bottom=medium,
            left=vertical if column in {4, 7, 9} else None,
        )
    ws.row_dimensions[3].height = 44

    for offset, (_, values) in enumerate(table.iterrows()):
        row = 4 + offset
        status = str(values.iloc[8])
        for column, value in enumerate(values, 1):
            cell = ws.cell(row, column, None if pd.isna(value) else value)
            cell.font = Font(name="Times New Roman", size=8.7)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = Border(
                bottom=thin,
                left=vertical if column in {4, 7, 9} else None,
            )
            if offset % 2:
                cell.fill = PatternFill("solid", fgColor=light)
        if "Primary" in status or "Current national" in status:
            fill = green
        elif "not final" in status.lower() or "not current" in status.lower():
            fill = gray
        elif "Legacy" in status or "Planned wording" in status:
            fill = yellow
        else:
            fill = None
        if fill:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=fill)
        ws.row_dimensions[row].height = 43

    notes_start = 5 + len(table)
    for offset, note in enumerate(notes):
        row = notes_start + offset
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell = ws.cell(row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.3, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 20

    widths = [31, 39, 17, 23, 18, 30, 46, 24, 31, 42]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:J{3 + len(table)}"
    ws.sheet_view.zoomScale = 55
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_to_height
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.15
    ws.page_margins.top = ws.page_margins.bottom = 0.22
    ws.print_area = f"A1:J{notes_start + len(notes) - 1}"
    if page_break_after_data_rows:
        ws.print_title_rows = "3:3"
        for data_row_count in page_break_after_data_rows:
            ws.row_breaks.append(Break(id=3 + data_row_count))

    wb.properties.title = title
    wb.properties.subject = "Climate-welfare variable definition and harmonization audit"
    wb.properties.creator = "Mike Li"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def validate_variable_dictionary(
    output: Path,
    *,
    sheet_name: str,
    expected_rows: int,
) -> None:
    wb = load_workbook(output, data_only=False)
    assert wb.sheetnames == [sheet_name]
    ws = wb[sheet_name]
    assert ws.max_column == 10
    assert ws.max_row >= expected_rows + 3
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(
                    ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                )
