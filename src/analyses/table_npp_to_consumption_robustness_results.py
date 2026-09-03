#!/usr/bin/env python3
"""NPP-to-Consumption Robustness Results.

Plan: compare the prespecified Stage-2 control, common-sample, land-cover,
buffer, timing, and questionnaire-regime sensitivities for total and food
consumption.
Framework: AnaSOP Sections 5-7, Stage-2 robustness workflow step 10.
"""

from __future__ import annotations

from pathlib import Path

from linearmodels.iv import AbsorbingLS
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
HOUSEHOLDS = ROOT / "data/processed/cses_household_cropland_npp_analysis_preprocessed.parquet"
STAGE1 = ROOT / "data/processed/cses_public_village_cropland_npp_absolute_climate_panel_preprocessed.parquet"
NATIONAL_EVIDENCE = ROOT / "data/exp/analysis/climate-npp/national-npp-to-consumption-regression-results"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/npp-to-consumption-robustness-results"
OUTPUT = ROOT / "data/results/tables/Table_npp_to_consumption_robustness_results.xlsx"

ID = "National Village Point ID"
WEIGHT = "Household Survey Weight"
YEAR = "Interview Calendar Year"
MONTH = "Interview Month"
TOTAL = "Log Real 2021 Annual Total Consumption per Capita"
FOOD = "Log Real 2021 Annual Food Consumption per Capita"
TOTAL_FLAG = "Stage 2 Total Consumption Complete Case"
FOOD_FLAG = "Stage 2 Food Consumption Complete Case"
REGIME = "Total Consumption Instrument Regime"

STRICT_5 = "Prior-Year Strict-Cropland NPP"
INCLUSIVE_5 = "Prior-Year Inclusive-Agriculture NPP"
STRICT_2 = "Prior-Year Strict-Cropland NPP at 2 km"
STRICT_10 = "Prior-Year Strict-Cropland NPP at 10 km"
CURRENT_5 = "Contemporaneous Strict-Cropland NPP"
NPP_INCREMENT = 0.1

COMPOSITION = [
    "Household Size",
    "Female Household Member Share",
    "Mean Household Member Age Years",
    "Child Age 0-14 Share",
    "Older Age 65 Plus Share",
    "Household Dependency Ratio",
]
URBAN = "Urban location"
AGRICULTURE = "Agricultural participation"
HEAD_SCHOOL = "Household head ever attended school"
SOCIOECONOMIC = [URBAN, AGRICULTURE, HEAD_SCHOOL]

REGIME_ORDER = [
    "2007 integrated housing recall",
    "2009-2013 recall plus housing",
    "2014-2017 expanded recall plus housing",
    "2019-2021 itemized recall plus housing and education",
]
REGIME_LABELS = {
    "2007 integrated housing recall": "2007",
    "2009-2013 recall plus housing": "2009–13",
    "2014-2017 expanded recall plus housing": "2014–17",
    "2019-2021 itemized recall plus housing and education": "2019–21",
}


def load_data() -> pd.DataFrame:
    columns = [
        ID,
        WEIGHT,
        YEAR,
        MONTH,
        TOTAL,
        FOOD,
        TOTAL_FLAG,
        FOOD_FLAG,
        REGIME,
        STRICT_5,
        INCLUSIVE_5,
        STRICT_2,
        STRICT_10,
        *COMPOSITION,
        "Urban Rural",
        "Agricultural Participation",
        "Household Head Ever Attended School",
    ]
    frame = pd.read_parquet(HOUSEHOLDS, columns=columns).copy()
    for column in [
        WEIGHT,
        YEAR,
        MONTH,
        TOTAL,
        FOOD,
        STRICT_5,
        INCLUSIVE_5,
        STRICT_2,
        STRICT_10,
        *COMPOSITION,
        "Household Head Ever Attended School",
    ]:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    frame[URBAN] = frame["Urban Rural"].astype("string").eq("1.0").astype(float)
    frame[AGRICULTURE] = frame["Agricultural Participation"].astype("boolean").astype(float)
    frame[HEAD_SCHOOL] = pd.to_numeric(frame["Household Head Ever Attended School"], errors="coerce")
    frame["Exact Survey Time"] = (
        frame[YEAR].astype("Int64").astype("string")
        + "-"
        + frame[MONTH].astype("Int64").astype("string").str.zfill(2)
    )

    current = pd.read_parquet(
        STAGE1,
        columns=[ID, "Year", "Annual Strict-Cropland Mean NPP kg C per m2"],
    ).rename(
        columns={
            "Year": YEAR,
            "Annual Strict-Cropland Mean NPP kg C per m2": CURRENT_5,
        }
    )
    frame = frame.merge(current, on=[ID, YEAR], how="left", validate="many_to_one")
    return frame


def fit_single(
    frame: pd.DataFrame,
    outcome: str,
    flag: str,
    exposure: str,
    controls: list[str],
    village_fixed_effects: bool = False,
    common_sample_controls: list[str] | None = None,
) -> dict[str, object]:
    required = [ID, WEIGHT, "Exact Survey Time", outcome, exposure, *controls]
    if common_sample_controls:
        required.extend(common_sample_controls)
    required = list(dict.fromkeys(required))
    sample = frame.loc[frame[flag].fillna(False).astype(bool)].dropna(subset=required).copy()
    exog = sample[[exposure, *controls]].astype(float).copy()
    absorb_columns = ["Exact Survey Time", *([ID] if village_fixed_effects else [])]
    absorb = pd.DataFrame(
        {column: sample[column].astype("category") for column in absorb_columns},
        index=sample.index,
    )
    clusters = pd.DataFrame(
        {"Village cluster": pd.Categorical(sample[ID]).codes}, index=sample.index
    )
    model = AbsorbingLS(
        sample[outcome].astype(float),
        exog,
        absorb=absorb,
        weights=sample[WEIGHT].astype(float),
        drop_absorbed=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)
    interval = model.conf_int(level=0.95).loc[exposure]
    return result_record(model, exposure, sample, float(interval["lower"]), float(interval["upper"]))


def result_record(
    model: object,
    parameter: str,
    sample: pd.DataFrame,
    lower: float,
    upper: float,
) -> dict[str, object]:
    beta = float(model.params[parameter])
    standard_error = float(model.std_errors[parameter])
    return {
        "Coefficient": beta,
        "Standard Error": standard_error,
        "95 Percent CI Lower": lower,
        "95 Percent CI Upper": upper,
        "Probability Value": float(model.pvalues[parameter]),
        "Percent Difference per 0.1 NPP": 100.0 * (np.exp(NPP_INCREMENT * beta) - 1.0),
        "Percent Difference 95 Percent CI Lower": 100.0 * (np.exp(NPP_INCREMENT * lower) - 1.0),
        "Percent Difference 95 Percent CI Upper": 100.0 * (np.exp(NPP_INCREMENT * upper) - 1.0),
        "Observations": int(model.nobs),
        "Villages": int(sample[ID].nunique()),
        "R Squared": float(model.rsquared),
    }


def fit_regime_slopes(frame: pd.DataFrame, outcome: str, flag: str) -> dict[str, dict[str, object]]:
    required = [ID, WEIGHT, "Exact Survey Time", outcome, STRICT_5, REGIME, *COMPOSITION]
    sample = frame.loc[frame[flag].fillna(False).astype(bool)].dropna(subset=required).copy()
    exog = sample[COMPOSITION].astype(float).copy()
    parameter_by_regime: dict[str, str] = {}
    for index, regime in enumerate(REGIME_ORDER, start=1):
        parameter = f"NPP slope: regime {index}"
        parameter_by_regime[regime] = parameter
        exog[parameter] = sample[STRICT_5] * sample[REGIME].eq(regime).astype(float)
    absorb = pd.DataFrame(
        {"Exact Survey Time": sample["Exact Survey Time"].astype("category")},
        index=sample.index,
    )
    clusters = pd.DataFrame(
        {"Village cluster": pd.Categorical(sample[ID]).codes}, index=sample.index
    )
    model = AbsorbingLS(
        sample[outcome].astype(float),
        exog,
        absorb=absorb,
        weights=sample[WEIGHT].astype(float),
        drop_absorbed=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)
    intervals = model.conf_int(level=0.95)
    results: dict[str, dict[str, object]] = {}
    for regime, parameter in parameter_by_regime.items():
        regime_sample = sample.loc[sample[REGIME].eq(regime)]
        record = result_record(
            model,
            parameter,
            regime_sample,
            float(intervals.loc[parameter, "lower"]),
            float(intervals.loc[parameter, "upper"]),
        )
        # The fitted model uses the full interaction sample; row support is regime-specific.
        record["Observations"] = int(len(regime_sample))
        record["Villages"] = int(regime_sample[ID].nunique())
        results[regime] = record
    return results


def build_results(frame: pd.DataFrame) -> pd.DataFrame:
    specifications = [
        ("A. Model and common-sample checks", "Primary", STRICT_5, COMPOSITION, False, None, "Composition; exact-time FE"),
        ("A. Model and common-sample checks", "+ village fixed effects", STRICT_5, COMPOSITION, True, None, "+ village FE"),
        ("A. Model and common-sample checks", "Common sample: composition", STRICT_5, COMPOSITION, False, SOCIOECONOMIC, "Expanded-control sample"),
        ("A. Model and common-sample checks", "Common sample: socioeconomic", STRICT_5, [*COMPOSITION, *SOCIOECONOMIC], False, SOCIOECONOMIC, "Expanded-control sample"),
        ("B. Exposure-definition and timing checks", "Inclusive agriculture, 5 km", INCLUSIVE_5, COMPOSITION, False, None, "Land-cover definition"),
        ("B. Exposure-definition and timing checks", "Strict cropland, 2 km", STRICT_2, COMPOSITION, False, None, "Narrow buffer"),
        ("B. Exposure-definition and timing checks", "Strict cropland, 10 km", STRICT_10, COMPOSITION, False, None, "Wide buffer"),
        ("B. Exposure-definition and timing checks", "Contemporaneous strict, 5 km", CURRENT_5, COMPOSITION, False, None, "Interview-year NPP"),
    ]
    rows: list[dict[str, object]] = []
    for section, specification, exposure, controls, village_fe, common_controls, change in specifications:
        total = fit_single(frame, TOTAL, TOTAL_FLAG, exposure, controls, village_fe, common_controls)
        food = fit_single(frame, FOOD, FOOD_FLAG, exposure, controls, village_fe, common_controls)
        row: dict[str, object] = {
            "Section": section,
            "Specification": specification,
            "Exposure": exposure,
            "Change from primary": change,
        }
        row.update({f"Total {key}": value for key, value in total.items()})
        row.update({f"Food {key}": value for key, value in food.items()})
        rows.append(row)

    total_regimes = fit_regime_slopes(frame, TOTAL, TOTAL_FLAG)
    food_regimes = fit_regime_slopes(frame, FOOD, FOOD_FLAG)
    for regime in REGIME_ORDER:
        row = {
            "Section": "C. Questionnaire-regime slope interactions",
            "Specification": REGIME_LABELS[regime],
            "Exposure": STRICT_5,
            "Change from primary": "Regime-specific NPP slope",
        }
        row.update({f"Total {key}": value for key, value in total_regimes[regime].items()})
        row.update({f"Food {key}": value for key, value in food_regimes[regime].items()})
        rows.append(row)

    results = pd.DataFrame(rows)
    assert len(results) == 12
    assert int(results.loc[results["Specification"].eq("Primary"), "Total Observations"].iloc[0]) == 43_120
    assert int(results.loc[results["Specification"].eq("Primary"), "Food Observations"].iloc[0]) == 43_365

    national = pd.read_csv(NATIONAL_EVIDENCE / "national_npp_to_consumption_regression_ladder.csv")
    national_primary_total = national.loc[national["Model Number"].eq(2)].iloc[0]
    national_primary_food = national.loc[national["Model Number"].eq(6)].iloc[0]
    robust_primary = results.loc[results["Specification"].eq("Primary")].iloc[0]
    assert np.isclose(
        robust_primary["Total Percent Difference per 0.1 NPP"],
        national_primary_total["Percent Difference per 0.1 NPP"],
    )
    assert np.isclose(
        robust_primary["Food Percent Difference per 0.1 NPP"],
        national_primary_food["Percent Difference per 0.1 NPP"],
    )

    EVIDENCE.mkdir(parents=True, exist_ok=True)
    results.to_csv(EVIDENCE / "npp_to_consumption_robustness_results.csv", index=False)
    return results


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def effect_text(row: pd.Series, prefix: str) -> str:
    return (
        f"{row[f'{prefix} Percent Difference per 0.1 NPP']:.2f}%"
        f"{stars(float(row[f'{prefix} Probability Value']))}\n"
        f"[{row[f'{prefix} Percent Difference 95 Percent CI Lower']:.2f}, "
        f"{row[f'{prefix} Percent Difference 95 Percent CI Upper']:.2f}]"
    )


def exposure_label(exposure: str) -> str:
    return {
        STRICT_5: "Strict, 5 km, t−1",
        INCLUSIVE_5: "Inclusive, 5 km, t−1",
        STRICT_2: "Strict, 2 km, t−1",
        STRICT_10: "Strict, 10 km, t−1",
        CURRENT_5: "Strict, 5 km, t",
    }[exposure]


def write_workbook(results: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stage 2 Robustness"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    blue = "3F7CAC"
    teal = "3A9D8F"
    gold = "D6A84B"
    white = "FFFFFF"
    light = "F4F6F7"
    primary_fill = "E2F0D9"
    thin = Side(style="thin", color="C8D5DE")
    medium = Side(style="medium", color="7F9DB9")

    ws.merge_cells("A1:I1")
    ws["A1"] = "NPP-to-Consumption Robustness Results"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 29

    ws.merge_cells("A3:A4")
    ws["A3"] = "Specification"
    ws.merge_cells("B3:D3")
    ws["B3"] = "Real per-capita total consumption"
    ws.merge_cells("E3:G3")
    ws["E3"] = "Real per-capita food consumption"
    ws.merge_cells("H3:H4")
    ws["H3"] = "NPP exposure"
    ws.merge_cells("I3:I4")
    ws["I3"] = "Model change"
    for column in [1, 2, 5, 8, 9]:
        cell = ws.cell(3, column)
        cell.fill = PatternFill(
            "solid", fgColor=navy if column in {1, 8, 9} else blue if column == 2 else teal
        )
        cell.font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=medium)
    ws.row_dimensions[3].height = 23

    subheaders = ["Difference per 0.1 NPP\n(%, 95% CI)", "N", "Villages"] * 2
    for column, header in enumerate(subheaders, start=2):
        cell = ws.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor=blue if column <= 4 else teal)
        cell.font = Font(name="Times New Roman", size=8.6, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[4].height = 37

    current_row = 5
    for section, section_frame in results.groupby("Section", sort=False):
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
        section_cell = ws.cell(current_row, 1, section)
        section_cell.fill = PatternFill("solid", fgColor=navy)
        section_cell.font = Font(name="Times New Roman", size=9.2, bold=True, color=white)
        section_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[current_row].height = 21
        current_row += 1

        for _, row in section_frame.iterrows():
            values = [
                row["Specification"],
                effect_text(row, "Total"),
                int(row["Total Observations"]),
                int(row["Total Villages"]),
                effect_text(row, "Food"),
                int(row["Food Observations"]),
                int(row["Food Villages"]),
                exposure_label(str(row["Exposure"])),
                row["Change from primary"],
            ]
            is_primary = row["Specification"] == "Primary"
            for column, value in enumerate(values, start=1):
                cell = ws.cell(current_row, column, value)
                cell.font = Font(
                    name="Times New Roman",
                    size=8.7,
                    bold=is_primary or column in {2, 5},
                )
                cell.alignment = Alignment(
                    horizontal="left" if column in {1, 8, 9} else "center",
                    vertical="center",
                    wrap_text=True,
                    indent=1 if column in {1, 8, 9} else 0,
                )
                cell.border = Border(bottom=thin)
                if is_primary:
                    cell.fill = PatternFill("solid", fgColor=primary_fill)
                elif current_row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=light)
                if column in {3, 4, 6, 7}:
                    cell.number_format = "#,##0"
            ws.row_dimensions[current_row].height = 35
            current_row += 1

    note_row = current_row + 1
    notes = [
        "Notes: Cells report 100[exp(0.1β)−1] and 95% confidence intervals for a 0.1 kg C m⁻² increase in NPP. All models use survey weights, exact survey-time fixed effects, and village-clustered standard errors.",
        "*** p<0.01, ** p<0.05, * p<0.10. The green row is the prespecified primary model. Common-sample rows distinguish sample loss from added socioeconomic controls.",
        "Questionnaire-regime rows come from one interacted model per outcome. Results are conditional associations and are not interpreted as causal mediation effects.",
    ]
    for note in notes:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=9)
        cell = ws.cell(note_row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.0, italic=True, color="3F4B52")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[note_row].height = 23
        note_row += 1

    widths = [28, 21, 12, 11, 21, 12, 11, 22, 25]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B6"
    ws.sheet_view.zoomScale = 75
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.22
    ws.page_margins.top = ws.page_margins.bottom = 0.22
    ws.print_area = f"A1:I{note_row - 1}"
    wb.properties.title = "NPP-to-Consumption Robustness Results"
    wb.properties.subject = "Stage-2 control, measurement, timing, and questionnaire-regime sensitivities"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Stage 2 Robustness"]
    ws = wb["Stage 2 Robustness"]
    assert ws.max_column == 9
    assert not ws.merged_cells.ranges
    assert ws["A1"].value == "NPP-to-Consumption Robustness Results"
    assert ws["A6"].value == "Primary"
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    frame = load_data()
    results = build_results(frame)
    write_workbook(results)
    validate_workbook()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    columns = [
        "Section",
        "Specification",
        "Total Percent Difference per 0.1 NPP",
        "Total Percent Difference 95 Percent CI Lower",
        "Total Percent Difference 95 Percent CI Upper",
        "Total Probability Value",
        "Total Observations",
        "Food Percent Difference per 0.1 NPP",
        "Food Percent Difference 95 Percent CI Lower",
        "Food Percent Difference 95 Percent CI Upper",
        "Food Probability Value",
        "Food Observations",
    ]
    print(results[columns].to_string(index=False))


if __name__ == "__main__":
    main()
