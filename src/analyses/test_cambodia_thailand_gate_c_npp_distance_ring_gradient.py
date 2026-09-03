#!/usr/bin/env python3
"""Estimate the approved Gate C NPP fixed distance-ring gradient diagnostic.

Plan: report every 0-10, 10-20, 20-40, and 40-60 km post-conflict
drought-sensitivity coefficient relative to the over-60-km frontier controls.
Framework: retain the frozen natural-unit NPP outcome, rainfall-deficit hazard,
2001-2007 versus 2012-2024 window, binary-support overlap weights, grid fixed
effects, climate-cell-by-year fixed effects, and 10 km spatial clustering.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from linearmodels.panel import PanelOLS
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.stats import t as student_t

from estimate_cambodia_thailand_gate_c_npp import (
    BLOCK,
    CELL,
    CLIMATE,
    DRY_RAIN,
    MAIN_PERIOD,
    NPP_NATURAL,
    PANEL,
    POST,
    TARGET,
    WEIGHT,
    YEAR,
    fit_model,
)


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data/exp/experiments/cambodia-thailand-gate-c-npp/distance-ring-gradient"
TIDY = OUT / "gate_c_npp_distance_ring_gradient_tidy.csv"
WORKBOOK = OUT / "gate_c_npp_distance_ring_gradient_summary.xlsx"
METADATA = OUT / "gate_c_npp_distance_ring_gradient_metadata.json"
MAIN_TIDY = (
    ROOT
    / "data/exp/experiments/cambodia-thailand-gate-c-npp/gate_c_npp_main_model_tidy.csv"
)
RING_SUPPORT = (
    ROOT
    / "data/exp/experiment-design/cambodia-thailand-gate-b/gate_b_binary_weight_ring_support.csv"
)

RING = "Conflict Distance Ring"
DOSE = "Continuous Conflict Dose"
RINGS = ["0-10 km", "10-20 km", "20-40 km", "40-60 km"]
REFERENCE = "Over 60 km"


def prepare_sample(frame: pd.DataFrame) -> pd.DataFrame:
    sample = frame.loc[frame[MAIN_PERIOD]].dropna(
        subset=[
            CELL,
            YEAR,
            CLIMATE,
            BLOCK,
            WEIGHT,
            POST,
            RING,
            NPP_NATURAL,
            DRY_RAIN,
        ]
    ).copy()
    if set(sample[RING].unique()) != set([*RINGS, REFERENCE]):
        raise RuntimeError(f"Unexpected distance-ring support: {sorted(sample[RING].unique())}")
    sample["Climate Cell Year"] = (
        sample[CLIMATE].astype("string") + "__" + sample[YEAR].astype("string")
    )
    if sample.duplicated([CELL, YEAR]).any():
        raise RuntimeError("Distance-ring sample is not unique by grid cell and year")
    if not sample[WEIGHT].gt(0).all():
        raise RuntimeError("Analysis weights must be strictly positive")
    return sample


def interaction_columns(
    sample: pd.DataFrame, groups: list[str], group_column: str
) -> tuple[list[str], dict[str, str]]:
    regressors: list[str] = []
    target_terms: dict[str, str] = {}
    for group in groups:
        indicator = sample[group_column].eq(group).astype(float)
        post_term = f"{group}: exposure x post"
        drought_term = f"{group}: exposure x drought"
        target_term = f"{group}: exposure x post x drought"
        sample[post_term] = indicator * sample[POST].astype(float)
        sample[drought_term] = indicator * sample[DRY_RAIN]
        sample[target_term] = indicator * sample[POST].astype(float) * sample[DRY_RAIN]
        regressors.extend([post_term, drought_term, target_term])
        target_terms[group] = target_term
    return regressors, target_terms


def fit_panel(sample: pd.DataFrame, regressors: list[str]):
    panel = sample.set_index([CELL, YEAR]).sort_index()
    other_effects = pd.DataFrame(
        {
            "Climate Cell Year": pd.Categorical(panel["Climate Cell Year"]).codes,
        },
        index=panel.index,
    )
    clusters = pd.DataFrame(
        {"10 km spatial block": pd.Categorical(panel[BLOCK]).codes},
        index=panel.index,
    )
    return PanelOLS(
        panel[NPP_NATURAL].astype(float),
        panel[regressors].astype(float),
        weights=panel[WEIGHT].astype(float),
        entity_effects=True,
        other_effects=other_effects,
        drop_absorbed=True,
        check_rank=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)


def coefficient_row(
    fitted,
    sample: pd.DataFrame,
    specification: str,
    group: str,
    term: str,
) -> dict[str, object]:
    confidence = fitted.conf_int(level=0.95)
    return {
        "Result Type": "coefficient",
        "Specification": specification,
        "Group or Contrast": group,
        "Term": term,
        "Estimate": float(fitted.params[term]),
        "Clustered Standard Error": float(fitted.std_errors[term]),
        "95% CI Lower": float(confidence.loc[term, "lower"]),
        "95% CI Upper": float(confidence.loc[term, "upper"]),
        "p-value": float(fitted.pvalues[term]),
        "Degrees of Freedom": float(fitted.df_resid),
        "Observations": int(fitted.nobs),
        "Grid Cells": int(sample[CELL].nunique()),
        "Climate Cells": int(sample[CLIMATE].nunique()),
        "Spatial Blocks": int(sample[BLOCK].nunique()),
        "Years": int(sample[YEAR].nunique()),
        "Fixed Effects": "grid cell; climate cell x year",
        "Inference": "10 km spatial-block clustered; debiased",
        "Weights": "frozen Gate B binary-support overlap weights",
    }


def contrast_row(
    fitted,
    sample: pd.DataFrame,
    specification: str,
    label: str,
    loadings: dict[str, float],
) -> dict[str, object]:
    vector = np.zeros(len(fitted.params))
    for term, loading in loadings.items():
        vector[fitted.params.index.get_loc(term)] = loading
    estimate = float(vector @ fitted.params.to_numpy())
    variance = float(vector @ fitted.cov.to_numpy() @ vector)
    standard_error = float(np.sqrt(max(variance, 0.0)))
    degrees_freedom = float(fitted.df_resid)
    statistic = estimate / standard_error
    critical = float(student_t.ppf(0.975, degrees_freedom))
    p_value = float(2 * student_t.sf(abs(statistic), degrees_freedom))
    return {
        "Result Type": "contrast",
        "Specification": specification,
        "Group or Contrast": label,
        "Term": "linear contrast of target coefficients",
        "Estimate": estimate,
        "Clustered Standard Error": standard_error,
        "95% CI Lower": estimate - critical * standard_error,
        "95% CI Upper": estimate + critical * standard_error,
        "p-value": p_value,
        "Degrees of Freedom": degrees_freedom,
        "Observations": int(fitted.nobs),
        "Grid Cells": int(sample[CELL].nunique()),
        "Climate Cells": int(sample[CLIMATE].nunique()),
        "Spatial Blocks": int(sample[BLOCK].nunique()),
        "Years": int(sample[YEAR].nunique()),
        "Fixed Effects": "grid cell; climate cell x year",
        "Inference": "10 km spatial-block clustered Wald contrast; debiased",
        "Weights": "frozen Gate B binary-support overlap weights",
    }


def joint_row(fitted, sample: pd.DataFrame, target_terms: dict[str, str]) -> dict[str, object]:
    restriction = np.zeros((len(target_terms), len(fitted.params)))
    for row_number, term in enumerate(target_terms.values()):
        restriction[row_number, fitted.params.index.get_loc(term)] = 1.0
    test = fitted.wald_test(restriction)
    return {
        "Result Type": "joint test",
        "Specification": "Four fixed distance rings",
        "Group or Contrast": "All four ring target coefficients jointly zero",
        "Term": "joint Wald test",
        "Estimate": float(test.stat),
        "Clustered Standard Error": np.nan,
        "95% CI Lower": np.nan,
        "95% CI Upper": np.nan,
        "p-value": float(test.pval),
        "Degrees of Freedom": int(len(target_terms)),
        "Observations": int(fitted.nobs),
        "Grid Cells": int(sample[CELL].nunique()),
        "Climate Cells": int(sample[CLIMATE].nunique()),
        "Spatial Blocks": int(sample[BLOCK].nunique()),
        "Years": int(sample[YEAR].nunique()),
        "Fixed Effects": "grid cell; climate cell x year",
        "Inference": "joint Wald chi-square test using clustered covariance",
        "Weights": "frozen Gate B binary-support overlap weights",
    }


def estimate_four_ring_model(sample: pd.DataFrame):
    frame = sample.copy()
    regressors, target_terms = interaction_columns(frame, RINGS, RING)
    fitted = fit_panel(frame, regressors)
    rows = [
        coefficient_row(
            fitted,
            frame,
            "Four fixed distance rings",
            ring,
            target_terms[ring],
        )
        for ring in RINGS
    ]
    for first, second in zip(RINGS[:-1], RINGS[1:], strict=True):
        rows.append(
            contrast_row(
                fitted,
                frame,
                "Four fixed distance rings",
                f"{first} minus {second}",
                {target_terms[first]: 1.0, target_terms[second]: -1.0},
            )
        )
    rows.append(joint_row(fitted, frame, target_terms))
    return fitted, target_terms, pd.DataFrame(rows)


def estimate_inner_outer_model(sample: pd.DataFrame):
    frame = sample.copy()
    frame["Collapsed Distance Band"] = np.select(
        [
            frame[RING].isin(["0-10 km", "10-20 km"]),
            frame[RING].isin(["20-40 km", "40-60 km"]),
        ],
        ["Inner 0-20 km", "Outer 20-60 km"],
        default="Over 60 km",
    )
    groups = ["Inner 0-20 km", "Outer 20-60 km"]
    regressors, target_terms = interaction_columns(
        frame, groups, "Collapsed Distance Band"
    )
    fitted = fit_panel(frame, regressors)
    rows = [
        coefficient_row(
            fitted,
            frame,
            "Predeclared inner-versus-outer model",
            group,
            target_terms[group],
        )
        for group in groups
    ]
    rows.append(
        contrast_row(
            fitted,
            frame,
            "Predeclared inner-versus-outer model",
            "Inner 0-20 km minus outer 20-60 km",
            {
                target_terms["Inner 0-20 km"]: 1.0,
                target_terms["Outer 20-60 km"]: -1.0,
            },
        )
    )
    return fitted, target_terms, pd.DataFrame(rows)


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def write_workbook(
    four_ring_rows: pd.DataFrame,
    collapsed_rows: pd.DataFrame,
    joint: pd.Series,
    support: pd.DataFrame,
    mean_dose: dict[str, float],
) -> None:
    four = four_ring_rows.set_index("Group or Contrast")
    collapsed = collapsed_rows.set_index("Group or Contrast")
    columns = [*RINGS, "Inner 0-20 km", "Outer 20-60 km", "Inner minus outer"]
    row_sources = {
        **{ring: four.loc[ring] for ring in RINGS},
        "Inner 0-20 km": collapsed.loc["Inner 0-20 km"],
        "Outer 20-60 km": collapsed.loc["Outer 20-60 km"],
        "Inner minus outer": collapsed.loc["Inner 0-20 km minus outer 20-60 km"],
    }
    support_index = support.set_index("Distance Ring")
    rows: list[dict[str, object]] = []
    for label, field, formatter in [
        (
            "Target coefficient",
            "Estimate",
            lambda value, row: f"{value:.5f}{stars(float(row['p-value']))}",
        ),
        (
            "Clustered SE",
            "Clustered Standard Error",
            lambda value, row: f"({value:.5f})",
        ),
        (
            "95% CI",
            None,
            lambda value, row: f"[{row['95% CI Lower']:.5f}, {row['95% CI Upper']:.5f}]",
        ),
        ("p-value", "p-value", lambda value, row: f"{value:.4f}"),
    ]:
        output: dict[str, object] = {"Term / statistic": label}
        for column in columns:
            source = row_sources[column]
            value = source[field] if field else np.nan
            output[column] = formatter(value, source)
        rows.append(output)

    diagnostic_rows = [
        ("Grid cells", "Grid Cells"),
        ("10 km spatial blocks", "Spatial Blocks"),
        ("Cell ESS under frozen weights", "Cell ESS under Binary Support Weights"),
        ("10 km block ESS under frozen weights", "10 km Block ESS under Binary Support Weights"),
        ("Mean continuous conflict dose", None),
    ]
    for label, support_field in diagnostic_rows:
        output = {"Term / statistic": label}
        for column in columns:
            if column in RINGS:
                if support_field:
                    value = float(support_index.loc[column, support_field])
                    output[column] = int(value) if label in {"Grid cells", "10 km spatial blocks"} else f"{value:.1f}"
                else:
                    output[column] = f"{mean_dose[column]:.3f}"
            else:
                output[column] = "—"
        rows.append(output)

    for label, value in [
        ("Joint Wald statistic: all four rings", f"{float(joint['Estimate']):.3f}"),
        ("Joint Wald p-value: all four rings", f"{float(joint['p-value']):.4f}"),
        ("Observations", f"{int(joint['Observations']):,}"),
        ("Grid fixed effects", "Yes"),
        ("Climate-cell x year fixed effects", "Yes"),
    ]:
        output = {"Term / statistic": label, columns[0]: value}
        output.update({column: "" for column in columns[1:]})
        rows.append(output)

    table = pd.DataFrame(rows, columns=["Term / statistic", *columns])
    note = pd.DataFrame(
        [
            {
                "Term / statistic": "Notes",
                columns[0]: (
                    "Outcome: annual NPP anomaly (kg C/m2). Target: distance-band "
                    "exposure x post-2011 x rainfall-deficit intensity. Reference: "
                    "controls over 60 km from every candidate event. Parentheses: "
                    "10 km block-clustered SEs. * p<0.10, ** p<0.05, *** p<0.01."
                ),
                **{column: "" for column in columns[1:]},
            }
        ]
    )
    table = pd.concat([table, note], ignore_index=True)
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        table.to_excel(writer, index=False, sheet_name="Distance Gradient")
        sheet = writer.book["Distance Gradient"]
        sheet.freeze_panes = "B2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions["A"].width = 40
        for column_number in range(2, len(columns) + 2):
            letter = sheet.cell(1, column_number).column_letter
            sheet.column_dimensions[letter].width = 22
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        note_row = sheet.max_row
        sheet.merge_cells(
            start_row=note_row,
            start_column=2,
            end_row=note_row,
            end_column=len(columns) + 1,
        )
        sheet.cell(note_row, 2).font = Font(size=9, italic=True, color="404040")
        sheet.cell(note_row, 2).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[note_row].height = 38
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_LEGAL
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    columns = [
        CELL,
        YEAR,
        CLIMATE,
        BLOCK,
        WEIGHT,
        POST,
        MAIN_PERIOD,
        RING,
        DOSE,
        NPP_NATURAL,
        DRY_RAIN,
    ]
    sample = prepare_sample(pd.read_parquet(PANEL, columns=columns))
    support = pd.read_csv(RING_SUPPORT)
    mean_dose = (
        sample.drop_duplicates(CELL).groupby(RING)[DOSE].mean().to_dict()
    )

    print("Reproducing the frozen continuous-dose coefficient for QA", flush=True)
    continuous_table, _ = fit_model(
        sample,
        NPP_NATURAL,
        DRY_RAIN,
        "Distance-ring QA: continuous-dose reproduction",
    )
    reproduced_continuous = float(
        continuous_table.loc[continuous_table["Term"].eq(TARGET), "Estimate"].iloc[0]
    )
    opened_main = pd.read_csv(MAIN_TIDY)
    opened_continuous = float(
        opened_main.loc[
            opened_main["Specification"].eq(
                "Primary: natural-unit NPP x dry-rainfall intensity"
            )
            & opened_main["Term"].eq(TARGET),
            "Estimate",
        ].iloc[0]
    )
    if not np.isclose(reproduced_continuous, opened_continuous, atol=1e-12, rtol=0):
        raise RuntimeError(
            "Distance-ring sample does not reproduce the opened continuous-dose estimate: "
            f"{reproduced_continuous} versus {opened_continuous}"
        )

    print("Estimating four fixed distance rings", flush=True)
    four_fitted, four_terms, four_results = estimate_four_ring_model(sample)
    print("Estimating predeclared inner-versus-outer contrast", flush=True)
    collapsed_fitted, collapsed_terms, collapsed_results = estimate_inner_outer_model(sample)
    tidy = pd.concat([four_results, collapsed_results], ignore_index=True)
    tidy.to_csv(TIDY, index=False)

    four_coefficients = four_results.loc[four_results["Result Type"].eq("coefficient")]
    joint = four_results.loc[four_results["Result Type"].eq("joint test")].iloc[0]
    write_workbook(
        four_coefficients,
        collapsed_results,
        joint,
        support,
        mean_dose,
    )

    ring_estimates = four_coefficients.set_index("Group or Contrast")["Estimate"]
    ordered_toward_zero = bool(
        ring_estimates["0-10 km"]
        <= ring_estimates["10-20 km"]
        <= ring_estimates["20-40 km"]
        <= ring_estimates["40-60 km"]
        <= 0
    )
    adjacent = four_results.loc[
        four_results["Result Type"].eq("contrast"),
        [
            "Group or Contrast",
            "Estimate",
            "Clustered Standard Error",
            "95% CI Lower",
            "95% CI Upper",
            "p-value",
        ],
    ].to_dict(orient="records")
    collapsed_contrast = collapsed_results.loc[
        collapsed_results["Result Type"].eq("contrast")
    ].iloc[0]
    metadata = {
        "status": "completed human-approved Gate C NPP fixed distance-ring gradient",
        "human_approval_record": "MILI-D-20260822-021",
        "reference_group": REFERENCE,
        "rings_reported_without_selection": RINGS,
        "sample_years": sorted(sample[YEAR].unique().astype(int).tolist()),
        "outcome": NPP_NATURAL,
        "hazard": DRY_RAIN,
        "fixed_effects": "grid cell and climate-cell by year",
        "cluster": "10 km spatial block; debiased covariance",
        "weights": "frozen Gate B binary-support overlap weights",
        "qa_reproduced_continuous_dose_estimate": reproduced_continuous,
        "qa_opened_continuous_dose_estimate": opened_continuous,
        "qa_exact_continuous_dose_reproduction": True,
        "four_ring_joint_wald_statistic": float(joint["Estimate"]),
        "four_ring_joint_wald_p_value": float(joint["p-value"]),
        "point_estimates_monotone_negative_toward_zero": ordered_toward_zero,
        "adjacent_ring_contrasts": adjacent,
        "inner_minus_outer_estimate": float(collapsed_contrast["Estimate"]),
        "inner_minus_outer_p_value": float(collapsed_contrast["p-value"]),
        "ring_support_warning": (
            "The frozen binary-support weights balance all exposed cells against controls, "
            "not each ring separately. The 0-10 km and 10-20 km ring-specific 10 km block "
            "ESS values are 4.1 and 7.8, so ring coefficients are functional-form "
            "diagnostics rather than separately powered causal estimates."
        ),
        "created_at_utc": datetime.now(UTC).isoformat(),
    }
    METADATA.write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    display = pd.concat(
        [four_coefficients, collapsed_results], ignore_index=True
    )[[
        "Specification",
        "Group or Contrast",
        "Estimate",
        "Clustered Standard Error",
        "95% CI Lower",
        "95% CI Upper",
        "p-value",
    ]]
    print("\nDistance-gradient estimates")
    print(display.to_string(index=False))
    print(
        f"\nJoint Wald p-value: {float(joint['p-value']):.6f}; "
        f"monotone negative toward zero: {ordered_toward_zero}"
    )
    print(f"Saved {TIDY.relative_to(ROOT)}")
    print(f"Saved {WORKBOOK.relative_to(ROOT)}")
    print(f"Saved {METADATA.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
