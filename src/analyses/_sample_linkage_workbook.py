"""Shared workbook writer for the split analytical-sample and survey-wave tables."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def write_single_table(
    table: pd.DataFrame,
    *,
    title: str,
    sheet_name: str,
    output: Path,
    notes: list[str],
    highlighted_rows: set[str] | None = None,
) -> None:
    highlighted_rows = highlighted_rows or set()
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    pale_green = "E2F0D9"
    pale_blue = "DDEBF7"
    light_gray = "F3F3F3"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    medium = Side(style="medium", color="7F9DB9")
    vertical = Side(style="thin", color="D7E0E6")

    ws.merge_cells("A1:J1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    header_row = 3
    for column, label in enumerate(table.columns, start=1):
        cell = ws.cell(header_row, column, label)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=medium,
            bottom=medium,
            left=vertical if column in {5, 7, 9} else None,
        )
    ws.row_dimensions[header_row].height = 42

    for row_offset, (_, values) in enumerate(table.iterrows(), start=0):
        row = header_row + 1 + row_offset
        label = str(values.iloc[0])
        for column, value in enumerate(values, start=1):
            if pd.isna(value):
                value = None
            cell = ws.cell(row, column, value)
            cell.font = Font(name="Times New Roman", size=9.5)
            cell.alignment = Alignment(
                horizontal="center" if column in {4, 6} else "left",
                vertical="center",
                wrap_text=True,
                indent=1 if column in {5, 7, 8, 9, 10} else 0,
            )
            cell.border = Border(
                bottom=thin,
                left=vertical if column in {5, 7, 9} else None,
            )
            if row_offset % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=light_gray)
        ws.cell(row, 4).number_format = "#,##0"
        ws.cell(row, 6).number_format = "0.0%"
        if label in highlighted_rows:
            fill = pale_blue if "cross-fit" in label.lower() else pale_green
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(name="Times New Roman", size=9.5, bold=True)
        ws.row_dimensions[row].height = 42

    last_data_row = header_row + len(table)
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row = notes_start + offset
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell = ws.cell(row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.8, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 20

    widths = [32, 19, 22, 15, 23, 15, 32, 25, 42, 33]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A{header_row}:J{last_data_row}"
    ws.sheet_view.zoomScale = 60
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.18
    ws.page_margins.right = 0.18
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.print_area = f"A1:J{notes_start + len(notes) - 1}"

    wb.properties.title = title
    wb.properties.subject = "Climate-welfare analytical support and linkage audit"
    wb.properties.creator = "Mike Li"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def validate_workbook(output: Path, sheet_name: str, expected_rows: int) -> None:
    wb = load_workbook(output, data_only=False)
    assert wb.sheetnames == [sheet_name]
    ws = wb[sheet_name]
    assert ws.max_column == 10
    assert ws.max_row >= expected_rows + 3
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(
                    ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                )
