#!/usr/bin/env python3
"""Annual Cropland-NPP Coverage and Quality.

Plan: Document annual strict-cropland and inclusive-agriculture pixel support,
NPP coverage, and quality gates for the 5 km Stage-1 village panel.
Framework: AnaSOP Sections 4-7 Stage-1 NPP measurement and workflow step 2.
This table audits the ecological outcome; it does not estimate climate slopes.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "data/processed/cses_public_village_cropland_npp_absolute_climate_panel_preprocessed.parquet"
FIGURE_EVIDENCE = ROOT / "data/exp/analysis/climate-npp/cropland-definition-and-npp-quality-support"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/annual-cropland-npp-coverage-and-quality"
OUTPUT = ROOT / "data/results/tables/Table_annual_cropland_npp_coverage_and_quality.xlsx"

ID = "National Village Point ID"
YEAR = "Year"
STRICT_CANDIDATE = "Strict-Cropland Candidate 500m Pixel Count"
STRICT_VALID = "Strict-Cropland Valid NPP 500m Pixel Count"
STRICT_NPP = "Annual Strict-Cropland Mean NPP kg C per m2"
STRICT_BAD_QC = "Strict-Cropland Original NPP QC Above 100 Pixel Count"
INCLUSIVE_CANDIDATE = "Inclusive-Agriculture Candidate 500m Pixel Count"
INCLUSIVE_VALID = "Inclusive-Agriculture Valid NPP 500m Pixel Count"
INCLUSIVE_NPP = "Annual Inclusive-Agriculture Mean NPP kg C per m2"
INCLUSIVE_BAD_QC = "Inclusive-Agriculture Original NPP QC Above 100 Pixel Count"


def build_table() -> tuple[pd.DataFrame, dict[str, float]]:
    columns = [
        ID,
        YEAR,
        STRICT_CANDIDATE,
        STRICT_VALID,
        STRICT_NPP,
        STRICT_BAD_QC,
        INCLUSIVE_CANDIDATE,
        INCLUSIVE_VALID,
        INCLUSIVE_NPP,
        INCLUSIVE_BAD_QC,
    ]
    frame = pd.read_parquet(INPUT, columns=columns).copy()
    table = (
        frame.groupby(YEAR, as_index=False)
        .agg(
            **{
                "Strict candidate pixels": (STRICT_CANDIDATE, "sum"),
                "Strict valid pixels": (STRICT_VALID, "sum"),
                "Strict mean NPP": (STRICT_NPP, "mean"),
                "Strict no-NPP villages": (STRICT_NPP, lambda values: int(values.isna().sum())),
                "Inclusive candidate pixels": (INCLUSIVE_CANDIDATE, "sum"),
                "Inclusive valid pixels": (INCLUSIVE_VALID, "sum"),
                "Inclusive mean NPP": (INCLUSIVE_NPP, "mean"),
                "Inclusive no-NPP villages": (INCLUSIVE_NPP, lambda values: int(values.isna().sum())),
            }
        )
        .rename(columns={YEAR: "Year"})
    )
    table.insert(
        3,
        "Strict valid share",
        table["Strict valid pixels"] / table["Strict candidate pixels"],
    )
    table.insert(
        8,
        "Inclusive valid share",
        table["Inclusive valid pixels"] / table["Inclusive candidate pixels"],
    )
    table = table[
        [
            "Year",
            "Strict candidate pixels",
            "Strict valid pixels",
            "Strict valid share",
            "Strict mean NPP",
            "Strict no-NPP villages",
            "Inclusive candidate pixels",
            "Inclusive valid pixels",
            "Inclusive valid share",
            "Inclusive mean NPP",
            "Inclusive no-NPP villages",
        ]
    ]
    summary = {
        "villages": float(frame[ID].nunique()),
        "village_years": float(len(frame)),
        "strict_candidate_pixels": float(frame[STRICT_CANDIDATE].sum()),
        "strict_valid_pixels": float(frame[STRICT_VALID].sum()),
        "strict_valid_share": float(frame[STRICT_VALID].sum() / frame[STRICT_CANDIDATE].sum()),
        "strict_no_npp_village_years": float(frame[STRICT_NPP].isna().sum()),
        "inclusive_candidate_pixels": float(frame[INCLUSIVE_CANDIDATE].sum()),
        "inclusive_valid_pixels": float(frame[INCLUSIVE_VALID].sum()),
        "inclusive_valid_share": float(frame[INCLUSIVE_VALID].sum() / frame[INCLUSIVE_CANDIDATE].sum()),
        "inclusive_no_npp_village_years": float(frame[INCLUSIVE_NPP].isna().sum()),
        "strict_original_qc_above_100_pixels": float(frame[STRICT_BAD_QC].sum()),
        "inclusive_original_qc_above_100_pixels": float(frame[INCLUSIVE_BAD_QC].sum()),
    }
    assert table["Year"].tolist() == list(range(2001, 2022))
    assert int(summary["villages"]) == 2_966
    assert int(summary["village_years"]) == 62_286
    assert int(summary["strict_no_npp_village_years"]) == 3_151
    assert int(summary["inclusive_no_npp_village_years"]) == 3_045
    assert int(summary["strict_candidate_pixels"] - summary["strict_valid_pixels"]) == int(
        summary["strict_original_qc_above_100_pixels"]
    )
    assert int(summary["inclusive_candidate_pixels"] - summary["inclusive_valid_pixels"]) == int(
        summary["inclusive_original_qc_above_100_pixels"]
    )

    figure_summary = pd.read_csv(
        FIGURE_EVIDENCE / "cropland_definition_npp_quality_summary.csv", index_col=0
    )
    assert int(figure_summary.loc[STRICT_NPP, "count"]) == int(
        summary["village_years"] - summary["strict_no_npp_village_years"]
    )

    EVIDENCE.mkdir(parents=True, exist_ok=True)
    table.to_csv(EVIDENCE / "annual_cropland_npp_coverage_and_quality.csv", index=False)
    pd.DataFrame([summary]).to_csv(EVIDENCE / "cropland_npp_coverage_quality_summary.csv", index=False)
    return table, summary


def write_workbook(table: pd.DataFrame, summary: dict[str, float]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Annual NPP Quality"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    strict_blue = "3F7CAC"
    inclusive_teal = "3A9D8F"
    white = "FFFFFF"
    light = "F4F6F7"
    thin = Side(style="thin", color="C8D5DE")
    medium = Side(style="medium", color="7F9DB9")

    ws.merge_cells("A1:K1")
    ws["A1"] = "Annual Cropland-NPP Coverage and Quality"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 29

    ws.merge_cells("A3:A4")
    ws["A3"] = "Year"
    ws.merge_cells("B3:F3")
    ws["B3"] = "Strict cropland"
    ws.merge_cells("G3:K3")
    ws["G3"] = "Inclusive agriculture"
    for cell in [ws["A3"], ws["B3"], ws["G3"]]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=10.1, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=medium, bottom=medium)
    ws.row_dimensions[3].height = 24

    subheaders = [
        "Candidate pixels",
        "Valid pixels",
        "Valid share",
        "Mean NPP",
        "No-NPP villages",
        "Candidate pixels",
        "Valid pixels",
        "Valid share",
        "Mean NPP",
        "No-NPP villages",
    ]
    for column, header in enumerate(subheaders, start=2):
        cell = ws.cell(4, column, header)
        fill = strict_blue if column <= 6 else inclusive_teal
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Times New Roman", size=8.8, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[4].height = 34

    for row_offset, record in enumerate(table.itertuples(index=False, name=None), start=5):
        for column, value in enumerate(record, start=1):
            cell = ws.cell(row_offset, column, value)
            cell.font = Font(name="Times New Roman", size=8.7, bold=column == 1)
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = Border(bottom=thin)
            if row_offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            if column == 1:
                cell.number_format = "0"
            elif column in {2, 3, 6, 7, 8, 11}:
                cell.number_format = "#,##0"
            elif column in {4, 9}:
                # Store computed values, rather than workbook formulas, so the
                # table renders consistently in Excel, LibreOffice, and previews.
                cell.number_format = "0.00%"
            elif column in {5, 10}:
                cell.number_format = "0.000"
        ws.row_dimensions[row_offset].height = 23

    summary_row = 26
    ws.cell(summary_row, 1, "All years")
    ws.cell(summary_row, 2, int(summary["strict_candidate_pixels"]))
    ws.cell(summary_row, 3, int(summary["strict_valid_pixels"]))
    ws.cell(summary_row, 4, float(summary["strict_valid_share"]))
    ws.cell(summary_row, 5, float(
        pd.read_parquet(INPUT, columns=[STRICT_NPP])[STRICT_NPP].mean()
    ))
    ws.cell(summary_row, 6, int(summary["strict_no_npp_village_years"]))
    ws.cell(summary_row, 7, int(summary["inclusive_candidate_pixels"]))
    ws.cell(summary_row, 8, int(summary["inclusive_valid_pixels"]))
    ws.cell(summary_row, 9, float(summary["inclusive_valid_share"]))
    ws.cell(summary_row, 10, float(
        pd.read_parquet(INPUT, columns=[INCLUSIVE_NPP])[INCLUSIVE_NPP].mean()
    ))
    ws.cell(summary_row, 11, int(summary["inclusive_no_npp_village_years"]))
    for column in range(1, 12):
        cell = ws.cell(summary_row, column)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=8.8, bold=True, color=white)
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = Border(top=medium, bottom=medium)
        if column in {1, 2, 3, 6, 7, 8, 11}:
            cell.number_format = "#,##0"
        elif column in {4, 9}:
            cell.number_format = "0.00%"
        elif column in {5, 10}:
            cell.number_format = "0.000"
    ws["A26"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[summary_row].height = 24

    note_row = 28
    notes = [
        f"Notes: The panel contains {int(summary['villages']):,} village points in every year from 2001 to 2021. Mean NPP is kg C m⁻² year⁻¹ within 5 km village buffers.",
        "Valid share equals valid annual NPP pixels divided by same-year candidate pixels. Candidate-minus-valid pixels exactly equal pixels whose original NPP QC exceeded 100; these pixels were excluded from NPP means.",
        "No-NPP villages almost entirely have zero candidate pixels under that land-cover definition. Missing values remain missing; the primary model uses strict cropland and inclusive agriculture is a sensitivity definition.",
    ]
    for note in notes:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=11)
        cell = ws.cell(note_row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.0, italic=True, color="3F4B52")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[note_row].height = 22
        note_row += 1

    widths = [10, 17, 17, 14, 15, 16, 17, 17, 14, 15, 16]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 75
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.20
    ws.page_margins.top = ws.page_margins.bottom = 0.22
    ws.print_area = "A1:K30"
    wb.properties.title = "Annual Cropland-NPP Coverage and Quality"
    wb.properties.subject = "Stage-1 annual NPP pixel support audit"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Annual NPP Quality"]
    ws = wb["Annual NPP Quality"]
    assert ws.max_row == 30
    assert ws.max_column == 11
    assert not ws.merged_cells.ranges
    assert ws["A5"].value == 2001
    assert ws["A25"].value == 2021
    assert abs(float(ws["D5"].value) - float(ws["C5"].value) / float(ws["B5"].value)) < 1e-12
    assert abs(float(ws["I25"].value) - float(ws["H25"].value) / float(ws["G25"].value)) < 1e-12
    assert not any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in ws.iter_rows()
        for cell in row
    )
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    table, summary = build_table()
    write_workbook(table, summary)
    validate_workbook()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(table.to_string(index=False))
    print(summary)


if __name__ == "__main__":
    main()
