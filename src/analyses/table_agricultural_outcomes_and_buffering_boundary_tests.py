#!/usr/bin/env python3
"""Build the standalone Appendix agricultural and buffering-boundary table."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from table_failed_pathways_and_buffering_tests import (
    COLUMNS,
    agricultural_rows,
    buffering_rows,
)


ROOT = Path(__file__).resolve().parents[2]
ANALYSIS_OUTPUT = (
    ROOT
    / "data/exp/analysis/climate-welfare/agricultural-buffering-boundary-tests/table_rows.csv"
)
OUTPUT = (
    ROOT
    / "data/exp/legacy-results/tables/Table_agricultural_outcomes_and_buffering_boundary_tests.xlsx"
)


def main() -> None:
    table = pd.DataFrame(agricultural_rows() + buffering_rows(), columns=COLUMNS)
    assert len(table) == 11
    ANALYSIS_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(ANALYSIS_OUTPUT, index=False)

    def one_row(family: str, outcome: str | None = None, exposure: str | None = None) -> pd.Series:
        selected = table.loc[table["Evidence family"].eq(family)]
        if outcome is not None:
            selected = selected.loc[selected["Outcome"].eq(outcome)]
        if exposure is not None:
            selected = selected.loc[selected["Exposure / modifier"].str.contains(exposure, regex=False)]
        if len(selected) != 1:
            raise ValueError(
                f"Expected one row for family={family!r}, outcome={outcome!r}, exposure={exposure!r}; found {len(selected)}"
            )
        return selected.iloc[0]

    wb = Workbook()
    ws = wb.active
    ws.title = "Agriculture buffers"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    header_fill = PatternFill("solid", fgColor="E5E5E5")
    agriculture_fill = PatternFill("solid", fgColor="F7F1E8")
    buffering_fill = PatternFill("solid", fgColor="F1ECF6")

    def add(row: int, column: int, value, *, bold: bool = False, fill=None) -> None:
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(name="Times New Roman", size=9.2, bold=bold)
        cell.alignment = Alignment(
            horizontal="left" if column == 1 else "center",
            vertical="center",
            wrap_text=True,
        )
        if fill is not None:
            cell.fill = fill
        if isinstance(value, (float, np.floating)):
            cell.number_format = "0.000"

    def merged(row: int, start: int, end: int, value, *, bold: bool = False, fill=None) -> None:
        if end > start:
            ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        add(row, start, value, bold=bold, fill=fill)
        if fill is not None:
            for column in range(start, end + 1):
                ws.cell(row=row, column=column).fill = fill

    ws.merge_cells("A1:I1")
    ws["A1"] = "Agricultural Outcomes and Buffering Boundary Tests"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    specs = [
        (2, 3, "Agricultural participation", one_row("Agricultural survey pathway", outcome="Agricultural participation", exposure="May–October rainfall"), agriculture_fill),
        (4, 5, "Crop production / ha", one_row("Agricultural survey pathway", outcome="Real crop production per cultivated ha (asinh)", exposure="May–October rainfall"), agriculture_fill),
        (6, 6, "Heat × irrigation", one_row("Buffering boundary test", exposure="Irrigation"), buffering_fill),
        (7, 7, "Heat × road access", one_row("Buffering boundary test", exposure="Historical Road Access"), buffering_fill),
        (8, 9, "Heat × connectivity", one_row("Buffering boundary test", exposure="Baseline Settlement Connectivity"), buffering_fill),
    ]
    add(3, 1, "Statistic", bold=True, fill=header_fill)
    ws.cell(3, 1).border = Border(top=medium, bottom=thin)
    for start, end, label, _, _ in specs:
        merged(3, start, end, label, bold=True, fill=header_fill)
        for column in range(start, end + 1):
            ws.cell(3, column).border = Border(top=medium, bottom=thin)
    ws.row_dimensions[3].height = 36

    rows = [
        (4, "Specification"),
        (6, "Rainfall coefficient"), (7, "95% interval"), (8, "Raw p-value"),
        (10, "Onset coefficient"), (11, "95% interval"), (12, "Raw p-value"),
        (14, "Dry-spell coefficient"), (15, "95% interval"), (16, "Raw p-value"),
        (18, "Relative hot-dry coefficient (legacy)"), (19, "95% interval"), (20, "Raw p-value"),
        (22, "Heat × modifier coefficient"), (23, "95% interval"), (24, "Raw p-value"),
        (25, "Holm-adjusted p-value"), (27, "Sample"),
        (28, "Fixed effects / inference"), (29, "Gate status"), (30, "Interpretation limit"),
    ]
    for row, label in rows:
        add(row, 1, label, bold=label in {"Gate status", "Interpretation limit"})
        ws.row_dimensions[row].height = 23 if row not in {28, 29, 30} else 32

    exposure_rows = [
        ("May–October rainfall", 6, 7, 8),
        ("Wet-season onset", 10, 11, 12),
        ("Longest dry-spell", 14, 15, 16),
        ("Relative hot-dry", 18, 19, 20),
    ]
    for start, end, _, representative, fill in specs:
        is_buffer = representative["Evidence family"] == "Buffering boundary test"
        values: dict[int, object] = {
            4: representative["Specification"], 27: representative["Sample"],
            28: representative["Fixed effects / inference"], 29: representative["Gate status"],
            30: representative["Interpretation limit"],
        }
        if is_buffer:
            values.update(
                {
                    22: representative["Estimate / metric"], 23: representative["95% interval"],
                    24: representative["Raw p-value"],
                    25: str(representative["Adjusted / gate metric"]).split(";")[0].replace("Holm p=", ""),
                }
            )
        else:
            outcome = str(representative["Outcome"])
            for exposure, estimate_row, interval_row, probability_row in exposure_rows:
                record = one_row("Agricultural survey pathway", outcome=outcome, exposure=exposure)
                values.update(
                    {
                        estimate_row: record["Estimate / metric"],
                        interval_row: record["95% interval"],
                        probability_row: record["Raw p-value"],
                    }
                )
        for row, _ in rows:
            merged(row, start, end, values.get(row, "—"), fill=fill)
    for column in range(1, 10):
        ws.cell(30, column).border = Border(bottom=medium)

    notes = [
        "Notes: Agricultural rows use CSES weights and 0.75° spatial-block-clustered uncertainty; relative hot-dry exposure is a legacy diagnostic and is not current absolute-heat evidence.",
        "The three heat-buffer interactions use Candidate B, ≥35°C, 5 km exposure. Holm adjustment is applied across the three tests; none supports a protective or causal adaptation claim.",
    ]
    for row, note in enumerate(notes, start=32):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = ws.cell(row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.5, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 18

    for column, width in enumerate([36] + [20] * 8, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.18
    ws.page_margins.right = 0.18
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.print_area = "A1:I33"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == ["Agriculture buffers"]
    assert len(table) == 11
    assert table.loc[
        table["Evidence family"].eq("Buffering boundary test"), "Adjusted / gate metric"
    ].str.contains("Holm p=1.000").all()
    for row in check.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not any(error in cell.value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?"))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
