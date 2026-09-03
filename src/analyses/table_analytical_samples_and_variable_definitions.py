"""Create the one-sheet analytical-sample and variable-definition workbook."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
STAGE1 = ROOT / "data/processed/cses_public_village_cropland_npp_absolute_climate_panel_preprocessed.parquet"
HOUSEHOLDS = ROOT / "data/processed/cses_household_cropland_npp_analysis_preprocessed.parquet"
REGIONS = ROOT / "data/processed/outcome_blind_spatial_regions_preprocessed.parquet"
OUTPUT = ROOT / "data/results/tables/Table_analytical_samples_and_variable_definitions.xlsx"
EVIDENCE_DIR = ROOT / "data/exp/analysis/climate-npp/analytical-samples-and-variable-definitions"


DETAIL_COLUMNS = [
    "Variable or analytical sample",
    "Definition / construction",
    "Analytical role and timing",
    "Unit",
    "Non-missing observations",
    "Spatial or household support",
    "Missingness / coverage",
    "Linkage / final status",
]

DISPLAY_COLUMNS = [
    "Item",
    "Construction / inclusion rule",
    "Role",
    "Timing",
    "Unit",
    "N",
    "Places / waves",
    "Status",
]


def pct(numerator: int, denominator: int) -> str:
    return f"{100 * numerator / denominator:.1f}%"


def build_rows() -> list[dict[str, object]]:
    stage1 = pd.read_parquet(STAGE1)
    households = pd.read_parquet(HOUSEHOLDS)
    regions = pd.read_parquet(REGIONS)
    main = households.loc[households["Main Linked Sample"].fillna(False).astype(bool)].copy()

    strict = "Annual Strict-Cropland Mean NPP kg C per m2"
    inclusive = "Annual Inclusive-Agriculture Mean NPP kg C per m2"
    total_complete = main["Stage 2 Total Consumption Complete Case"].fillna(False).astype(bool)
    food_complete = main["Stage 2 Food Consumption Complete Case"].fillna(False).astype(bool)
    point_matched = main["Public Village Point Matched"].fillna(False).astype(bool)
    prior_npp = main["Prior-Year Strict-Cropland NPP"].notna()

    n_release = len(households)
    n_main = len(main)
    n_stage1 = int(stage1[strict].notna().sum())
    n_strict_villages = int(stage1.loc[stage1[strict].notna(), "National Village Point ID"].nunique())
    n_inclusive = int(stage1[inclusive].notna().sum())
    n_inclusive_villages = int(stage1.loc[stage1[inclusive].notna(), "National Village Point ID"].nunique())
    n_total = int(total_complete.sum())
    n_food = int(food_complete.sum())
    total_villages = int(main.loc[total_complete, "National Village Point ID"].nunique())
    food_villages = int(main.loc[food_complete, "National Village Point ID"].nunique())
    total_outcome = int(main["Real 2021 Annual Total Consumption per Capita Riels"].notna().sum())
    food_outcome = int(main["Real 2021 Annual Food Consumption per Capita Riels"].notna().sum())
    dependency = int(main["Household Dependency Ratio"].notna().sum())
    school_observed = int(main["Household Head Ever Attended School"].notna().sum())
    school_yes = int((main["Household Head Ever Attended School"] == 1).sum())
    weight_observed = int(main["Household Survey Weight"].notna().sum())
    valid_share = int(stage1["Strict-Cropland Valid NPP Pixel Share"].notna().sum())
    npp_qc = "Mean Strict-Cropland Recoded NPP QC Filled Days Percent"
    n_qc = int(stage1[npp_qc].notna().sum())
    qc_p50 = float(stage1[npp_qc].dropna().quantile(0.50))
    qc_p75 = float(stage1[npp_qc].dropna().quantile(0.75))
    region_villages = int(regions["SKATER Region ID"].notna().sum())

    rows: list[dict[str, object]] = []

    def add(section: str, *values: object) -> None:
        rows.append({"Section": section, **dict(zip(DETAIL_COLUMNS, values, strict=True))})

    add(
        "A. Analytical samples",
        "Harmonised CSES household release",
        "Ten available survey waves; 2004 retained for construction diagnostics but lacks a public village-point link.",
        "Source universe before spatial-link eligibility screening",
        "household-wave",
        n_release,
        "10 waves; 2004–2021",
        "No row deletion at release-audit stage",
        "Source universe; 2004 is not used in linked regressions",
    )
    add(
        "A. Analytical samples",
        "Main 2007–2021 household analysis frame",
        "Nine waves eligible for deterministic linkage to the national public village-point frame.",
        "Stage-2 source frame before complete-case restrictions",
        "household-wave",
        n_main,
        "9 waves; interview years 2007–2021",
        f"{pct(n_main, n_release)} of the harmonised release",
        "Final source frame for household analyses",
    )
    add(
        "A. Analytical samples",
        "Stage-1 primary complete sample",
        "Village-years with strict-cropland NPP and all four climate shocks plus annual rainfall.",
        "Climate → NPP fixed-effects models, same calendar year t",
        "village-year",
        n_stage1,
        f"{n_strict_villages:,} villages; 2001–2021",
        f"{pct(n_stage1, len(stage1))} of {len(stage1):,} candidate village-years",
        "Primary stage-1 estimation sample",
    )
    add(
        "A. Analytical samples",
        "Stage-2 total-consumption complete sample",
        "Positive survey weight, public-point match, prior-year strict-cropland NPP, primary outcome, and household composition controls.",
        "Prior-year NPP → real per-capita total consumption",
        "household-wave",
        n_total,
        f"{total_villages:,} villages; 9 waves",
        f"{pct(n_total, n_main)} of the 2007–2021 frame",
        "Primary stage-2 estimation sample",
    )
    add(
        "A. Analytical samples",
        "Stage-2 food-consumption complete sample",
        "Same linkage, timing, survey-weight, and household-control rules as the primary sample, using food consumption.",
        "Prior-year NPP → real per-capita food consumption",
        "household-wave",
        n_food,
        f"{food_villages:,} villages; 9 waves",
        f"{pct(n_food, n_main)} of the 2007–2021 frame",
        "Secondary stage-2 estimation sample",
    )

    climate_support = f"{len(stage1):,} village-years; {stage1['National Village Point ID'].nunique():,} villages"
    add(
        "B. Stage-1 climate and ecological variables",
        "Annual Strict-Cropland Mean NPP",
        "Mean annual MODIS NPP among same-year strict-cropland 500 m pixels within the 5 km village buffer.",
        "Primary ecological outcome in year t",
        "kg C m⁻² year⁻¹",
        n_stage1,
        f"{n_strict_villages:,} villages; 2001–2021",
        f"{pct(n_stage1, len(stage1))} observed; no imputation",
        "Primary stage-1 outcome; also lagged stage-2 exposure",
    )
    add(
        "B. Stage-1 climate and ecological variables",
        "Annual Inclusive-Agriculture Mean NPP",
        "Mean annual MODIS NPP including strict cropland and cropland–natural-vegetation mosaic pixels within 5 km.",
        "Alternative land-cover definition in year t",
        "kg C m⁻² year⁻¹",
        n_inclusive,
        f"{n_inclusive_villages:,} villages; 2001–2021",
        f"{pct(n_inclusive, len(stage1))} observed; no imputation",
        "Sensitivity variable",
    )
    add(
        "B. Stage-1 climate and ecological variables",
        "Annual Heat Days at or Above 35 °C",
        "Count of days in a calendar year with daily maximum temperature ≥35 °C, averaged across 1 km cells in the 5 km buffer.",
        "Primary heat exposure in year t; coefficient reported per 10 days",
        "days year⁻¹",
        len(stage1),
        climate_support,
        "100.0% observed",
        "Primary stage-1 exposure",
    )
    add(
        "B. Stage-1 climate and ecological variables",
        "Annual Heat Degree-Days Above 35 °C",
        "Annual sum of daily maximum-temperature exceedance above 35 °C, averaged across buffer cells.",
        "Alternative heat-intensity exposure in year t; reported per 10 degree-days",
        "°C-days year⁻¹",
        len(stage1),
        climate_support,
        "100.0% observed",
        "Alternative heat specification; not entered with heat days",
    )
    add(
        "B. Stage-1 climate and ecological variables",
        "Annual Maximum Consecutive Five-Day Precipitation (Rx5day)",
        "Largest rolling five-day precipitation total in each calendar year, averaged across buffer cells.",
        "Extreme-rainfall exposure in year t; coefficient reported per 10 mm",
        "mm",
        len(stage1),
        climate_support,
        "100.0% observed",
        "Primary stage-1 exposure",
    )
    add(
        "B. Stage-1 climate and ecological variables",
        "Annual Maximum Consecutive Dry Days Below 1 mm",
        "Longest annual run of days with precipitation <1 mm, averaged across buffer cells.",
        "Dry-spell exposure in year t; coefficient reported per 10 days",
        "days",
        len(stage1),
        climate_support,
        "100.0% observed",
        "Primary stage-1 exposure",
    )
    add(
        "B. Stage-1 climate and ecological variables",
        "Annual Precipitation Total",
        "Sum of daily precipitation within the calendar year, averaged across buffer cells.",
        "Annual-water control in year t; coefficient reported per 100 mm",
        "mm year⁻¹",
        len(stage1),
        climate_support,
        "100.0% observed",
        "Included with Rx5day and dry-spell measures",
    )
    add(
        "B. Stage-1 climate and ecological variables",
        "NPP Pixel Support",
        "Candidate and valid 500 m pixel counts, valid-pixel share, and QC fill-days; original NPP QC values >100 are recoded to 0 before aggregation.",
        "Quality gate and sensitivity diagnostics",
        "count; share; percent",
        valid_share,
        f"{stage1['National Village Point ID'].nunique():,} candidate villages; 2001–2021",
        f"Strict-cropland valid-share field observed for {pct(valid_share, len(stage1))} of village-years",
        "Materialised; no missing-value imputation",
    )
    add(
        "B. Stage-1 climate and ecological variables",
        "Mean Annual Strict-Cropland NPP Filled-Days Percentage",
        "Mean percentage of growing-season days supported by gap-filled FPAR or LAI inputs among annual strict-cropland NPP pixels in the 5 km buffer.",
        "Continuous NPP-quality adjustment and pooled quality-support restrictions",
        "percent",
        n_qc,
        f"{n_strict_villages:,} villages; 2001–2021",
        f"P50={qc_p50:.2f}%; P75={qc_p75:.2f}%; lower values indicate less gap filling",
        "Final quality variable; thresholds frozen on the pooled primary sample",
    )

    add(
        "C. Stage-2 household variables",
        "Real 2021 Annual Total Consumption per Capita",
        "Annual food + recall non-food + housing services + applicable education expenditure, CPI-deflated to 2021 and divided by household size.",
        "Primary household outcome; natural log used in regression",
        "2021 riels person⁻¹ year⁻¹",
        total_outcome,
        "2007–2021 household frame",
        f"{n_main - total_outcome:,} missing ({100 * (n_main-total_outcome)/n_main:.1f}%)",
        "Primary stage-2 outcome; positive values only",
    )
    add(
        "C. Stage-2 household variables",
        "Real 2021 Annual Food Consumption per Capita",
        "Seven-day food consumption annualised by ×52, deflated with the applicable food CPI, and divided by household size.",
        "Secondary household outcome; natural log used in regression",
        "2021 riels person⁻¹ year⁻¹",
        food_outcome,
        "2007–2021 household frame",
        "No missing values in the eligible frame",
        "Secondary stage-2 outcome; positive values only",
    )
    add(
        "C. Stage-2 household variables",
        "Prior-Year Strict-Cropland NPP",
        "Strict-cropland mean NPP within the household village's 5 km buffer in interview year minus one.",
        "Primary stage-2 ecological exposure; coefficient reported per 0.1 kg C m⁻² year⁻¹",
        "kg C m⁻² year⁻¹",
        int(prior_npp.sum()),
        f"{main.loc[prior_npp, 'National Village Point ID'].nunique():,} linked villages",
        f"Observed for {pct(int(prior_npp.sum()), n_main)} of the 2007–2021 frame",
        "Exact interview-year lag; no interpolation",
    )
    add(
        "C. Stage-2 household variables",
        "Interview Calendar Year",
        "Actual household interview year; the 2019 survey wave contains interviews in both 2019 and 2020.",
        "Defines survey-period effects and the ecological lag",
        "calendar year",
        int(main["Interview Calendar Year"].notna().sum()),
        "9 survey waves; interview years 2007–2021",
        "100.0% observed in the main frame",
        "Recovered from raw timing fields where required",
    )
    add(
        "C. Stage-2 household variables",
        "Prior NPP Calendar Year",
        "Interview Calendar Year − 1.",
        "Selects the ecological exposure preceding each interview",
        "calendar year",
        int(main["Prior NPP Calendar Year"].notna().sum()),
        "Household-specific linkage year",
        "100.0% constructed in the main frame",
        "Deterministic timing rule",
    )
    add(
        "C. Stage-2 household variables",
        "Household Composition Vector",
        "Household size, female share, mean age, child share, older-person share, and dependency ratio.",
        "Pre-specified stage-2 demographic controls measured at interview",
        "count; share; years; ratio",
        dependency,
        "2007–2021 household frame",
        f"Dependency ratio is the limiting field: {n_main-dependency:,} missing ({100*(n_main-dependency)/n_main:.1f}%)",
        "Entered as controls, not fixed effects",
    )
    add(
        "C. Stage-2 household variables",
        "Socioeconomic Control Vector",
        "Urban–rural residence, agricultural participation, and whether the household head ever attended school.",
        "Expanded-control sensitivity model at interview",
        "binary indicators",
        school_observed,
        f"{school_yes:,} heads coded as ever attended school",
        f"Head-schooling field has {n_main-school_observed:,} missing ({100*(n_main-school_observed)/n_main:.2f}%)",
        "Sensitivity only; not required by the primary model",
    )
    add(
        "C. Stage-2 household variables",
        "Household Survey Weight",
        "Harmonised released household sampling weight.",
        "Applied to all national and regional stage-2 regressions",
        "survey weight",
        weight_observed,
        "All households in the 2007–2021 frame",
        "100.0% observed; estimation requires a positive weight",
        "Final survey design weight",
    )

    add(
        "D. Linkage and spatial design",
        "CSES Public-Point Linkage",
        "Deterministic village-code match from CSES geography to a unique national public village point.",
        "Attaches 5 km satellite buffers without publishing confidential coordinates",
        "binary match status",
        int(point_matched.sum()),
        f"{main.loc[point_matched, 'National Village Point ID'].nunique():,} national village points",
        f"{pct(int(point_matched.sum()), n_main)} of households matched; unmatched rows retained in linkage audits",
        "Required for stage-2 satellite linkage",
    )
    add(
        "D. Linkage and spatial design",
        "Consumption Instrument Regime",
        "Four questionnaire-construction regimes: 2007; 2009–2013; 2014–2017; and 2019–2021.",
        "Survey-wave harmonisation and regime sensitivity",
        "category",
        int(main["Total Consumption Instrument Regime"].notna().sum()),
        "9 survey waves",
        "100.0% classified in the main frame",
        "Instrument-regime effects used in sensitivity analysis",
    )
    add(
        "D. Linkage and spatial design",
        "SKATER Region ID",
        "Frozen six-region minimum-spanning-tree partition based only on standardised, outcome-blind spatial and baseline features.",
        "Discrete spatial heterogeneity analysis; fixed before outcome comparison",
        "category (1–6)",
        region_villages,
        "2,966 national public village points; six contiguous regions",
        "100.0% of the national public-point frame assigned",
        "Final outcome-blind regionalisation",
    )

    assert len(rows) == 25
    return rows


COMPACT_TEXT = {
    "Harmonised CSES household release": (
        "10 released survey waves",
        "Source universe",
        "2004–2021",
        "10 waves",
        "2004 not spatially linked",
    ),
    "Main 2007–2021 household analysis frame": (
        "Waves with public village-point support",
        "Stage-2 source",
        "2007–2021",
        "9 waves",
        "Final frame",
    ),
    "Stage-1 primary complete sample": (
        "NPP + four shocks + rainfall observed",
        "Stage-1 FE",
        "2001–2021",
        "2,899 villages",
        "Primary sample",
    ),
    "Stage-2 total-consumption complete sample": (
        "Match + lagged NPP + outcome + controls + weight >0",
        "Stage-2 total",
        "2007–2021",
        "2,848 villages; 9 waves",
        "Primary sample",
    ),
    "Stage-2 food-consumption complete sample": (
        "Primary rules, replacing total with food consumption",
        "Stage-2 food",
        "2007–2021",
        "2,848 villages; 9 waves",
        "Secondary sample",
    ),
    "Annual Strict-Cropland Mean NPP": (
        "MODIS NPP; strict cropland; 5 km mean",
        "Stage-1 Y; Stage-2 X",
        "t / t−1",
        "2,899 villages",
        "Primary; 94.9%",
    ),
    "Annual Inclusive-Agriculture Mean NPP": (
        "Strict + mosaic agriculture; 5 km mean",
        "Alternative NPP",
        "t / t−1",
        "2,900 villages",
        "Sensitivity; 95.1%",
    ),
    "Annual Heat Days at or Above 35 °C": (
        "Count days Tmax ≥35 °C; 5 km mean",
        "Heat exposure",
        "year t",
        "2,966 villages",
        "Primary; complete",
    ),
    "Annual Heat Degree-Days Above 35 °C": (
        "Σ max(Tmax −35, 0); 5 km mean",
        "Heat intensity",
        "year t",
        "2,966 villages",
        "Alternative; complete",
    ),
    "Annual Maximum Consecutive Five-Day Precipitation (Rx5day)": (
        "Maximum rolling 5-day rainfall; 5 km mean",
        "Heavy-rain exposure",
        "year t",
        "2,966 villages",
        "Primary; complete",
    ),
    "Annual Maximum Consecutive Dry Days Below 1 mm": (
        "Longest run with rain <1 mm; 5 km mean",
        "Dry-spell exposure",
        "year t",
        "2,966 villages",
        "Primary; complete",
    ),
    "Annual Precipitation Total": (
        "Annual rainfall sum; 5 km mean",
        "Water control",
        "year t",
        "2,966 villages",
        "Control; complete",
    ),
    "NPP Pixel Support": (
        "Valid / candidate 500 m pixels; QA >100 → 0",
        "Quality gate",
        "year t",
        "2,966 villages",
        "95.0% observed",
    ),
    "Mean Annual Strict-Cropland NPP Filled-Days Percentage": (
        "Mean gap-filled FPAR/LAI days across strict-cropland pixels",
        "Quality adjustment",
        "year t",
        "2,899 villages",
        "P50 62.05%; P75 67.86%",
    ),
    "Real 2021 Annual Total Consumption per Capita": (
        "Food + nonfood + housing + education†; CPI / household size",
        "Stage-2 Y (log)",
        "interview year",
        "9 waves",
        "Primary; 0.6% missing",
    ),
    "Real 2021 Annual Food Consumption per Capita": (
        "7-day food ×52; food CPI / household size",
        "Stage-2 Y (log)",
        "interview year",
        "9 waves",
        "Secondary; complete",
    ),
    "Prior-Year Strict-Cropland NPP": (
        "Strict-cropland NPP; 5 km; one-year lag",
        "Stage-2 X",
        "t−1",
        "2,848 villages",
        "Primary; 70.8% linked",
    ),
    "Interview Calendar Year": (
        "Actual interview year",
        "Timing / period FE",
        "year t",
        "9 waves",
        "Final; complete",
    ),
    "Prior NPP Calendar Year": (
        "Interview year − 1",
        "Satellite join key",
        "year t−1",
        "9 waves",
        "Final; complete",
    ),
    "Household Composition Vector": (
        "Size; female share; age; child; older; dependency",
        "Primary controls",
        "interview",
        "9 waves",
        "2.4% missing††",
    ),
    "Socioeconomic Control Vector": (
        "Urban; agriculture; head schooling",
        "Expanded controls",
        "interview",
        "9 waves",
        "Sensitivity; <0.1% missing",
    ),
    "Household Survey Weight": (
        "Released household sampling weight",
        "Survey weighting",
        "interview",
        "9 waves",
        "Final; complete",
    ),
    "CSES Public-Point Linkage": (
        "Village code → unique public point",
        "Satellite linkage",
        "time-invariant",
        "2,966 points",
        "73.8% households matched",
    ),
    "Consumption Instrument Regime": (
        "2007 | 2009–13 | 2014–17 | 2019–21",
        "Instrument check",
        "survey wave",
        "9 waves",
        "Final; complete",
    ),
    "SKATER Region ID": (
        "Outcome-blind SKATER; K=6",
        "Regional heterogeneity",
        "frozen",
        "2,966 villages",
        "Final; contiguous",
    ),
}


def compact_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    compact: list[dict[str, object]] = []
    for row in rows:
        item = str(row["Variable or analytical sample"])
        construction, role, timing, support, status = COMPACT_TEXT[item]
        compact.append(
            {
                "Section": row["Section"],
                "Item": item,
                "Construction / inclusion rule": construction,
                "Role": role,
                "Timing": timing,
                "Unit": row["Unit"],
                "N": row["Non-missing observations"],
                "Places / waves": support,
                "Status": status,
            }
        )
    return compact


def write_workbook(rows: list[dict[str, object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Samples and variables"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    section_blue = "5B9BD5"
    pale_green = "E2F0D9"
    pale_blue = "DDEBF7"
    light_gray = "F5F7F8"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    medium = Side(style="medium", color="7F9DB9")
    vertical = Side(style="thin", color="D7E0E6")

    ws.merge_cells("A1:H1")
    ws["A1"] = "Table 1. Analytical Samples and Variable Definitions"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    for col, label in enumerate(DISPLAY_COLUMNS, 1):
        cell = ws.cell(3, col, label)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=9.2, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=medium,
            bottom=medium,
            left=vertical if col in {6, 7, 8} else None,
        )
    ws.row_dimensions[3].height = 42

    current_section = None
    row_number = 4
    data_index = 0
    for record in rows:
        section = str(record["Section"])
        if section != current_section:
            ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=8)
            cell = ws.cell(row_number, 1, section)
            cell.fill = PatternFill("solid", fgColor=section_blue)
            cell.font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[row_number].height = 19
            current_section = section
            row_number += 1

        status = str(record["Status"])
        is_primary = "Primary" in status or "Final" in status
        is_sensitivity = "Sensitivity" in status or "Alternative" in status
        for col, label in enumerate(DISPLAY_COLUMNS, 1):
            value = record[label]
            cell = ws.cell(row_number, col, value)
            cell.font = Font(name="Times New Roman", size=9.2, bold=(col == 1))
            cell.alignment = Alignment(
                horizontal="right" if col == 6 else "left",
                vertical="center",
                wrap_text=True,
                indent=1 if col in {7, 8} else 0,
            )
            cell.border = Border(
                bottom=thin,
                left=vertical if col in {6, 7, 8} else None,
            )
            if data_index % 2:
                cell.fill = PatternFill("solid", fgColor=light_gray)
            if col == 8 and is_primary:
                cell.fill = PatternFill("solid", fgColor=pale_green)
                cell.font = Font(name="Times New Roman", size=9.2, bold=True)
            elif col == 8 and is_sensitivity:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
                cell.font = Font(name="Times New Roman", size=9.2, bold=True)
        ws.cell(row_number, 6).number_format = "#,##0"
        ws.row_dimensions[row_number].height = 31
        row_number += 1
        data_index += 1

    notes = [
        "Notes: N is unweighted. Stage 1 uses same-year climate and NPP; Stage 2 uses prior-year 5 km NPP. Variables remain in natural units; no standardisation, winsorisation, clipping, or imputation.",
        "2004 is excluded only because no public village-point link is available. † Education enters total consumption where required by the survey instrument. †† Dependency ratio is the limiting household-control field.",
        "Green status cells mark primary/final quantities; blue cells mark alternatives or sensitivities. Household characteristics are controls, not household fixed effects.",
    ]
    row_number += 1
    for note in notes:
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=8)
        cell = ws.cell(row_number, 1, note)
        cell.font = Font(name="Times New Roman", size=8.1, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row_number].height = 22
        row_number += 1

    widths = [34, 43, 25, 17, 23, 16, 25, 27]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A4"
    ws.sheet_view.zoomScale = 55
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.14
    ws.page_margins.top = ws.page_margins.bottom = 0.18
    ws.print_area = f"A1:H{row_number - 1}"
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    wb.properties.title = "Analytical Samples and Variable Definitions"
    wb.properties.subject = "Two-stage climate, cropland NPP, and household-consumption analysis"
    wb.properties.creator = "Mike Li"

    # Preserve cell content while removing presentation-only merges so the
    # worksheet can be exported as one deterministic rectangular Word table.
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    for note_row in range(row_number - len(notes), row_number):
        ws.cell(note_row, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        ws.row_dimensions[note_row].height = 17
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate(rows: list[dict[str, object]]) -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Samples and variables"]
    ws = wb["Samples and variables"]
    assert ws.max_column == 8
    assert not ws.merged_cells.ranges
    labels = [record["Item"] for record in rows]
    for label in labels:
        assert any(cell.value == label for row in ws.iter_rows() for cell in row)
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    detailed_rows = build_rows()
    rows = compact_rows(detailed_rows)
    evidence = pd.DataFrame(detailed_rows)
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    evidence.to_csv(EVIDENCE_DIR / "analytical_samples_and_variable_definitions.csv", index=False)
    pd.DataFrame(rows).to_csv(EVIDENCE_DIR / "analytical_samples_and_variable_definitions_display.csv", index=False)
    write_workbook(rows)
    validate(rows)
    print(f"Wrote {OUTPUT}")
    print(f"Data rows: {len(rows)}; sheets: 1")


if __name__ == "__main__":
    main()
