#!/usr/bin/env python3
"""Create a one-sheet table of activated high-frequency response dynamics."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/exp/high-frequency-boundary-dynamics/distributed_lag_summary_estimates.csv"
OUTPUT = ROOT / "data/exp/legacy-results/tables/Table_high_frequency_rainfall_response_dynamics.xlsx"


def interpretation(row: pd.Series) -> str:
    if row["ci_low"] >= -0.20 and row["ci_high"] <= 0.20 and "Cumulative" not in row["estimand"]:
        return "Entire 95% CI lies within the ±0.20 SD equivalence region"
    if row["ci_low"] <= 0 <= row["ci_high"]:
        return "Confidence interval includes no historical-side response difference"
    return "Southwest-side response is larger" if row["estimate"] > 0 else "Southwest-side response is smaller"


def main() -> None:
    data = pd.read_csv(SOURCE)
    data["Outcome"] = data["outcome_family"]
    data["Shock intensity"] = data["shock_family"]
    data["Estimand"] = data["estimand"]
    data["Specification"] = data["specification"]
    data["Scale"] = data["scale"]
    data["Estimate"] = data["estimate"]
    data["95% CI"] = data.apply(lambda row: f"[{row['ci_low']:.3f}, {row['ci_high']:.3f}]", axis=1)
    data["p-value"] = data["p_value"]
    data["Lead-test p-value"] = data["pretrend_p_value"]
    data["Observations"] = data["observations"].astype(int)
    data["Dates"] = data["dates"].astype(int)
    data["Villages"] = data["villages"].astype(int)
    data["Inference"] = data["inference"]
    data["Interpretation"] = data.apply(interpretation, axis=1)
    columns = [
        "Outcome", "Shock intensity", "Estimand", "Specification", "Scale",
        "Estimate", "95% CI", "p-value", "Lead-test p-value", "Observations",
        "Dates", "Villages", "Inference", "Interpretation",
    ]
    estimand_order = {
        "Immediate average response, periods 0-2": 0,
        "Cumulative response, periods 0-5": 1,
        "Late response, periods 6-8": 2,
    }
    data["_estimand_order"] = data["Estimand"].map(estimand_order)
    data = data.sort_values(["Outcome", "Shock intensity", "_estimand_order", "Specification"])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Response Dynamics"
    sheet.append(columns)
    for row in data[columns].itertuples(index=False, name=None):
        sheet.append(list(row))
    navy, light_blue, white = "1F4E78", "D9EAF7", "FFFFFF"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(2, sheet.max_row + 1):
        if row % 2 == 0:
            for cell in sheet[row]:
                cell.fill = PatternFill("solid", fgColor=light_blue)
        for cell in sheet[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D9D9D9"))
        for column in (6, 8, 9):
            sheet.cell(row, column).number_format = "0.000"
    widths = [10, 17, 38, 18, 27, 12, 20, 12, 18, 14, 10, 11, 34, 48]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 38
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:1"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows: {len(data):,}; columns: {len(columns)}; sheets: 1")


if __name__ == "__main__":
    main()
