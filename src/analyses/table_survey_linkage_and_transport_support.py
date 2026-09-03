#!/usr/bin/env python3
"""Survey Linkage and Transport Support.

Plan: Quantify matched and unmatched CSES geography, repeated-village and
wave-level support, and the national covariate-support boundary for welfare
interpretation.
Framework: AnaSOP workflow step 1 and the national extrapolation gate.  The
support score is outcome-blind and does not restrict national hazard mapping.
"""

from __future__ import annotations

from pathlib import Path

import duckdb
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.neighbors import NearestNeighbors


ROOT = Path(__file__).resolve().parents[2]
CSES = ROOT / "data/processed/cses_climate_ecology_welfare_frame_preprocessed.parquet"
PSEUDO = ROOT / (
    "data/processed/"
    "cses_commune_survey_time_climate_food_pseudopanel_preprocessed.parquet"
)
VILLAGES = ROOT / "data/processed/cambodia_public_village_points_preprocessed.parquet"
MEMBERSHIP = ROOT / "data/processed/cambodia_public_village_buffer_grid_crosswalk/radius_5_km.parquet"
CONTEXT = ROOT / "data/processed/cambodia_national_predetermined_covariates_preprocessed.parquet"
RISK = ROOT / "data/exp/analysis/climate-welfare/national-dryspell-absolute-heat-risk/village_risk_metrics_5km.parquet"
ANALYSIS_DIR = ROOT / "data/exp/analysis/climate-welfare/survey-linkage-transport-support"
OUTPUT = ROOT / "data/exp/internal_output_archive/obsolete-formal-results-20260824/tables/Table_survey_linkage_and_transport_support.xlsx"

ID = "National Village Point ID"
FOOD = "Real 2021 Food Consumption Value per Household Member Riels"
HEAT = "Last Complete Season Absolute Heat Day Count 35 C Candidate B"
WEIGHT = "Household Survey Weight"
ROAD = "Historical Road Distance km"

COLUMNS = [
    "Linkage group, survey wave, or support stratum",
    "Period / unit",
    "Households",
    "Villages / communes",
    "Weighted / national share",
    "Primary outcome support",
    "Historical road distance, median [P10, P90] km",
    "Support score, median [P10, P90]",
    "Mapping eligibility",
    "Limitation / analytical role",
]

SUPPORT_VARIABLES = [
    "Historical Mean Longest Dry Spell Days",
    "Historical Mean Post-Onset Heat Days 35 C",
    "Mean Elevation m",
    "Mean Slope Degrees",
    "Distance to Nearest Historical Road Proxy km",
    "Log Baseline Population 2000",
    "Baseline Cropland Share",
]


def q(path: Path) -> str:
    return path.as_posix().replace("'", "''")


def weighted_share(mask: pd.Series, weights: pd.Series) -> float:
    valid = weights.notna() & weights.gt(0)
    if not valid.any():
        return np.nan
    return float(weights.loc[valid & mask].sum() / weights.loc[valid].sum())


def interval_text(values: pd.Series, digits: int = 2) -> str:
    clean = values.dropna()
    if clean.empty:
        return "—"
    p10, median, p90 = clean.quantile([0.10, 0.50, 0.90])
    return f"{median:.{digits}f} [{p10:.{digits}f}, {p90:.{digits}f}]"


def outcome_text(frame: pd.DataFrame) -> str:
    if frame.empty:
        return "—"
    complete = frame[FOOD].gt(0) & frame[HEAT].notna()
    return f"Food + heat complete: {complete.mean():.1%}"


def linked_outcome_text(frame: pd.DataFrame) -> str:
    if frame.empty:
        return "—"
    complete = frame[FOOD].gt(0) & frame[HEAT].notna()
    return f"Among linked: {complete.mean():.1%}"


def public_context() -> pd.DataFrame:
    con = duckdb.connect()
    context = con.execute(
        f"""
        SELECT m."National Village Point ID",
          avg(c."Mean Elevation m") AS "Mean Elevation m",
          avg(c."Mean Slope Degrees") AS "Mean Slope Degrees",
          avg(c."Distance to Nearest Historical Road Proxy km")
            AS "Distance to Nearest Historical Road Proxy km",
          avg(c."Log Baseline Population 2000") AS "Log Baseline Population 2000",
          avg(c."Baseline Cropland Share") AS "Baseline Cropland Share"
        FROM read_parquet('{q(MEMBERSHIP)}') m
        INNER JOIN read_parquet('{q(CONTEXT)}') c USING ("National Grid Cell ID")
        GROUP BY 1
        """
    ).df()
    con.close()
    villages = pd.read_parquet(VILLAGES, columns=[ID])
    risk = pd.read_parquet(
        RISK,
        columns=[
            ID,
            "Historical Mean Longest Dry Spell Days",
            "Historical Mean Post-Onset Heat Days 35 C",
            "CSES Linked Village",
        ],
    )
    frame = villages.merge(context, on=ID, how="left", validate="one_to_one")
    frame = frame.merge(risk, on=ID, how="left", validate="one_to_one")
    assert frame[SUPPORT_VARIABLES].notna().all(axis=1).all()
    return frame


def add_support_score(frame: pd.DataFrame) -> tuple[pd.DataFrame, float, float]:
    linked = frame["CSES Linked Village"].eq(1)
    reference = frame.loc[linked, SUPPORT_VARIABLES]
    median = reference.median()
    scale = (reference.quantile(0.75) - reference.quantile(0.25)).replace(0, 1.0)
    standardized = ((frame[SUPPORT_VARIABLES] - median) / scale).to_numpy()
    linked_values = standardized[linked.to_numpy()]

    leave_one_out = NearestNeighbors(n_neighbors=2).fit(linked_values)
    linked_distance = leave_one_out.kneighbors(linked_values)[0][:, 1]
    threshold95 = float(np.quantile(linked_distance, 0.95))
    threshold99 = float(np.quantile(linked_distance, 0.99))

    nearest = NearestNeighbors(n_neighbors=1).fit(linked_values)
    distance = nearest.kneighbors(standardized)[0][:, 0]
    result = frame.copy()
    result["Nearest Linked Covariate Distance"] = distance
    result["Survey Support Score"] = np.exp(-distance)
    result["Primary 95 Percent Support"] = distance <= threshold95
    result["Sensitivity 99 Percent Support"] = distance <= threshold99
    return result, threshold95, threshold99


def survey_row(
    label: str,
    period: str,
    frame: pd.DataFrame,
    denominator: pd.DataFrame,
    eligibility: str,
    limitation: str,
) -> dict[str, object]:
    linked = frame.loc[frame["Climate Ecology Link Available"].eq(1)]
    return {
        COLUMNS[0]: label,
        COLUMNS[1]: period,
        COLUMNS[2]: len(frame),
        COLUMNS[3]: frame["Village Code"].nunique(),
        COLUMNS[4]: weighted_share(frame.index.to_series().isin(frame.index), frame[WEIGHT])
        if frame.index.equals(denominator.index)
        else float(frame[WEIGHT].sum() / denominator[WEIGHT].sum()),
        COLUMNS[5]: outcome_text(frame) if frame["Climate Ecology Link Available"].eq(1).all() else linked_outcome_text(linked),
        COLUMNS[6]: interval_text(linked[ROAD]),
        COLUMNS[7]: "1.00 [1.00, 1.00]" if len(frame) and frame["Climate Ecology Link Available"].eq(1).all() else "—",
        COLUMNS[8]: eligibility,
        COLUMNS[9]: limitation,
    }


def national_row(
    label: str,
    frame: pd.DataFrame,
    total: int,
    eligibility: str,
    limitation: str,
) -> dict[str, object]:
    return {
        COLUMNS[0]: label,
        COLUMNS[1]: "National village point, 5 km context",
        COLUMNS[2]: None,
        COLUMNS[3]: len(frame),
        COLUMNS[4]: len(frame) / total,
        COLUMNS[5]: "Not a household sample",
        COLUMNS[6]: interval_text(frame["Distance to Nearest Historical Road Proxy km"]),
        COLUMNS[7]: interval_text(frame["Survey Support Score"]),
        COLUMNS[8]: eligibility,
        COLUMNS[9]: limitation,
    }


def build_table() -> tuple[pd.DataFrame, list[int], pd.DataFrame, float, float]:
    cses = pd.read_parquet(
        CSES,
        columns=[
            "Survey Year", "Household ID", "Village Code", WEIGHT,
            "Climate Ecology Link Available", ID, FOOD, HEAT, ROAD,
        ],
    ).reset_index(drop=True)
    linked_mask = cses["Climate Ecology Link Available"].eq(1)
    linked = cses.loc[linked_mask]
    unlinked = cses.loc[~linked_mask]
    wave_counts = linked.groupby("Village Code", observed=True)["Survey Year"].nunique()
    repeated_ids = wave_counts.loc[lambda x: x > 1].index
    repeated = linked.loc[linked["Village Code"].isin(repeated_ids)]
    single = linked.loc[~linked["Village Code"].isin(repeated_ids)]
    pseudo = pd.read_parquet(PSEUDO)
    repeated_cells = pseudo.loc[pseudo["Repeated Commune Indicator"]].copy()
    primary_cells = repeated_cells.loc[
        repeated_cells["Primary Minimum Five Households"]
    ].copy()
    robust_cells = repeated_cells.loc[
        repeated_cells["Robustness Minimum Ten Households"]
    ].copy()
    total_pseudo_weight = float(pseudo["Survey Weight Sum"].sum())

    def pseudo_row(
        label: str,
        frame: pd.DataFrame,
        eligibility: str,
        limitation: str,
    ) -> dict[str, object]:
        return {
            COLUMNS[0]: label,
            COLUMNS[1]: "2007-2021 / commune-survey-time cell",
            COLUMNS[2]: int(frame["Household Count"].sum()),
            COLUMNS[3]: int(frame["Commune Code"].nunique()),
            COLUMNS[4]: float(frame["Survey Weight Sum"].sum() / total_pseudo_weight),
            COLUMNS[5]: "Survey-weighted food + approved climate complete: 100.0%",
            COLUMNS[6]: "—",
            COLUMNS[7]: "—",
            COLUMNS[8]: eligibility,
            COLUMNS[9]: limitation,
        }

    rows: list[dict[str, object]] = []
    rows.extend(
        [
            survey_row(
                "Harmonized CSES frame", "2007–2021 / household", cses, cses,
                "Mixed; link required for household interpretation",
                f"Coverage denominator; {100 * len(unlinked) / len(cses):.1f}% of households lack a public-point link",
            ),
            survey_row(
                "Public-point linked households", "2007–2021 / household", linked, cses,
                "Observed household-support locations",
                "Primary Model B linkage frame",
            ),
            survey_row(
                "Unlinked or ambiguous households", "2007–2021 / household", unlinked, cses,
                "Not eligible for climate-linked household analysis",
                "No unambiguous national public-village point; exposures remain missing",
            ),
            survey_row(
                "Repeated-village linked sample", "2007–2021 / household", repeated, cses,
                "Eligible for village-effect confirmation",
                "Village observed in more than one survey wave",
            ),
            survey_row(
                "Single-wave linked sample", "2007–2021 / household", single, cses,
                "Eligible for repeated-cross-section model only",
                "Cannot contribute to village-effect confirmation",
            ),
            pseudo_row(
                "All linked commune survey-time cells", pseudo,
                "Descriptive commune-cell frame",
                "Includes communes observed in only one survey wave",
            ),
            pseudo_row(
                "Repeated-commune survey-time cells", repeated_cells,
                "Eligible for within-commune estimation before cell-size gate",
                "Commune observed in more than one survey wave",
            ),
            pseudo_row(
                "Primary pseudo-panel cells (at least 5 households)", primary_cells,
                "Primary within-commune Model B sample",
                "Repeated commune and at least 5 households per survey-time cell",
            ),
            pseudo_row(
                "Pseudo-panel sensitivity cells (at least 10 households)", robust_cells,
                "Cell-size robustness sample",
                "Repeated commune and at least 10 households per survey-time cell",
            ),
        ]
    )
    section_breaks = [0, len(rows)]

    for wave, frame in cses.groupby("Survey Year", observed=True, sort=True):
        linked_wave = frame.loc[frame["Climate Ecology Link Available"].eq(1)]
        rows.append(
            {
                COLUMNS[0]: f"CSES {int(wave)}",
                COLUMNS[1]: f"{int(wave)} wave / household",
                COLUMNS[2]: len(frame),
                COLUMNS[3]: frame["Village Code"].nunique(),
                COLUMNS[4]: float(linked_wave[WEIGHT].sum() / frame[WEIGHT].sum()),
                COLUMNS[5]: linked_outcome_text(linked_wave),
                COLUMNS[6]: interval_text(linked_wave[ROAD]),
                COLUMNS[7]: "1.00 [1.00, 1.00]",
                COLUMNS[8]: "Eligible after successful public-point link",
                COLUMNS[9]: f"{len(frame) - len(linked_wave):,} households unlinked or ambiguous",
            }
        )

    national, threshold95, threshold99 = add_support_score(public_context())
    section_breaks.append(len(rows))
    primary = national["Primary 95 Percent Support"]
    sensitivity = national["Sensitivity 99 Percent Support"]
    linked_national = national["CSES Linked Village"].eq(1)
    rows.extend(
        [
            national_row(
                "National village frame", national, len(national),
                "Hazard mapping unrestricted; welfare interpretation conditional",
                "All villages retain continuous hazard exposure regardless of survey support",
            ),
            national_row(
                "CSES-linked national villages", national.loc[linked_national], len(national),
                "Observed household-support locations",
                "Reference set used to construct the outcome-blind support score",
            ),
            national_row(
                "Primary transport-support eligible", national.loc[primary], len(national),
                "Eligible for bounded household interpretation",
                "Nearest-linked distance no greater than linked leave-one-out P95",
            ),
            national_row(
                "Outside primary transport support", national.loc[~primary], len(national),
                "Hazard map only; no household-loss label",
                "Beyond linked leave-one-out P95 in seven-variable covariate space",
            ),
            national_row(
                "P99 sensitivity: eligible", national.loc[sensitivity], len(national),
                "Sensitivity boundary only",
                "Looser nearest-linked boundary; not used to define the primary mask",
            ),
            national_row(
                "P99 sensitivity: outside support", national.loc[~sensitivity], len(national),
                "Hazard map only under sensitivity boundary",
                "Extreme covariate-distance tail beyond linked leave-one-out P99",
            ),
        ]
    )
    table = pd.DataFrame(rows, columns=COLUMNS)
    return table, section_breaks, national, threshold95, threshold99


def write_workbook(table: pd.DataFrame, section_breaks: list[int], threshold95: float, threshold99: float) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Linkage and support"
    ws.sheet_view.showGridLines = False

    navy, blue, light, green, white = "1F4E78", "5B9BD5", "F3F3F3", "E2F0D9", "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    medium = Side(style="medium", color="7F9DB9")
    vertical = Side(style="thin", color="D7E0E6")

    ws.merge_cells("A1:J1")
    ws["A1"] = "Survey Linkage and Transport Support"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    for column, label in enumerate(COLUMNS, 1):
        cell = ws.cell(3, column, label)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=9.2, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=medium, left=vertical if column in {5, 7, 9} else None)
    ws.row_dimensions[3].height = 45

    labels = {
        section_breaks[0]: "A. Overall linkage, sample structure, and pseudo-panel support",
        section_breaks[1]: "B. Survey-wave linkage support",
        section_breaks[2]: "C. National covariate-transport support",
    }
    cursor = 4
    for index, values in table.iterrows():
        if index in labels:
            ws.merge_cells(start_row=cursor, start_column=1, end_row=cursor, end_column=10)
            cell = ws.cell(cursor, 1, labels[index])
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(name="Times New Roman", size=10, bold=True, color=white)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = Border(top=medium, bottom=thin)
            ws.row_dimensions[cursor].height = 20
            cursor += 1
        for column, value in enumerate(values, 1):
            if pd.isna(value):
                value = None
            cell = ws.cell(cursor, column, value)
            cell.font = Font(name="Times New Roman", size=8.8)
            cell.alignment = Alignment(
                horizontal="center" if column in {3, 4, 5} else "left",
                vertical="center", wrap_text=True,
            )
            cell.border = Border(bottom=thin, left=vertical if column in {5, 7, 9} else None)
            if index % 2:
                cell.fill = PatternFill("solid", fgColor=light)
        ws.cell(cursor, 3).number_format = "#,##0"
        ws.cell(cursor, 4).number_format = "#,##0"
        ws.cell(cursor, 5).number_format = "0.0%"
        if values.iloc[0] in {
            "Public-point linked households",
            "Primary pseudo-panel cells (at least 5 households)",
            "Primary transport-support eligible",
        }:
            for cell in ws[cursor]:
                cell.fill = PatternFill("solid", fgColor=green)
                cell.font = Font(name="Times New Roman", size=8.8, bold=True)
        ws.row_dimensions[cursor].height = 35
        cursor += 1

    notes = [
        "Notes: Survey and commune-cell shares use released household weights; national shares are unweighted shares of 13,042 public village points. Commune-cell household counts sum represented linked households and do not imply longitudinal household follow-up.",
        "The support score is exp(−d), where d is nearest-neighbour distance to a linked CSES village after robust IQR scaling of dry spell, heat, elevation, slope, historical-road distance, baseline log population, and cropland share.",
        f"Primary and sensitivity distance limits are the linked-village leave-one-out P95 ({threshold95:.3f}) and P99 ({threshold99:.3f}). The mask limits household interpretation only; it never removes villages from national hazard maps.",
    ]
    notes_start = cursor + 1
    for offset, note in enumerate(notes):
        row = notes_start + offset
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell = ws.cell(row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.2, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 19

    widths = [31, 22, 13, 13, 18, 25, 28, 25, 31, 40]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:J{cursor - 1}"
    ws.sheet_view.zoomScale = 58
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.16
    ws.page_margins.top = ws.page_margins.bottom = 0.22
    ws.print_area = f"A1:J{notes_start + len(notes) - 1}"
    wb.properties.title = "Survey Linkage and Transport Support"
    wb.properties.subject = "Outcome-blind survey linkage and national covariate-support audit"
    wb.properties.creator = "Mike Li"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate(table: pd.DataFrame, national: pd.DataFrame) -> None:
    assert table.shape == (24, 10), table.shape
    assert len(national) == 13_042
    assert int(table.loc[table[COLUMNS[0]].eq("Harmonized CSES frame"), COLUMNS[2]].iloc[0]) == 62_920
    assert int(table.loc[table[COLUMNS[0]].eq("Public-point linked households"), COLUMNS[2]].iloc[0]) == 46_445
    assert int(table.loc[table[COLUMNS[0]].eq("Repeated-village linked sample"), COLUMNS[2]].iloc[0]) == 21_724
    assert int(table.loc[table[COLUMNS[0]].eq("Primary pseudo-panel cells (at least 5 households)"), COLUMNS[2]].iloc[0]) == 42_870
    assert int(table.loc[table[COLUMNS[0]].eq("Primary pseudo-panel cells (at least 5 households)"), COLUMNS[3]].iloc[0]) == 994
    assert national["Primary 95 Percent Support"].mean() > 0.90
    assert national["Sensitivity 99 Percent Support"].mean() >= national["Primary 95 Percent Support"].mean()
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Linkage and support"]
    assert wb["Linkage and support"].max_column == 10
    for row in wb["Linkage and support"].iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    table, section_breaks, national, threshold95, threshold99 = build_table()
    ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)
    table.to_csv(ANALYSIS_DIR / "table_rows.csv", index=False)
    national.to_parquet(ANALYSIS_DIR / "national_village_support.parquet", index=False, compression="zstd")
    write_workbook(table, section_breaks, threshold95, threshold99)
    validate(table, national)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Rows={len(table)}; primary support={national['Primary 95 Percent Support'].mean():.1%}; P99 support={national['Sensitivity 99 Percent Support'].mean():.1%}")


if __name__ == "__main__":
    main()
