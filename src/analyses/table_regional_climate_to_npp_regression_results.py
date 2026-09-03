#!/usr/bin/env python3
"""Regional Climate-to-NPP Regression Results.

Plan: Compare national and frozen six-region Stage-1 climate slopes and report
joint regional slope-equality tests in one compact regression table.
Framework: AnaSOP Sections 5-7 outcome-blind SKATER regionalisation, a common
interaction model, village and year fixed effects, annual rainfall control,
village-clustered inference, and pre-specified joint Wald tests.
"""

from __future__ import annotations

from pathlib import Path

from linearmodels.iv import AbsorbingLS
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
PANEL = ROOT / "data/processed/cses_public_village_cropland_npp_absolute_climate_panel_preprocessed.parquet"
REGIONS = ROOT / "data/processed/outcome_blind_spatial_regions_preprocessed.parquet"
OUTPUT = ROOT / "data/results/tables/Table_regional_climate_to_npp_regression_results.xlsx"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/regional-climate-to-npp-regression-results"

ID = "National Village Point ID"
YEAR = "Year"
OUTCOME = "Annual Strict-Cropland Mean NPP kg C per m2"
REGION = "SKATER Region ID"
HEAT = "Village Buffer Mean Annual Heat Days at or Above 35 C"
RX5DAY = "Village Buffer Mean Annual Maximum Consecutive Five-Day Precipitation Rx5day mm"
DRY = "Village Buffer Mean Annual Maximum Consecutive Dry Days Below 1 mm"
RAIN = "Village Buffer Mean Annual Precipitation Total mm"

HEAT_10 = "Heat days ≥35 °C / 10"
RX5DAY_10 = "Rx5day / 10 mm"
DRY_10 = "Maximum dry spell / 10 days"
RAIN_100 = "Annual precipitation / 100 mm"
EXPOSURES = [HEAT_10, RX5DAY_10, DRY_10]
DISPLAY = {
    HEAT_10: "Heat days ≥35 °C (per 10 days)",
    RX5DAY_10: "Rx5day (per 10 mm)",
    DRY_10: "Maximum dry spell (per 10 days)",
}


def prepare_sample() -> pd.DataFrame:
    panel = pd.read_parquet(PANEL, columns=[ID, YEAR, OUTCOME, HEAT, RX5DAY, DRY, RAIN])
    regions = pd.read_parquet(REGIONS, columns=[ID, REGION])
    sample = panel.merge(regions, on=ID, validate="many_to_one")
    sample = sample.loc[sample[YEAR].between(2001, 2021)].copy()
    for column in [OUTCOME, HEAT, RX5DAY, DRY, RAIN, REGION]:
        sample[column] = pd.to_numeric(sample[column], errors="coerce")
    sample = sample.dropna(subset=[ID, YEAR, OUTCOME, HEAT, RX5DAY, DRY, RAIN, REGION]).reset_index(drop=True)
    sample[REGION] = sample[REGION].astype(int)
    sample[HEAT_10] = sample[HEAT] / 10.0
    sample[RX5DAY_10] = sample[RX5DAY] / 10.0
    sample[DRY_10] = sample[DRY] / 10.0
    sample[RAIN_100] = sample[RAIN] / 100.0
    return sample


def fit_absorbed(sample: pd.DataFrame, exog: pd.DataFrame) -> object:
    absorb = pd.DataFrame(
        {
            "Village fixed effect": sample[ID].astype("category"),
            "Calendar-year fixed effect": sample[YEAR].astype("category"),
        },
        index=sample.index,
    )
    clusters = pd.DataFrame(
        {"Village cluster": pd.Categorical(sample[ID]).codes},
        index=sample.index,
    )
    return AbsorbingLS(
        sample[OUTCOME].astype(float),
        exog.astype(float),
        absorb=absorb,
        drop_absorbed=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)


def estimate(sample: pd.DataFrame) -> tuple[object, object, dict[str, list[str]], pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    national_exog = sample[[HEAT_10, RX5DAY_10, DRY_10, RAIN_100]].copy()
    national = fit_absorbed(sample, national_exog)

    regional_columns: dict[str, np.ndarray] = {}
    terms_by_exposure: dict[str, list[str]] = {}
    for exposure in EXPOSURES:
        terms: list[str] = []
        for region_id in range(1, 7):
            term = f"{exposure} - R{region_id}"
            regional_columns[term] = sample[exposure].to_numpy(dtype=float) * sample[REGION].eq(region_id)
            terms.append(term)
        terms_by_exposure[exposure] = terms
    regional_columns[RAIN_100] = sample[RAIN_100].to_numpy(dtype=float)
    regional_exog = pd.DataFrame(regional_columns, index=sample.index)
    regional = fit_absorbed(sample, regional_exog)

    coefficient_rows: list[dict[str, object]] = []
    national_ci = national.conf_int(level=0.95)
    for exposure in EXPOSURES:
        coefficient_rows.append(
            {
                "Exposure": exposure,
                "Geography": "National",
                "Estimate": float(national.params[exposure]),
                "Clustered Standard Error": float(national.std_errors[exposure]),
                "95 Percent CI Lower": float(national_ci.loc[exposure, "lower"]),
                "95 Percent CI Upper": float(national_ci.loc[exposure, "upper"]),
                "Probability Value": float(national.pvalues[exposure]),
            }
        )
    regional_ci = regional.conf_int(level=0.95)
    for exposure, terms in terms_by_exposure.items():
        for region_id, term in enumerate(terms, start=1):
            coefficient_rows.append(
                {
                    "Exposure": exposure,
                    "Geography": f"R{region_id}",
                    "Estimate": float(regional.params[term]),
                    "Clustered Standard Error": float(regional.std_errors[term]),
                    "95 Percent CI Lower": float(regional_ci.loc[term, "lower"]),
                    "95 Percent CI Upper": float(regional_ci.loc[term, "upper"]),
                    "Probability Value": float(regional.pvalues[term]),
                }
            )

    test_rows: list[dict[str, object]] = []
    regressors = list(regional_exog.columns)
    for exposure, terms in terms_by_exposure.items():
        restriction = np.zeros((5, len(regressors)))
        for row_index, comparison in enumerate(terms[1:]):
            restriction[row_index, regressors.index(comparison)] = 1.0
            restriction[row_index, regressors.index(terms[0])] = -1.0
        test = regional.wald_test(restriction=restriction, value=np.zeros(5))
        test_rows.append(
            {
                "Exposure": exposure,
                "Null Hypothesis": "All six regional slopes are equal",
                "Wald Statistic": float(test.stat),
                "Degrees of Freedom": int(test.df),
                "Probability Value": float(test.pval),
            }
        )

    support = (
        sample.groupby(REGION, as_index=False)
        .agg(Observations=(ID, "size"), Villages=(ID, "nunique"))
        .rename(columns={REGION: "Region ID"})
    )
    return national, regional, terms_by_exposure, pd.DataFrame(coefficient_rows), pd.DataFrame(test_rows), support


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def coefficient_text(frame: pd.DataFrame, exposure: str, geography: str) -> str:
    row = frame.loc[(frame["Exposure"] == exposure) & (frame["Geography"] == geography)].iloc[0]
    return f"{row['Estimate']:.5f}{stars(float(row['Probability Value']))}\n({row['Clustered Standard Error']:.5f})"


def p_text(value: float) -> str:
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def write_workbook(
    sample: pd.DataFrame,
    coefficients: pd.DataFrame,
    tests: pd.DataFrame,
    support: pd.DataFrame,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Regional Stage 1"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    mid_blue = "5B9BD5"
    light_gray = "F4F6F7"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    medium = Side(style="medium", color="7F9DB9")
    vertical = Side(style="thin", color="D7E0E6")
    region_colors = ["173F5F", "2F80A2", "3A9D8F", "D4A72C", "D97757", "9B4A63"]

    ws.merge_cells("A1:I1")
    ws["A1"] = "Table 3. Regional Climate-to-NPP Regression Results"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A3"] = "Dependent variable"
    ws.merge_cells("B3:I3")
    ws["B3"] = "Annual strict-cropland mean NPP (kg C m⁻² year⁻¹)"
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=medium, bottom=medium)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 25

    headers = ["Climate slope", "National", "R1", "R2", "R3", "R4", "R5", "R6", "Equality p"]
    for column, label in enumerate(headers, start=1):
        cell = ws.cell(4, column, label)
        fill = mid_blue
        if 3 <= column <= 8:
            fill = region_colors[column - 3]
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Times New Roman", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=vertical if column > 1 else None)
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[4].height = 28

    row_number = 5
    for index, exposure in enumerate(EXPOSURES):
        label_cell = ws.cell(row_number, 1, DISPLAY[exposure])
        label_cell.font = Font(name="Times New Roman", size=9.7, bold=True)
        label_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        label_cell.border = Border(bottom=thin)
        if index % 2:
            label_cell.fill = PatternFill("solid", fgColor=light_gray)
        for column, geography in enumerate(["National", "R1", "R2", "R3", "R4", "R5", "R6"], start=2):
            cell = ws.cell(row_number, column, coefficient_text(coefficients, exposure, geography))
            cell.font = Font(name="Times New Roman", size=9.5, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=vertical, bottom=thin)
            if index % 2:
                cell.fill = PatternFill("solid", fgColor=light_gray)
        test_p = float(tests.loc[tests["Exposure"] == exposure, "Probability Value"].iloc[0])
        joint = ws.cell(row_number, 9, p_text(test_p))
        joint.font = Font(name="Times New Roman", size=9.7, bold=True)
        joint.alignment = Alignment(horizontal="center", vertical="center")
        joint.border = Border(left=vertical, bottom=thin)
        if index % 2:
            joint.fill = PatternFill("solid", fgColor=light_gray)
        ws.row_dimensions[row_number].height = 38
        row_number += 1

    ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=9)
    ws.cell(row_number, 1, "Model structure and regional support")
    ws.cell(row_number, 1).fill = PatternFill("solid", fgColor=mid_blue)
    ws.cell(row_number, 1).font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
    ws.cell(row_number, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row_number].height = 20
    row_number += 1

    regional_obs = support.set_index("Region ID")["Observations"].to_dict()
    regional_villages = support.set_index("Region ID")["Villages"].to_dict()
    structure_rows = [
        ("Single common interaction model", ["Reference"] + ["Yes"] * 6 + [None], None),
        ("Annual precipitation control", ["Yes"] * 7 + [None], None),
        ("Village fixed effects", ["Yes"] * 7 + [None], None),
        ("Calendar-year fixed effects", ["Yes"] * 7 + [None], None),
        ("Standard errors clustered by village", ["Yes"] * 7 + [None], None),
        (
            "Contributing observations",
            [len(sample)] + [int(regional_obs[i]) for i in range(1, 7)] + [None],
            "#,##0",
        ),
        (
            "Contributing villages",
            [sample[ID].nunique()] + [int(regional_villages[i]) for i in range(1, 7)] + [None],
            "#,##0",
        ),
        ("Years", [sample[YEAR].nunique()] * 7 + [None], "0"),
        ("Joint-test degrees of freedom", [None] * 7 + [5], "0"),
    ]
    for index, (label, values, number_format) in enumerate(structure_rows):
        label_cell = ws.cell(row_number, 1, label)
        label_cell.font = Font(name="Times New Roman", size=9.3, bold=index < 5)
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        label_cell.border = Border(bottom=thin)
        if index % 2:
            label_cell.fill = PatternFill("solid", fgColor=light_gray)
        for column, value in enumerate(values, start=2):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="Times New Roman", size=9.3)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=vertical, bottom=thin)
            if number_format:
                cell.number_format = number_format
            if index % 2:
                cell.fill = PatternFill("solid", fgColor=light_gray)
        ws.row_dimensions[row_number].height = 24
        row_number += 1

    row_number += 1
    notes = [
        "Notes: Coefficient cells report estimates and village-clustered standard errors in parentheses. National values reproduce the full primary model; R1–R6 are slopes from one common regional-interaction model.",
        "All models use 59,135 village-year observations, village and calendar-year fixed effects, and annual precipitation as a common control. Regional support counts describe contributions to the common model, not separate regressions.",
        "Equality p tests the null that all six regional slopes are equal (5 df). *** p<0.01; ** p<0.05; * p<0.10. SKATER regions were frozen using outcome-blind features before coefficient inspection.",
    ]
    for note in notes:
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=9)
        cell = ws.cell(row_number, 1, note)
        cell.font = Font(name="Times New Roman", size=8.3, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row_number].height = 22
        row_number += 1

    widths = [39, 18, 15, 15, 15, 15, 15, 15, 17]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 75
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.20
    ws.page_margins.top = ws.page_margins.bottom = 0.25
    ws.print_area = f"A1:I{row_number - 1}"
    wb.properties.title = "Regional Climate-to-NPP Regression Results"
    wb.properties.subject = "Frozen six-region Stage-1 interaction model"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Regional Stage 1"]
    ws = wb["Regional Stage 1"]
    assert ws.max_column == 9
    assert not ws.merged_cells.ranges
    assert [ws.cell(4, column).value for column in range(1, 10)] == [
        "Climate slope", "National", "R1", "R2", "R3", "R4", "R5", "R6", "Equality p"
    ]
    for sheet_row in ws.iter_rows():
        for cell in sheet_row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    sample = prepare_sample()
    _, _, _, coefficients, tests, support = estimate(sample)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    coefficients.to_csv(EVIDENCE / "regional_climate_to_npp_regression_coefficients.csv", index=False)
    tests.to_csv(EVIDENCE / "regional_climate_to_npp_slope_equality_tests.csv", index=False)
    support.to_csv(EVIDENCE / "regional_climate_to_npp_support.csv", index=False)
    write_workbook(sample, coefficients, tests, support)
    validate()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Common sample: {len(sample):,} village-years; {sample[ID].nunique():,} villages")
    print(coefficients.to_string(index=False))
    print(tests.to_string(index=False))
    print(support.to_string(index=False))


if __name__ == "__main__":
    main()
