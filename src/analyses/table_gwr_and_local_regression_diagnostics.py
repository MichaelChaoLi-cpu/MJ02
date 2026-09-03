#!/usr/bin/env python3
"""GWR and Local-Regression Diagnostics.

Plan: summarize adaptive bandwidth selection, effective local sample support,
local coefficient dispersion, pointwise uncertainty, multiplicity, and
adjacent-bandwidth stability for the four continuous spatial models.
Framework: AnaSOP Sections 5-7, continuous-spatial workflow steps 7 and 12.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import norm
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/exp/analysis/climate-npp/continuous-spatial-heterogeneity"
FIGURE_EVIDENCE = ROOT / "data/exp/analysis/climate-npp/gwr-bandwidth-and-effective-sample-diagnostics"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/gwr-and-local-regression-diagnostics"
OUTPUT = ROOT / "data/results/tables/Table_gwr_and_local_regression_diagnostics.xlsx"

BANDWIDTH = SOURCE / "adaptive_bandwidth_aicc_and_prediction_diagnostics.csv"
SURFACES = SOURCE / "continuous_local_slope_surfaces.parquet"

MODEL_ORDER = [
    "Heat to cropland NPP",
    "Dry spell to cropland NPP",
    "Cropland NPP to total consumption",
    "Cropland NPP to food consumption",
]
MODEL_LABELS = {
    "Heat to cropland NPP": "Heat → NPP (per 10 heat days)",
    "Dry spell to cropland NPP": "Dry spell → NPP (per 10 dry days)",
    "Cropland NPP to total consumption": "NPP → total consumption (% per 0.1 NPP)",
    "Cropland NPP to food consumption": "NPP → food consumption (% per 0.1 NPP)",
}
STATUS = {
    "Heat to cropland NPP": "Main-text supported",
    "Dry spell to cropland NPP": "Main-text supported",
    "Cropland NPP to total consumption": "Appendix diagnostic only",
    "Cropland NPP to food consumption": "Appendix diagnostic only",
}


def benjamini_hochberg(p_values: np.ndarray) -> np.ndarray:
    p_values = np.asarray(p_values, dtype=float)
    order = np.argsort(p_values)
    ranked = p_values[order] * len(p_values) / np.arange(1, len(p_values) + 1)
    adjusted_sorted = np.minimum.accumulate(ranked[::-1])[::-1]
    adjusted = np.empty(len(p_values), dtype=float)
    adjusted[order] = np.minimum(adjusted_sorted, 1.0)
    return adjusted


def build_tables() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    bandwidth = pd.read_csv(BANDWIDTH)
    surfaces = pd.read_parquet(SURFACES)
    selected = (
        bandwidth.loc[bandwidth.groupby("Model")["AICc"].idxmin()]
        .set_index("Model")
        .loc[MODEL_ORDER]
        .reset_index()
    )

    support_rows: list[dict[str, object]] = []
    coefficient_rows: list[dict[str, object]] = []
    combined_rows: list[dict[str, object]] = []
    for model in MODEL_ORDER:
        local = surfaces.loc[surfaces["Model"].eq(model)].copy()
        selected_row = selected.loc[selected["Model"].eq(model)].iloc[0]
        model_bandwidth = bandwidth.loc[bandwidth["Model"].eq(model)]
        minimum_rmse = float(model_bandwidth["Leave-One-Village-Out RMSE"].min())
        selected_rmse = float(selected_row["Leave-One-Village-Out RMSE"])

        p_values = 2.0 * norm.sf(
            np.abs(local["Local Slope"].to_numpy(float) / local["Local Standard Error"].to_numpy(float))
        )
        bh_adjusted = benjamini_hochberg(p_values)
        interval_excludes_zero = (
            local["Local 95 Percent CI Lower"].gt(0)
            | local["Local 95 Percent CI Upper"].lt(0)
        )
        effective = local["Effective Local Sample"]
        slopes = local["Local Slope"]
        stable = local["Slope Sign Stable Across Adjacent Bandwidths"].astype(bool)
        underlying_observations = 59_135 if "to cropland NPP" in model else 43_120 if "total" in model else 43_365

        support = {
            "Model": MODEL_LABELS[model],
            "Locations": int(len(local)),
            "Underlying observations": underlying_observations,
            "Selected neighbours": int(selected_row["Adaptive Neighbor Bandwidth"]),
            "Selected AICc": float(selected_row["AICc"]),
            "LOVO RMSE": selected_rmse,
            "RMSE above minimum": selected_rmse / minimum_rmse - 1.0,
            "Effective N P5": float(effective.quantile(0.05)),
            "Effective N median": float(effective.median()),
            "Effective N P95": float(effective.quantile(0.95)),
            "Status": STATUS[model],
        }
        coefficients = {
            "Model": MODEL_LABELS[model],
            "Local slope P5": float(slopes.quantile(0.05)),
            "Local slope median": float(slopes.median()),
            "Local slope P95": float(slopes.quantile(0.95)),
            "Negative slope share": float(slopes.lt(0).mean()),
            "Pointwise 95% CI excludes zero": float(interval_excludes_zero.mean()),
            "BH-FDR q<0.05 share": float(np.mean(bh_adjusted < 0.05)),
            "Adjacent-bandwidth sign stable": float(stable.mean()),
            "Interpretation": STATUS[model],
        }
        support_rows.append(support)
        coefficient_rows.append(coefficients)
        combined_rows.append({**support, **{key: value for key, value in coefficients.items() if key != "Model"}})

    support_table = pd.DataFrame(support_rows)
    coefficient_table = pd.DataFrame(coefficient_rows)
    combined = pd.DataFrame(combined_rows)

    assert support_table["Selected neighbours"].tolist() == [40, 30, 50, 60]
    assert np.allclose(
        support_table["Effective N median"].to_numpy(),
        [471.9301825892437, 350.18513831079315, 22.43324657818586, 26.865766864454653],
    )
    assert coefficient_table.loc[2, "BH-FDR q<0.05 share"] == 0
    assert coefficient_table.loc[3, "BH-FDR q<0.05 share"] == 0

    reviewed_selected = pd.read_csv(FIGURE_EVIDENCE / "selected_adaptive_bandwidth_diagnostics.csv")
    reviewed_selected = reviewed_selected.set_index("Model").loc[MODEL_ORDER]
    assert reviewed_selected["Adaptive Neighbor Bandwidth"].astype(int).tolist() == [40, 30, 50, 60]

    EVIDENCE.mkdir(parents=True, exist_ok=True)
    support_table.to_csv(EVIDENCE / "gwr_bandwidth_and_effective_sample_diagnostics.csv", index=False)
    coefficient_table.to_csv(EVIDENCE / "local_coefficient_and_multiplicity_diagnostics.csv", index=False)
    combined.to_csv(EVIDENCE / "gwr_and_local_regression_diagnostics_full.csv", index=False)
    return support_table, coefficient_table, combined


def write_workbook(support: pd.DataFrame, coefficients: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "GWR Diagnostics"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    blue = "3F7CAC"
    teal = "3A9D8F"
    gold = "D6A84B"
    white = "FFFFFF"
    light = "F4F6F7"
    stage1_fill = "E8F0F7"
    stage2_fill = "FFF2CC"
    thin = Side(style="thin", color="C8D5DE")
    medium = Side(style="medium", color="7F9DB9")

    ws.merge_cells("A1:J1")
    ws["A1"] = "GWR and Local-Regression Diagnostics"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 29

    ws.merge_cells("A3:J3")
    ws["A3"] = "A. Adaptive bandwidth, prediction, and effective local sample"
    ws["A3"].fill = PatternFill("solid", fgColor=navy)
    ws["A3"].font = Font(name="Times New Roman", size=9.7, bold=True, color=white)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 23

    support_columns = [
        "Model",
        "Locations",
        "Underlying observations",
        "Selected neighbours",
        "Selected AICc",
        "LOVO RMSE",
        "RMSE above minimum",
        "Effective N P5",
        "Effective N median",
        "Effective N P95",
    ]
    support_headers = [
        "Model",
        "Locations",
        "Underlying\nobservations",
        "Selected\nneighbours",
        "Selected\nAICc",
        "LOVO\nRMSE",
        "RMSE above\nminimum",
        "Effective N\nP5",
        "Effective N\nmedian",
        "Effective N\nP95",
    ]
    for column, header in enumerate(support_headers, start=1):
        cell = ws.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor=blue if column <= 7 else teal)
        cell.font = Font(name="Times New Roman", size=8.6, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[4].height = 37

    for row_number, record in enumerate(support[support_columns].itertuples(index=False, name=None), start=5):
        for column, value in enumerate(record, start=1):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="Times New Roman", size=8.8, bold=column == 1)
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "right",
                vertical="center",
                wrap_text=True,
                indent=1 if column == 1 else 0,
            )
            cell.border = Border(bottom=thin)
            cell.fill = PatternFill("solid", fgColor=stage1_fill if row_number <= 6 else stage2_fill)
            if column in {2, 3, 4}:
                cell.number_format = "#,##0"
            elif column == 5:
                cell.number_format = "#,##0.0"
            elif column == 6:
                cell.number_format = "0.0000"
            elif column == 7:
                cell.number_format = "0.00%"
            elif column in {8, 9, 10}:
                cell.number_format = "0.0"
        ws.row_dimensions[row_number].height = 28

    ws.merge_cells("A10:I10")
    ws["A10"] = "B. Local coefficient dispersion, uncertainty, and multiplicity"
    ws["A10"].fill = PatternFill("solid", fgColor=navy)
    ws["A10"].font = Font(name="Times New Roman", size=9.7, bold=True, color=white)
    ws["A10"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws["J10"] = "Interpretation"
    ws["J10"].fill = PatternFill("solid", fgColor=navy)
    ws["J10"].font = Font(name="Times New Roman", size=9.7, bold=True, color=white)
    ws["J10"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[10].height = 23

    coefficient_headers = [
        "Model",
        "Local slope\nP5",
        "Local slope\nmedian",
        "Local slope\nP95",
        "Negative slope\nshare",
        "Pointwise 95% CI\nexcludes zero",
        "BH-FDR q<0.05\nshare",
        "Adjacent-bandwidth\nsign stable",
        "",
        "Interpretation",
    ]
    for column, header in enumerate(coefficient_headers, start=1):
        if column == 9:
            continue
        cell = ws.cell(11, column, header)
        cell.fill = PatternFill("solid", fgColor=teal if column in {5, 6, 7, 8} else blue if column < 9 else gold)
        cell.font = Font(name="Times New Roman", size=8.5, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.merge_cells("H11:I11")
    ws["H11"] = "Adjacent-bandwidth\nsign stable"
    ws["H11"].fill = PatternFill("solid", fgColor=teal)
    ws["H11"].font = Font(name="Times New Roman", size=8.5, bold=True, color=white)
    ws["H11"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[11].height = 38

    coefficient_columns = [
        "Model",
        "Local slope P5",
        "Local slope median",
        "Local slope P95",
        "Negative slope share",
        "Pointwise 95% CI excludes zero",
        "BH-FDR q<0.05 share",
        "Adjacent-bandwidth sign stable",
        "Interpretation",
    ]
    for row_number, record in enumerate(coefficients[coefficient_columns].itertuples(index=False, name=None), start=12):
        values = list(record[:8]) + [record[8]]
        target_columns = [1, 2, 3, 4, 5, 6, 7, 8, 10]
        ws.merge_cells(start_row=row_number, start_column=8, end_row=row_number, end_column=9)
        for column, value in zip(target_columns, values, strict=True):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="Times New Roman", size=8.8, bold=column in {1, 10})
            cell.alignment = Alignment(
                horizontal="left" if column in {1, 10} else "right",
                vertical="center",
                wrap_text=True,
                indent=1 if column in {1, 10} else 0,
            )
            cell.border = Border(bottom=thin)
            cell.fill = PatternFill("solid", fgColor=stage1_fill if row_number <= 13 else stage2_fill)
            if column in {2, 3, 4}:
                cell.number_format = "0.000" if row_number <= 13 else "0.00"
            elif column in {5, 6, 7, 8}:
                cell.number_format = "0.0%"
        ws.row_dimensions[row_number].height = 30

    note_row = 17
    notes = [
        "Notes: Bandwidth is selected by minimum AICc without reference to local coefficient signs or significance. LOVO is leave-one-village-out prediction; effective N is the Kish-equivalent local sample under kernel and survey weights.",
        "Pointwise intervals are unadjusted. BH-FDR applies Benjamini–Hochberg adjustment across mapped locations within each model and is reported only as a multiplicity diagnostic because local estimates are spatially dependent.",
        "Blue rows are Stage 1 and support the main continuous-spatial result. Gold rows are Stage 2; low effective local samples and zero BH-FDR support restrict them to appendix diagnostics.",
    ]
    for note in notes:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
        cell = ws.cell(note_row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.0, italic=True, color="3F4B52")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[note_row].height = 23
        note_row += 1

    widths = [37, 13, 17, 15, 16, 14, 17, 14, 14, 24]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 78
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.22
    ws.page_margins.top = ws.page_margins.bottom = 0.22
    ws.print_area = "A1:J19"
    wb.properties.title = "GWR and Local-Regression Diagnostics"
    wb.properties.subject = "Adaptive bandwidth, local support, uncertainty, and multiplicity audit"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["GWR Diagnostics"]
    ws = wb["GWR Diagnostics"]
    assert ws.max_row == 19
    assert ws.max_column == 10
    assert not ws.merged_cells.ranges
    assert ws["A5"].value.startswith("Heat → NPP")
    assert ws["D5"].value == 40
    assert ws["D8"].value == 60
    assert ws["G14"].value == 0
    assert ws["G15"].value == 0
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    support, coefficients, _ = build_tables()
    write_workbook(support, coefficients)
    validate_workbook()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print("\nBandwidth and support")
    print(support.to_string(index=False))
    print("\nLocal coefficient and multiplicity")
    print(coefficients.to_string(index=False))


if __name__ == "__main__":
    main()
