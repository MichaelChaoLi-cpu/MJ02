#!/usr/bin/env python3
"""Build the Appendix table of failed pathways and buffering boundary tests."""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break


ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "data/exp/analysis/climate-welfare"
FAILED = BASE / "failed-annual-full-season-pathways"
AGRICULTURE = BASE / "cses-direct-monsoon-diagnostic"
BUFFERING = BASE / "absolute-heat-food-buffering"
ANALYSIS_OUTPUT = BASE / "failed-pathways-buffering-tests/table_rows.csv"
OUTPUT = (
    ROOT
    / "data/exp/internal_output_archive/tables"
    / "Table_failed_pathways_and_buffering_tests_superseded.xlsx"
)

COLUMNS = [
    "Evidence family",
    "Outcome",
    "Exposure / modifier",
    "Specification",
    "Sample",
    "Estimate / metric",
    "95% interval",
    "Raw p-value",
    "Adjusted / gate metric",
    "Fixed effects / inference",
    "Gate status",
    "Interpretation limit",
]


def exposure_label(variable: str) -> str:
    if "Longest Intraseasonal Dry Spell" in variable or "Longest Dry Spell" in variable:
        return "Longest dry-spell anomaly (1 SD)"
    if "Absolute Heat" in variable:
        return "Absolute heat days ≥35°C (per 10)"
    if "May-October Precipitation" in variable:
        return "May–October rainfall anomaly (1 SD)"
    if "Wet-Season Onset" in variable:
        return "Wet-season onset anomaly (1 SD)"
    if "Hot-Dry Day Count" in variable:
        return "Relative hot-dry anomaly (1 SD; legacy)"
    return variable


def ecological_rows() -> list[dict[str, object]]:
    npp = pd.read_csv(FAILED / "annual_npp_coefficients.csv")
    npp_gate = pd.read_csv(FAILED / "annual_npp_prediction_gate.csv")
    vegetation = pd.read_csv(FAILED / "full_season_vegetation_coefficients.csv")
    vegetation_gate = pd.read_csv(FAILED / "full_season_evi_prediction_gate.csv")
    npp_summary = json.loads((BASE / "model-a-monsoon-npp/results_summary.json").read_text())
    vegetation_summary = json.loads(
        (BASE / "model-a-seasonal-vegetation/results_summary.json").read_text()
    )

    rows: list[dict[str, object]] = []
    focal_npp = npp.loc[
        npp["Variable"].str.contains("Longest Intraseasonal Dry Spell|Absolute Heat", regex=True)
    ]
    for record in focal_npp.to_dict("records"):
        all_land = record["Outcome"] == "All-land NPP"
        observations = (
            int(npp_summary["primary_sample_observations"])
            if all_land else int(record["Observations"])
        )
        villages = (
            int(npp_summary["primary_sample_villages"])
            if all_land else int(record["Villages"])
        )
        rows.append(
            {
                "Evidence family": "Annual NPP pathway",
                "Outcome": record["Outcome"],
                "Exposure / modifier": exposure_label(str(record["Variable"])),
                "Specification": "Candidate B; 35°C; 5 km; joint climate model",
                "Sample": f"N={observations:,}; {villages:,} villages",
                "Estimate / metric": f'{float(record["Coefficient"]):.5f} kg C m⁻²',
                "95% interval": f'[{float(record["95 Percent CI Lower"]):.5f}, {float(record["95 Percent CI Upper"]):.5f}]',
                "Raw p-value": float(record["Probability Value"]),
                "Adjusted / gate metric": "Joint held-out RMSE change = −1.53%",
                "Fixed effects / inference": "Village + year FE; spatial-block clustered SE",
                "Gate status": "Not promoted: annual-NPP prediction gate failed",
                "Interpretation limit": "A significant component cannot rescue worse held-out prediction",
            }
        )
    for record in npp_gate.to_dict("records"):
        improvement = float(record["Improvement versus rainfall only percent"])
        rows.append(
            {
                "Evidence family": "Annual NPP prediction",
                "Outcome": "Held-out annual NPP",
                "Exposure / modifier": record["Model"],
                "Specification": "Strict spatial × temporal cross-fit",
                "Sample": "N=312,072 strict-support village-years",
                "Estimate / metric": f'RMSE={float(record["RMSE"]):.5f}',
                "95% interval": "—",
                "Raw p-value": np.nan,
                "Adjusted / gate metric": f"RMSE improvement={improvement:+.2f}%",
                "Fixed effects / inference": "Same folds and observations as rainfall benchmark",
                "Gate status": "Failed: prediction is worse than rainfall only",
                "Interpretation limit": "Annual aggregation is not the promoted ecological window",
            }
        )

    focal_vegetation = vegetation.loc[
        vegetation["Variable"].str.contains("Longest Intraseasonal Dry Spell|Absolute Heat", regex=True)
    ]
    for record in focal_vegetation.to_dict("records"):
        rows.append(
            {
                "Evidence family": "Full-season vegetation pathway",
                "Outcome": record["Outcome"],
                "Exposure / modifier": exposure_label(str(record["Variable"])),
                "Specification": "Candidate B; 35°C; 5 km; May–February outcome",
                "Sample": f'N={int(vegetation_summary["sample_observations"]):,}; {int(vegetation_summary["sample_villages"]):,} villages',
                "Estimate / metric": f'{float(record["Coefficient"]):.4f} outcome SD',
                "95% interval": f'[{float(record["95 Percent CI Lower"]):.4f}, {float(record["95 Percent CI Upper"]):.4f}]',
                "Raw p-value": float(record["Probability Value"]),
                "Adjusted / gate metric": "Joint Wald p=0.410; held-out change=−0.54%",
                "Fixed effects / inference": "Village + year FE; spatial-block clustered SE",
                "Gate status": "Not promoted: joint full-season gate failed",
                "Interpretation limit": "Full-season averaging is sensor-sensitive and masks phase timing",
            }
        )
    for record in vegetation_gate.to_dict("records"):
        improvement = float(record["Improvement versus rainfall only percent"])
        rows.append(
            {
                "Evidence family": "Full-season EVI prediction",
                "Outcome": "Held-out May–February EVI",
                "Exposure / modifier": record["Model"],
                "Specification": "Strict spatial × temporal cross-fit",
                "Sample": f'N={int(vegetation_summary["sample_observations"]):,}; {int(vegetation_summary["sample_villages"]):,} villages',
                "Estimate / metric": f'RMSE={float(record["RMSE"]):.4f}',
                "95% interval": "—",
                "Raw p-value": np.nan,
                "Adjusted / gate metric": f"RMSE improvement={improvement:+.2f}%",
                "Fixed effects / inference": "Same folds and observations as rainfall benchmark",
                "Gate status": (
                    "Insufficient alone: joint model and Wald gates fail"
                    if improvement > 0 else "Failed: prediction is worse than rainfall only"
                ),
                "Interpretation limit": "Does not replace the supported post-monsoon outcome",
            }
        )
    return rows


def agricultural_rows() -> list[dict[str, object]]:
    coefficients = pd.read_csv(AGRICULTURE / "coefficients.csv")
    models = pd.read_csv(AGRICULTURE / "model_summary.csv").set_index("Outcome")
    coefficients = coefficients.loc[
        coefficients["Outcome"].isin(
            ["Agricultural Participation Probability", "Asinh Real Crop Production per ha"]
        )
    ]
    rows: list[dict[str, object]] = []
    for record in coefficients.to_dict("records"):
        outcome = str(record["Outcome"])
        summary = models.loc[outcome]
        participation = outcome == "Agricultural Participation Probability"
        scale = 100.0 if participation else 1.0
        unit = "percentage points" if participation else "asinh units"
        legacy = "Hot-Dry Day Count" in str(record["Exposure"])
        rows.append(
            {
                "Evidence family": "Agricultural survey pathway",
                "Outcome": (
                    "Agricultural participation" if participation
                    else "Real crop production per cultivated ha (asinh)"
                ),
                "Exposure / modifier": exposure_label(str(record["Exposure"])),
                "Specification": "Last completed season; Candidate B; 5 km",
                "Sample": f'N={int(record["Observations"]):,}; {int(record["Villages"]):,} villages',
                "Estimate / metric": f'{scale * float(record["Coefficient"]):.3f} {unit}',
                "95% interval": f'[{scale * float(record["95 Percent CI Lower"]):.3f}, {scale * float(record["95 Percent CI Upper"]):.3f}]',
                "Raw p-value": float(record["Probability Value"]),
                "Adjusted / gate metric": f'Joint monsoon-family Wald p={float(summary["Joint monsoon-structure Wald probability"]):.3f}',
                "Fixed effects / inference": "District + wave×month FE; weights; block-clustered SE",
                "Gate status": "Human-link agricultural gate not passed",
                "Interpretation limit": (
                    "Legacy relative hot-dry diagnostic; not current absolute-heat evidence"
                    if legacy else "Mixed agricultural coefficients do not establish a production pathway"
                ),
            }
        )
    return rows


def buffering_rows() -> list[dict[str, object]]:
    interactions = pd.read_csv(BUFFERING / "buffer_interactions.csv")
    rows: list[dict[str, object]] = []
    for record in interactions.to_dict("records"):
        expected = str(record["Expected Protective Direction"])
        consistent = "yes" if bool(record["Direction Consistent"]) else "no"
        rows.append(
            {
                "Evidence family": "Buffering boundary test",
                "Outcome": "Log real food consumption/member",
                "Exposure / modifier": f'Heat × {record["Buffer"]}',
                "Specification": "Candidate B; ≥35°C; 5 km; per 10 heat days",
                "Sample": f'N={int(record["Observations"]):,}; {int(record["Villages"]):,} villages',
                "Estimate / metric": f'{float(record["Coefficient Log Points"]):.4f} log points',
                "95% interval": f'[{float(record["95 Percent CI Lower"]):.4f}, {float(record["95 Percent CI Upper"]):.4f}]',
                "Raw p-value": float(record["Raw Probability Value"]),
                "Adjusted / gate metric": f'Holm p={float(record["Holm Adjusted Probability Value"]):.3f}; expected {expected}; consistent={consistent}',
                "Fixed effects / inference": "District + wave×month FE; weights; block-clustered SE",
                "Gate status": "Protective-buffer gate not passed",
                "Interpretation limit": "Heterogeneity association only; not an intervention or adaptation effect",
            }
        )
    return rows


def main() -> None:
    table = pd.DataFrame(
        ecological_rows() + agricultural_rows() + buffering_rows(), columns=COLUMNS
    )
    assert len(table) == 23, f"Expected 23 rows, found {len(table)}"
    ANALYSIS_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(ANALYSIS_OUTPUT, index=False)

    wb = Workbook()
    ws = wb.active
    ws.title = "Failed pathways"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    header_fill = PatternFill("solid", fgColor="E5E5E5")
    ecology_fill = PatternFill("solid", fgColor="EAF3EE")
    prediction_fill = PatternFill("solid", fgColor="EAF1F8")
    agriculture_fill = PatternFill("solid", fgColor="F7F1E8")
    buffering_fill = PatternFill("solid", fgColor="F1ECF6")

    def styled_cell(row: int, column: int, value, *, bold: bool = False, fill=None) -> None:
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(name="Times New Roman", size=9.2, bold=bold)
        cell.alignment = Alignment(
            horizontal="left" if column == 1 else "center",
            vertical="center",
            wrap_text=True,
        )
        if fill is not None:
            cell.fill = fill
        if isinstance(value, (float, np.floating)):
            cell.number_format = "0.000"

    def merged_value(row: int, start: int, end: int, value, *, bold: bool = False, fill=None) -> None:
        if end > start:
            ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        styled_cell(row, start, value, bold=bold, fill=fill)
        if fill is not None:
            for column in range(start, end + 1):
                ws.cell(row=row, column=column).fill = fill

    def one_row(family: str, outcome: str | None = None, exposure: str | None = None) -> pd.Series:
        selected = table.loc[table["Evidence family"].eq(family)]
        if outcome is not None:
            selected = selected.loc[selected["Outcome"].eq(outcome)]
        if exposure is not None:
            selected = selected.loc[selected["Exposure / modifier"].str.contains(exposure, regex=False)]
        if len(selected) != 1:
            raise ValueError(
                f"Expected one row for family={family!r}, outcome={outcome!r}, exposure={exposure!r}; found {len(selected)}"
            )
        return selected.iloc[0]

    # Page 1: ecological pathways and their held-out prediction gates.
    ws.merge_cells("A1:I1")
    ws["A1"] = "Failed Pathways and Buffering Tests"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25
    ecology_specs = [
        ("All-land NPP", "Annual NPP pathway", "All-land NPP", None, ecology_fill),
        ("Cropland-weighted NPP", "Annual NPP pathway", "Cropland-weighted NPP", None, ecology_fill),
        ("Full-season EVI", "Full-season vegetation pathway", "Full-season EVI", None, ecology_fill),
        ("Full-season NDVI", "Full-season vegetation pathway", "Full-season NDVI", None, ecology_fill),
        ("NPP: monsoon only", "Annual NPP prediction", None, "Monsoon structure only", prediction_fill),
        ("NPP: joint model", "Annual NPP prediction", None, "Joint rainfall and monsoon structure", prediction_fill),
        ("EVI: monsoon only", "Full-season EVI prediction", None, "Monsoon structure only", prediction_fill),
        ("EVI: joint model", "Full-season EVI prediction", None, "Joint rainfall and monsoon structure", prediction_fill),
    ]
    styled_cell(3, 1, "Statistic", bold=True, fill=header_fill)
    for column, (label, _, _, _, fill) in enumerate(ecology_specs, start=2):
        styled_cell(3, column, label, bold=True, fill=header_fill)
        ws.cell(3, column).border = Border(top=medium, bottom=thin)
    ws.cell(3, 1).border = Border(top=medium, bottom=thin)
    ws.row_dimensions[3].height = 36

    coefficient_records: dict[int, tuple[pd.Series, pd.Series]] = {}
    prediction_records: dict[int, pd.Series] = {}
    for column, (_, family, outcome, exposure, _) in enumerate(ecology_specs, start=2):
        if "prediction" in family.lower():
            prediction_records[column] = one_row(family, exposure=exposure)
        else:
            coefficient_records[column] = (
                one_row(family, outcome=outcome, exposure="Longest dry-spell"),
                one_row(family, outcome=outcome, exposure="Absolute heat"),
            )

    page1_rows = [
        (4, "Specification"),
        (6, "Longest dry-spell coefficient"),
        (7, "95% interval"),
        (8, "Raw p-value"),
        (10, "Absolute-heat coefficient"),
        (11, "95% interval"),
        (12, "Raw p-value"),
        (14, "Held-out RMSE"),
        (15, "Improvement vs rainfall only"),
        (17, "Sample"),
        (18, "Fixed effects / inference"),
        (19, "Gate status"),
        (20, "Interpretation limit"),
    ]
    for row, label in page1_rows:
        styled_cell(row, 1, label, bold=label in {"Gate status", "Interpretation limit"})
        ws.row_dimensions[row].height = 24 if row not in {18, 19, 20} else 32
    for column, (_, _, _, _, fill) in enumerate(ecology_specs, start=2):
        if column in coefficient_records:
            dry, heat = coefficient_records[column]
            values = {
                4: dry["Specification"],
                6: dry["Estimate / metric"], 7: dry["95% interval"], 8: dry["Raw p-value"],
                10: heat["Estimate / metric"], 11: heat["95% interval"], 12: heat["Raw p-value"],
                17: dry["Sample"], 18: dry["Fixed effects / inference"],
                19: dry["Gate status"], 20: dry["Interpretation limit"],
            }
        else:
            prediction = prediction_records[column]
            values = {
                4: prediction["Specification"], 14: prediction["Estimate / metric"],
                15: str(prediction["Adjusted / gate metric"]).replace("RMSE improvement=", ""),
                17: prediction["Sample"],
                18: prediction["Fixed effects / inference"], 19: prediction["Gate status"],
                20: prediction["Interpretation limit"],
            }
        for row, _ in page1_rows:
            styled_cell(row, column, values.get(row, "—"), fill=fill)
    for column in range(1, 10):
        ws.cell(20, column).border = Border(bottom=medium)

    page1_notes = [
        "Notes: Coefficient significance does not override a failed held-out prediction or cross-sensor promotion gate.",
        "Annual NPP is measured in kg C m⁻²; full-season vegetation coefficients are outcome SD. Prediction rows use strict spatial-by-temporal cross-fitting.",
    ]
    for offset, note in enumerate(page1_notes, start=22):
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=9)
        cell = ws.cell(offset, 1, note)
        cell.font = Font(name="Times New Roman", size=8.5, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[offset].height = 18

    # Page 2: agricultural outcomes and buffering interactions.
    page2_title = 26
    page2_header = 28
    ws.merge_cells(start_row=page2_title, start_column=1, end_row=page2_title, end_column=9)
    ws.cell(page2_title, 1, "Failed Pathways and Buffering Tests (continued)")
    ws.cell(page2_title, 1).font = Font(name="Times New Roman", size=14, bold=True)
    ws.cell(page2_title, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[page2_title].height = 25

    page2_specs = [
        (2, 3, "Agricultural participation", one_row("Agricultural survey pathway", outcome="Agricultural participation", exposure="May–October rainfall"), agriculture_fill),
        (4, 5, "Crop production / ha", one_row("Agricultural survey pathway", outcome="Real crop production per cultivated ha (asinh)", exposure="May–October rainfall"), agriculture_fill),
        (6, 6, "Heat × irrigation", one_row("Buffering boundary test", exposure="Irrigation"), buffering_fill),
        (7, 7, "Heat × road access", one_row("Buffering boundary test", exposure="Historical Road Access"), buffering_fill),
        (8, 9, "Heat × connectivity", one_row("Buffering boundary test", exposure="Baseline Settlement Connectivity"), buffering_fill),
    ]
    styled_cell(page2_header, 1, "Statistic", bold=True, fill=header_fill)
    ws.cell(page2_header, 1).border = Border(top=medium, bottom=thin)
    for start, end, label, _, _ in page2_specs:
        merged_value(page2_header, start, end, label, bold=True, fill=header_fill)
        for column in range(start, end + 1):
            ws.cell(page2_header, column).border = Border(top=medium, bottom=thin)
    ws.row_dimensions[page2_header].height = 36

    page2_rows = [
        (29, "Specification"),
        (31, "Rainfall coefficient"), (32, "95% interval"), (33, "Raw p-value"),
        (35, "Onset coefficient"), (36, "95% interval"), (37, "Raw p-value"),
        (39, "Dry-spell coefficient"), (40, "95% interval"), (41, "Raw p-value"),
        (43, "Relative hot-dry coefficient (legacy)"), (44, "95% interval"), (45, "Raw p-value"),
        (47, "Heat × modifier coefficient"), (48, "95% interval"), (49, "Raw p-value"),
        (50, "Holm-adjusted p-value"), (52, "Sample"), (53, "Fixed effects / inference"),
        (54, "Gate status"), (55, "Interpretation limit"),
    ]
    for row, label in page2_rows:
        styled_cell(row, 1, label, bold=label in {"Gate status", "Interpretation limit"})
        ws.row_dimensions[row].height = 23 if row not in {53, 54, 55} else 32

    exposure_rows = [
        ("May–October rainfall", 31, 32, 33),
        ("Wet-season onset", 35, 36, 37),
        ("Longest dry-spell", 39, 40, 41),
        ("Relative hot-dry", 43, 44, 45),
    ]
    for start, end, _, representative, fill in page2_specs:
        is_buffer = representative["Evidence family"] == "Buffering boundary test"
        values: dict[int, object] = {
            29: representative["Specification"], 52: representative["Sample"],
            53: representative["Fixed effects / inference"], 54: representative["Gate status"],
            55: representative["Interpretation limit"],
        }
        if is_buffer:
            values.update(
                {47: representative["Estimate / metric"], 48: representative["95% interval"],
                 49: representative["Raw p-value"], 50: str(representative["Adjusted / gate metric"]).split(";")[0].replace("Holm p=", "")}
            )
        else:
            outcome = str(representative["Outcome"])
            for exposure, estimate_row, interval_row, probability_row in exposure_rows:
                record = one_row("Agricultural survey pathway", outcome=outcome, exposure=exposure)
                values.update(
                    {estimate_row: record["Estimate / metric"], interval_row: record["95% interval"], probability_row: record["Raw p-value"]}
                )
        for row, _ in page2_rows:
            merged_value(row, start, end, values.get(row, "—"), fill=fill)
    for column in range(1, 10):
        ws.cell(55, column).border = Border(bottom=medium)

    page2_notes = [
        "Notes: Agricultural rows use CSES weights and 0.75° spatial-block-clustered uncertainty; relative hot-dry exposure is a legacy diagnostic and is not current absolute-heat evidence.",
        "The three heat-buffer interactions use Candidate B, ≥35°C, 5 km exposure. Holm adjustment is applied across the three tests; none supports a protective or causal adaptation claim.",
    ]
    for offset, note in enumerate(page2_notes, start=57):
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=9)
        cell = ws.cell(offset, 1, note)
        cell.font = Font(name="Times New Roman", size=8.5, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[offset].height = 18

    widths = [36] + [20] * 8
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    ws.page_margins.left = 0.18
    ws.page_margins.right = 0.18
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.print_area = "A1:I58"
    ws.row_breaks.append(Break(id=24))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == ["Failed pathways"]
    assert table["Evidence family"].value_counts().to_dict() == {
        "Agricultural survey pathway": 8,
        "Annual NPP pathway": 4,
        "Full-season vegetation pathway": 4,
        "Buffering boundary test": 3,
        "Annual NPP prediction": 2,
        "Full-season EVI prediction": 2,
    }
    assert table.loc[
        table["Evidence family"].eq("Buffering boundary test"), "Adjusted / gate metric"
    ].str.contains("Holm p=1.000").all()
    assert not table["Gate status"].str.contains(
        r"(?:^|\b)(?:gate passed|promoted)$", case=False, regex=True
    ).any()
    for row in check.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not any(error in cell.value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?"))

    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(table.groupby("Evidence family", sort=False).size().to_string())


if __name__ == "__main__":
    main()
