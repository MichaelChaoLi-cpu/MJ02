#!/usr/bin/env python3
"""Food-Consumption Response and Validation.

Plan: Present the eight frozen household and commune pseudo-panel specifications
in one Stargazer-style regression table.
Framework: AnaSOP Model B with area and survey-time fixed effects, survey
weights, and spatial-block-clustered standard errors.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / (
    "data/exp/analysis/climate-welfare/household-heat-identification-repair/"
    "coefficients.csv"
)
OUTPUT = ROOT / "data/exp/internal_output_archive/obsolete-formal-results-20260824/tables/Table_food_consumption_response_and_validation.xlsx"

MODELS = [
    ("(1)", "All linked households, district FE", "All linked", "District"),
    ("(2)", "All linked households, commune FE", "All linked", "Commune"),
    ("(3)", "Repeated-village sample, district FE", "Repeated villages", "District"),
    ("(4)", "Repeated-village sample, village FE", "Repeated villages", "Village"),
    ("(5)", "Single-wave-village sample, district FE", "Single-wave villages", "District"),
    ("(6)", "Commune pseudo-panel, minimum 5 households", "Commune cells >=5", "Commune"),
    ("(7)", "Commune pseudo-panel, minimum 10 households", "Commune cells >=10", "Commune"),
    (
        "(8)",
        "Commune pseudo-panel, minimum 5 households, composition adjusted",
        "Commune cells >=5",
        "Commune",
    ),
]

EXPOSURES = [
    ("May-October precipitation anomaly", "Rainfall Anomaly Z"),
    ("Wet-season onset-date anomaly", "Onset Anomaly Z"),
    ("Longest intraseasonal dry-spell anomaly", "Dry Spell Anomaly Z"),
    ("Post-onset >=35 C days / 10", "Absolute Heat Days per 10"),
]


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def result_row(frame: pd.DataFrame, specification: str, exposure: str) -> pd.Series:
    rows = frame.loc[
        frame["Specification"].eq(specification) & frame["Exposure"].eq(exposure)
    ]
    if len(rows) != 1:
        raise ValueError(
            f"Expected one result for {specification!r}, {exposure!r}; found {len(rows)}"
        )
    return rows.iloc[0]


def style_cell(
    cell,
    *,
    bold: bool = False,
    italic: bool = False,
    horizontal: str = "center",
) -> None:
    cell.font = Font(
        name="Times New Roman", size=10.5, bold=bold, italic=italic, color="000000"
    )
    cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)


def main() -> None:
    results = pd.read_csv(INPUT)
    wb = Workbook()
    ws = wb.active
    ws.title = "Regression results"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B7"

    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    header_fill = PatternFill("solid", fgColor="EDEDED")
    heat_fill = PatternFill("solid", fgColor="EAF3EE")
    last_column = len(MODELS) + 1
    last_letter = get_column_letter(last_column)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    ws["A1"] = "Food-Consumption Response and Validation"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=6)
    ws.merge_cells(start_row=2, start_column=7, end_row=2, end_column=9)
    ws["B2"] = "Household regressions"
    ws["G2"] = "Commune survey-time pseudo-panel"
    for cell in (ws["B2"], ws["G2"]):
        style_cell(cell, bold=True)
    for column in range(1, last_column + 1):
        ws.cell(2, column).border = Border(top=medium, bottom=thin)

    for column, (number, _, sample_label, area_label) in enumerate(MODELS, start=2):
        ws.cell(3, column, number)
        ws.cell(4, column, sample_label)
        ws.cell(5, column, f"{area_label} FE")
        style_cell(ws.cell(3, column), bold=True)
        style_cell(ws.cell(4, column))
        style_cell(ws.cell(5, column))
        ws.cell(3, column).fill = header_fill
    ws["A4"] = "Analytical sample"
    ws["A5"] = "Area fixed effects"
    for cell in (ws["A4"], ws["A5"]):
        style_cell(cell, horizontal="left")
    for column in range(1, last_column + 1):
        ws.cell(5, column).border = Border(bottom=thin)

    current_row = 7
    heat_row = None
    for label, exposure in EXPOSURES:
        ws.cell(current_row, 1, label)
        style_cell(ws.cell(current_row, 1), horizontal="left")
        is_heat = exposure == "Absolute Heat Days per 10"
        if is_heat:
            heat_row = current_row
        for column, (_, specification, _, _) in enumerate(MODELS, start=2):
            row = result_row(results, specification, exposure)
            estimate = float(row["Coefficient Log Points"])
            probability = float(row["Probability Value"])
            ws.cell(current_row, column, f"{estimate:.3f}{stars(probability)}")
            ws.cell(
                current_row + 1,
                column,
                f'({float(row["Clustered Standard Error"]):.3f})',
            )
            style_cell(ws.cell(current_row, column), bold=is_heat)
            style_cell(ws.cell(current_row + 1, column), italic=True)
        if is_heat:
            for cell in ws[current_row] + ws[current_row + 1]:
                cell.fill = heat_fill
        current_row += 2

    if heat_row is None:
        raise RuntimeError("Heat row was not generated")

    percent_row = current_row
    ws.cell(percent_row, 1, "Implied change for 10 heat days (%)")
    style_cell(ws.cell(percent_row, 1), bold=True, horizontal="left")
    for column, (_, specification, _, _) in enumerate(MODELS, start=2):
        row = result_row(results, specification, "Absolute Heat Days per 10")
        ws.cell(
            column=column,
            row=percent_row,
            value=float(row["Percent Change for Exposure Unit"]),
        )
        ws.cell(percent_row, column).number_format = "0.00"
        style_cell(ws.cell(percent_row, column), bold=True)
    for cell in ws[percent_row]:
        cell.fill = heat_fill

    summary_start = percent_row + 2
    summary_rows = [
        "Survey-time fixed effects",
        "Composition controls",
        "Survey weights",
        "Observations",
        "Spatial units absorbed",
        "Spatial blocks",
    ]
    for offset, label in enumerate(summary_rows):
        row_number = summary_start + offset
        ws.cell(row_number, 1, label)
        style_cell(ws.cell(row_number, 1), horizontal="left")
        for column, (_, specification, _, _) in enumerate(MODELS, start=2):
            representative = result_row(results, specification, "Absolute Heat Days per 10")
            if label == "Survey-time fixed effects":
                value = "Yes"
            elif label == "Composition controls":
                value = "Yes" if column == 9 else "No"
            elif label == "Survey weights":
                value = "Yes"
            elif label == "Observations":
                value = int(representative["Observations"])
            elif label == "Spatial units absorbed":
                value = int(representative["Spatial Units"])
            else:
                value = int(representative["Spatial Blocks"])
            ws.cell(row_number, column, value)
            style_cell(ws.cell(row_number, column))
            if label in {"Observations", "Spatial units absorbed", "Spatial blocks"}:
                ws.cell(row_number, column).number_format = "#,##0"

    for column in range(1, last_column + 1):
        ws.cell(summary_start - 1, column).border = Border(top=thin)
        ws.cell(summary_start + len(summary_rows) - 1, column).border = Border(bottom=medium)

    contrast = results.loc[
        results["Specification"].eq("Formal repeated-versus-single-wave heat contrast")
    ].iloc[0]
    notes_row = summary_start + len(summary_rows) + 2
    notes = [
        "Notes: Coefficients are log points. Heat is measured per 10 additional post-onset days at or above 35 C; the other climate measures are in SD units.",
        "All models use survey weights and survey-time fixed effects; standard errors clustered by 0.75-degree spatial block are in parentheses. ***, **, and * denote p < 0.01, p < 0.05, and p < 0.10.",
        f'The formal repeated-versus-single-wave heat-slope contrast is {float(contrast["Coefficient Log Points"]):.3f} (p = {float(contrast["Probability Value"]):.3f}); separate subgroup estimates therefore do not establish heterogeneity.',
        "The percentage row equals 100 x [exp(coefficient) - 1]. Models (6)-(8) use repeated communes and are weighted by the sum of household survey weights in each cell.",
    ]
    for offset, note in enumerate(notes):
        row = notes_row + offset
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
        ws.cell(row, 1, note)
        ws.cell(row, 1).font = Font(
            name="Times New Roman", size=9, italic=True, color="333333"
        )
        ws.cell(row, 1).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        ws.row_dimensions[row].height = 22

    ws.column_dimensions["A"].width = 42
    for column in range(2, last_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = 16
    for row in range(2, notes_row):
        if ws.row_dimensions[row].height is None:
            ws.row_dimensions[row].height = 20
    ws.row_dimensions[4].height = 30

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.18
    ws.page_margins.right = 0.18
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.print_area = f"A1:{last_letter}{notes_row + len(notes) - 1}"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == ["Regression results"]
    assert check["Regression results"].max_column == 9
    for row in check["Regression results"].iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(
                    ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                )
    commune = result_row(
        results, "All linked households, commune FE", "Absolute Heat Days per 10"
    )
    pseudo = result_row(
        results, "Commune pseudo-panel, minimum 5 households", "Absolute Heat Days per 10"
    )
    assert float(commune["Probability Value"]) < 0.05
    assert float(pseudo["Probability Value"]) < 0.05
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
