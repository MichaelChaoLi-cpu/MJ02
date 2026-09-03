#!/usr/bin/env python3
"""Build the Appendix monsoon-definition and exposure-support table.

The table reports frozen, outcome-blind definitions and national natural-unit
support. It does not select definitions using ecological or household outcomes.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "data/processed/cambodia_national_monsoon_timing_preprocessed.parquet"
ANALYSIS_OUTPUT = (
    ROOT
    / "data/exp/analysis/climate-welfare/monsoon-definition-exposure-support"
    / "table_exposure_support_rows.csv"
)
OUTPUT = ROOT / "data/exp/internal_output_archive/obsolete-formal-results-20260824/tables/Table_monsoon_definition_and_exposure_support.xlsx"


VARIABLES = [
    {
        "measure": "Wet-season onset (day of year)",
        "candidate": "Approved",
        "variable": "Wet-Season Onset DOY Candidate B",
        "definition": "First 3-day rainfall total ≥30 mm",
        "window": "15 Apr–31 Aug; no 14-day total <5 mm in next 30 days",
        "role": "Frozen primary onset definition",
    },
    {
        "measure": "Longest intraseasonal dry spell (days)",
        "candidate": "Approved",
        "variable": "Longest Intraseasonal Dry Spell Days Candidate B",
        "definition": "Longest run with daily rainfall <1 mm",
        "window": "Approved onset–31 Oct; local 1991–2020 z-score in regressions",
        "role": "Primary ecological exposure",
    },
]

for threshold in (33, 35, 37):
    role = "Threshold robustness"
    if threshold == 35:
        role = "Primary household heat exposure"
    if threshold == 37:
        role += "; sparse upper threshold"
    VARIABLES.append(
        {
            "measure": f"Absolute heat days ≥{threshold}°C",
            "candidate": "Approved",
            "variable": f"Post-Onset Absolute Heat Day Count {threshold} C Candidate B",
            "definition": f"Count of days with daily Tmax ≥{threshold}°C",
            "window": "Approved onset–31 Oct; natural count; effect per 10 days",
            "role": role,
        }
    )

VARIABLES.append(
    {
        "measure": "Heat degree-days above 35°C",
        "candidate": "Approved",
        "variable": "Post-Onset Heat Degree-Days Above 35 C Candidate B",
        "definition": "Sum of max(daily Tmax − 35°C, 0)",
        "window": "Approved onset–31 Oct; natural °C-days",
        "role": "Secondary heat-intensity diagnostic",
    }
)


def build_rows(frame: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for spec in VARIABLES:
        observed = frame.loc[frame[spec["variable"]].notna(), [
            "Climate Cell ID", "Year", spec["variable"]
        ]].copy()
        values = observed[spec["variable"]].astype(float)
        rows.append(
            {
                "Measure": spec["measure"],
                "Rule": spec["candidate"],
                "Frozen definition": spec["definition"],
                "Window / transformation": spec["window"],
                "Final role / limitation": spec["role"],
                "Observed cell-years": len(values),
                "Coverage (%)": 100 * len(values) / len(frame),
                "Climate cells": observed["Climate Cell ID"].nunique(),
                "Years": observed["Year"].nunique(),
                "Mean": values.mean(),
                "P10": values.quantile(0.10),
                "Median": values.median(),
                "P90": values.quantile(0.90),
            }
        )
    return pd.DataFrame(rows)


def main() -> None:
    columns = ["Climate Cell ID", "Year"] + [str(spec["variable"]) for spec in VARIABLES]
    frame = pd.read_parquet(INPUT, columns=columns)
    frame = frame.loc[frame["Year"].between(1991, 2024)].copy()
    table = build_rows(frame)

    ANALYSIS_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(ANALYSIS_OUTPUT, index=False)

    wb = Workbook()
    ws = wb.active
    ws.title = "Exposure support"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    header_fill = PatternFill("solid", fgColor="E5E5E5")
    primary_fill = PatternFill("solid", fgColor="EAF3EE")
    sparse_fill = PatternFill("solid", fgColor="F9EFE8")

    ws.merge_cells("A1:M1")
    title = ws["A1"]
    title.value = "Monsoon Definition and Exposure Support"
    title.font = Font(name="Times New Roman", size=14, bold=True)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    headers = list(table.columns)
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=column, value=header)
        cell.font = Font(name="Times New Roman", size=8.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = header_fill
        cell.border = Border(top=medium, bottom=thin)
    ws.row_dimensions[3].height = 34

    for row_index, record in enumerate(table.to_dict("records"), start=4):
        is_primary = record["Final role / limitation"] in {
            "Frozen primary onset definition",
            "Primary ecological exposure",
            "Primary household heat exposure",
        }
        is_sparse = "sparse upper threshold" in str(record["Final role / limitation"])
        for column, header in enumerate(headers, start=1):
            value = record[header]
            cell = ws.cell(row=row_index, column=column, value=value)
            cell.font = Font(name="Times New Roman", size=8.5)
            cell.alignment = Alignment(
                horizontal="left" if column <= 5 else "center",
                vertical="center",
                wrap_text=column <= 5,
            )
            if is_primary:
                cell.fill = primary_fill
            elif is_sparse:
                cell.fill = sparse_fill
            if header in {"Observed cell-years", "Climate cells", "Years"}:
                cell.number_format = "#,##0"
            elif header == "Coverage (%)":
                cell.number_format = "0.0"
            elif header in {"Mean", "P10", "Median", "P90"}:
                cell.number_format = "0.0"
        ws.row_dimensions[row_index].height = 33

    final_table_row = 3 + len(table)
    for column in range(1, len(headers) + 1):
        ws.cell(final_table_row, column).border = Border(bottom=medium)

    notes = [
        "Notes: National support covers 6,252 climate cells in 1991–2024; the same climate surfaces are aggregated to 13,042 public village points at 2, 5, and 10 km.",
        "All definitions were constructed without reading ecological or household outcomes. Daily maximum air temperature is from CHIRTS-ERA5; rainfall-based anomalies use a cell-local 1991–2020 reference.",
        "The single approved monsoon definition is used throughout. Green rows are primary measures; the orange row is the sparse 37°C boundary check.",
    ]
    notes_start = final_table_row + 2
    for offset, note in enumerate(notes):
        row = notes_start + offset
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
        cell = ws.cell(row=row, column=1, value=note)
        cell.font = Font(name="Times New Roman", size=8.5, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 18

    widths = [34, 9, 35, 46, 34, 15, 12, 12, 8, 10, 9, 10, 9]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.30
    ws.page_margins.bottom = 0.30
    ws.print_area = f"A1:M{notes_start + len(notes) - 1}"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    # Re-open and validate the materialized workbook and key support facts.
    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == ["Exposure support"]
    assert len(table) == 6
    assert frame["Climate Cell ID"].nunique() == 6_252
    assert frame["Year"].nunique() == 34
    primary_heat = table.loc[
        table["Measure"].eq("Absolute heat days ≥35°C")
        & table["Rule"].eq("Approved")
    ].iloc[0]
    assert abs(float(primary_heat["Mean"]) - 7.9741) < 1e-4
    assert float(primary_heat["Coverage (%)"]) > 99.6
    for row in check.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not any(error in cell.value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?"))

    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
