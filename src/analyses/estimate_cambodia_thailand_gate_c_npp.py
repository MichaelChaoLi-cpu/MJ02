#!/usr/bin/env python3
"""Estimate the human-approved NPP-only Cambodia-Thailand Gate C model.

The target is the continuous-conflict-dose x post-2011 x drought interaction.
The primary outcome is annual NPP anomaly in kg C per m2. The standardized NPP
outcome is a scale robustness check, and positive consecutive-dry-day anomaly
is the mandatory alternative drought definition. No specification search is
performed in this script.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from linearmodels.panel import PanelOLS
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
PANEL = ROOT / "data/processed/cambodia_thailand_gate_c_annual_panel_preprocessed.parquet"
OUT = ROOT / "data/exp/experiments/cambodia-thailand-gate-c-npp"
TIDY = OUT / "gate_c_npp_main_model_tidy.csv"
WORKBOOK = OUT / "gate_c_npp_regression_summary.xlsx"
METADATA = OUT / "gate_c_npp_main_model_metadata.json"

CELL = "National Grid Cell ID"
YEAR = "Year"
CLIMATE = "Climate Cell ID"
BLOCK = "Spatial Block ID"
WEIGHT = "Gate B Binary Support Overlap Weight"
DOSE = "Continuous Conflict Dose"
POST = "Post-conflict 2012-2024"
MAIN_PERIOD = "Gate C Main Pre-Post Period"
NPP_NATURAL = "Annual Land NPP Anomaly kg C per m2"
NPP_STANDARDIZED = "Annual Land NPP Anomaly Z 2001-2020"
DRY_RAIN = "May October Dry Rainfall Intensity"
DRY_SPELL_ANOMALY = "May October Maximum Consecutive Dry Days Anomaly Z"
TARGET = "Conflict dose x post x drought"


def prepare_sample(frame: pd.DataFrame, outcome: str, hazard: str) -> pd.DataFrame:
    sample = frame.loc[frame[MAIN_PERIOD]].copy()
    sample = sample.dropna(
        subset=[outcome, hazard, CELL, YEAR, CLIMATE, BLOCK, WEIGHT, DOSE, POST]
    ).copy()
    sample["Climate Cell Year"] = (
        sample[CLIMATE].astype("string") + "__" + sample[YEAR].astype("string")
    )
    sample["Conflict dose x post"] = sample[DOSE] * sample[POST].astype(float)
    sample["Conflict dose x drought"] = sample[DOSE] * sample[hazard]
    sample[TARGET] = (
        sample[DOSE] * sample[POST].astype(float) * sample[hazard]
    )
    if not sample[WEIGHT].gt(0).all():
        raise RuntimeError("Analysis weights must be strictly positive")
    if sample.duplicated([CELL, YEAR]).any():
        raise RuntimeError("NPP analysis sample is not unique by grid cell and year")
    return sample


def fit_model(
    frame: pd.DataFrame,
    outcome: str,
    hazard: str,
    specification: str,
) -> tuple[pd.DataFrame, dict[str, object]]:
    sample = prepare_sample(frame, outcome, hazard)
    panel = sample.set_index([CELL, YEAR]).sort_index()
    regressors = [
        "Conflict dose x post",
        "Conflict dose x drought",
        TARGET,
    ]
    other_effects = pd.DataFrame(
        {
            "Climate Cell Year": pd.Categorical(panel["Climate Cell Year"]).codes,
        },
        index=panel.index,
    )
    clusters = pd.DataFrame(
        {"10 km spatial block": pd.Categorical(panel[BLOCK]).codes},
        index=panel.index,
    )
    fitted = PanelOLS(
        panel[outcome].astype(float),
        panel[regressors].astype(float),
        weights=panel[WEIGHT].astype(float),
        entity_effects=True,
        other_effects=other_effects,
        drop_absorbed=True,
        check_rank=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)
    confidence = fitted.conf_int(level=0.95)
    rows = []
    for term in regressors:
        rows.append(
            {
                "Specification": specification,
                "Outcome": outcome,
                "Hazard": hazard,
                "Term": term,
                "Estimate": float(fitted.params[term]),
                "Clustered Standard Error": float(fitted.std_errors[term]),
                "95% CI Lower": float(confidence.loc[term, "lower"]),
                "95% CI Upper": float(confidence.loc[term, "upper"]),
                "p-value": float(fitted.pvalues[term]),
                "Significant at 5%": bool(fitted.pvalues[term] < 0.05),
                "Observations": int(fitted.nobs),
                "Grid Cells": int(sample[CELL].nunique()),
                "Climate Cells": int(sample[CLIMATE].nunique()),
                "Spatial Blocks": int(sample[BLOCK].nunique()),
                "Years": int(sample[YEAR].nunique()),
                "Weighted R-squared Within": float(fitted.rsquared_within),
                "Fixed Effects": "grid cell; climate cell x year",
                "Inference": "10 km spatial-block clustered; debiased",
                "Weights": "frozen outcome-independent Gate B overlap weights",
            }
        )
    diagnostics = {
        "specification": specification,
        "outcome": outcome,
        "hazard": hazard,
        "observations": int(fitted.nobs),
        "grid_cells": int(sample[CELL].nunique()),
        "climate_cells": int(sample[CLIMATE].nunique()),
        "spatial_blocks": int(sample[BLOCK].nunique()),
        "years": [int(sample[YEAR].min()), int(sample[YEAR].max())],
        "included_year_count": int(sample[YEAR].nunique()),
        "pre_years": sorted(sample.loc[~sample[POST], YEAR].unique().astype(int).tolist()),
        "post_years": sorted(sample.loc[sample[POST], YEAR].unique().astype(int).tolist()),
        "fixed_effects": "grid cell and climate-cell by year",
        "cluster": "10 km spatial block",
        "cluster_count": int(sample[BLOCK].nunique()),
        "weights": "frozen outcome-independent Gate B overlap weights",
        "target": TARGET,
    }
    return pd.DataFrame(rows), diagnostics


def write_workbook(results: pd.DataFrame) -> None:
    specifications = list(dict.fromkeys(results["Specification"]))
    summary_rows: list[dict[str, object]] = []
    for term in ["Conflict dose x post", "Conflict dose x drought", TARGET]:
        estimate_row: dict[str, object] = {"Term / statistic": term}
        se_row: dict[str, object] = {"Term / statistic": "Clustered SE"}
        p_row: dict[str, object] = {"Term / statistic": "p-value"}
        for specification in specifications:
            row = results.loc[
                results["Specification"].eq(specification) & results["Term"].eq(term)
            ].iloc[0]
            stars = "***" if row["p-value"] < 0.01 else "**" if row["p-value"] < 0.05 else "*" if row["p-value"] < 0.10 else ""
            estimate_row[specification] = f"{row['Estimate']:.5f}{stars}"
            se_row[specification] = f"({row['Clustered Standard Error']:.5f})"
            p_row[specification] = f"{row['p-value']:.4f}"
        summary_rows.extend([estimate_row, se_row, p_row])
    for statistic, column in [
        ("Observations", "Observations"),
        ("Grid cells", "Grid Cells"),
        ("10 km spatial blocks", "Spatial Blocks"),
        ("Grid fixed effects", None),
        ("Climate-cell x year fixed effects", None),
        ("Overlap weights", None),
    ]:
        row = {"Term / statistic": statistic}
        for specification in specifications:
            first = results.loc[results["Specification"].eq(specification)].iloc[0]
            if column:
                row[specification] = int(first[column])
            else:
                row[specification] = "Yes"
        summary_rows.append(row)
    summary = pd.DataFrame(summary_rows)
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        summary.to_excel(writer, index=False, sheet_name="Regression Summary")
        sheet = writer.book["Regression Summary"]
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions["A"].width = 38
        for column in range(2, len(specifications) + 2):
            sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = 29
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.sheet_view.showGridLines = False


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    columns = [
        CELL,
        YEAR,
        CLIMATE,
        BLOCK,
        WEIGHT,
        DOSE,
        POST,
        MAIN_PERIOD,
        NPP_NATURAL,
        NPP_STANDARDIZED,
        DRY_RAIN,
        DRY_SPELL_ANOMALY,
    ]
    frame = pd.read_parquet(PANEL, columns=columns)
    frame["May October Long Dry Spell Intensity"] = np.maximum(
        frame[DRY_SPELL_ANOMALY], 0.0
    )
    model_plan = [
        (NPP_NATURAL, DRY_RAIN, "Primary: natural-unit NPP x dry-rainfall intensity"),
        (NPP_STANDARDIZED, DRY_RAIN, "Scale check: standardized NPP x dry-rainfall intensity"),
        (NPP_NATURAL, "May October Long Dry Spell Intensity", "Mandatory confirmation: natural-unit NPP x long dry-spell intensity"),
    ]
    tables = []
    diagnostics = []
    for outcome, hazard, specification in model_plan:
        print(f"Estimating {specification}", flush=True)
        table, diagnostic = fit_model(frame, outcome, hazard, specification)
        tables.append(table)
        diagnostics.append(diagnostic)
    results = pd.concat(tables, ignore_index=True)
    results.to_csv(TIDY, index=False)
    write_workbook(results)
    metadata = {
        "status": "first opened NPP-only Gate C estimates",
        "human_approval_record": "MILI-D-20260822-014",
        "specifications": diagnostics,
        "interpretation_limit": (
            "Clustered p-values are model-based. Causal promotion additionally requires "
            "the predeclared spatial-placebo and leave-one-sector diagnostics."
        ),
        "created_at_utc": datetime.now(UTC).isoformat(),
    }
    METADATA.write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    target = results.loc[results["Term"].eq(TARGET)]
    print("\nTarget coefficients")
    print(
        target[
            [
                "Specification",
                "Estimate",
                "Clustered Standard Error",
                "95% CI Lower",
                "95% CI Upper",
                "p-value",
                "Significant at 5%",
                "Observations",
                "Spatial Blocks",
            ]
        ].to_string(index=False)
    )
    print(f"\nSaved {TIDY.relative_to(ROOT)}")
    print(f"Saved {WORKBOOK.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
