#!/usr/bin/env python3
"""Regional Ecological-Economic Priority Evidence.

Plan: Report the six frozen regions on dry-spell sensitivity, P10-to-P90 NPP
translations, expanded-control food-consumption relevance, sample support,
and the fixed national-reference evidence typology.
Framework: AnaSOP Sections 5-7; the ecological and household coefficients are
displayed side by side and are not multiplied into a mediation estimate.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/exp/analysis/climate-npp/regional-dry-spell-sensitivity-and-food-consumption-relevance"
OUTPUT = ROOT / "data/results/tables/Table_regional_ecological_economic_priority_evidence.xlsx"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/regional-ecological-economic-priority-evidence"


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def build_rows() -> tuple[pd.DataFrame, dict[str, float]]:
    evidence = pd.read_csv(SOURCE / "regional_priority_evidence.csv")
    metadata = json.loads((SOURCE / "regional_priority_metadata.json").read_text(encoding="utf-8"))
    if set(evidence["Region"]) != {f"R{i}" for i in range(1, 7)}:
        raise ValueError("Expected exactly six frozen SKATER regions")
    rows = []
    for _, row in evidence.sort_values("Region ID").iterrows():
        rows.append(
            {
                "Region": row["Region"],
                "Dry-spell slope": (
                    f"{row['Dry_Spell_Coefficient_per_10_Days']:.5f}"
                    f"{stars(float(row['Dry_Spell_P_Value']))}\n"
                    f"({row['Dry_Spell_Clustered_SE']:.5f})"
                ),
                "P10-P90 NPP change": f"{row['P10_P90_NPP_Change_kg_C_per_ha']:.1f}",
                "NPP decline": (
                    f"{row['P10_P90_NPP_Loss_Percent']:.2f}%\n"
                    f"[{row['P10_P90_NPP_Loss_CI_Lower']:.2f}, "
                    f"{row['P10_P90_NPP_Loss_CI_Upper']:.2f}]"
                ),
                "Regional mean NPP": f"{row['Region_Mean_NPP']:.3f}",
                "Food-consumption difference": (
                    f"{row['Food_Percent_Difference']:.2f}%"
                    f"{stars(float(row['Food_NPP_P_Value']))}\n"
                    f"[{row['Food_Percent_CI_Lower']:.2f}, "
                    f"{row['Food_Percent_CI_Upper']:.2f}]"
                ),
                "Stage-1 villages": int(row["Stage1_Villages"]),
                "Households": int(row["Households"]),
                "Evidence type": row["Evidence_Type"],
            }
        )
    table = pd.DataFrame(rows)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    evidence.to_csv(EVIDENCE / "regional_priority_evidence_full.csv", index=False)
    table.to_csv(EVIDENCE / "regional_priority_table_display.csv", index=False)
    return table, metadata


def write_workbook(table: pd.DataFrame, metadata: dict[str, float]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Regional Priority"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    blue = "3F7CAC"
    teal = "3A9D8F"
    gold = "D4A72C"
    pale_green = "E2F0D9"
    pale_orange = "FCE4D6"
    light = "F4F6F7"
    white = "FFFFFF"
    thin = Side(style="thin", color="C8D5DE")
    medium = Side(style="medium", color="7F9DB9")

    ws["A1"] = "Table 5. Regional Ecological-Economic Priority Evidence"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A3"] = (
        f"Pooled dry-spell contrast: {metadata['dry_spell_p10_days']:.2f} to "
        f"{metadata['dry_spell_p90_days']:.2f} days "
        f"({metadata['dry_spell_contrast_days']:.2f} days)"
    )
    ws["A3"].font = Font(name="Times New Roman", size=9.2, italic=True, color="3F4B52")

    headers = [
        "Region",
        "Dry-spell slope\n(per 10 days)",
        "P10-P90 NPP change\n(kg C/ha)",
        "NPP decline\n(% regional mean) [95% CI]",
        "Regional mean NPP\n(kg C/m²)",
        "Food-consumption difference\n(per 0.1 NPP) [95% CI]",
        "Stage-1\nvillages",
        "Stage-2\nhouseholds",
        "Evidence type",
    ]
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor=navy if column in {1, 9} else blue if column <= 5 else teal)
        cell.font = Font(name="Times New Roman", size=8.8, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=medium)
    ws.row_dimensions[4].height = 42

    columns = list(table.columns)
    for row_number, record in enumerate(table.itertuples(index=False, name=None), start=5):
        for column, value in enumerate(record, start=1):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="Times New Roman", size=9.0, bold=column in {1, 9})
            cell.alignment = Alignment(
                horizontal="left" if column == 9 else "center",
                vertical="center",
                wrap_text=True,
                indent=1 if column == 9 else 0,
            )
            cell.border = Border(bottom=thin)
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            if columns[column - 1] == "Evidence type":
                if value == "Jointly elevated":
                    cell.fill = PatternFill("solid", fgColor=pale_green)
                elif value == "Discordant":
                    cell.fill = PatternFill("solid", fgColor=pale_orange)
        ws.row_dimensions[row_number].height = 38

    note_row = 12
    notes = [
        "Notes: Dry-spell slopes come from one common regional-interaction model with village and year fixed effects, annual-rainfall control, and village-clustered standard errors.",
        "NPP translations apply the pooled P10-P90 dry-spell contrast to each regional slope; percentages use each region's mean annual strict-cropland NPP.",
        "Food-consumption associations use survey weights, exact survey-time effects, household-composition controls, and socioeconomic controls. The two evidence dimensions are not multiplied.",
        (
            "Jointly elevated requires both point estimates to exceed the corresponding national magnitude and both 95% intervals to exclude zero; "
            f"national references are {metadata['national_dry_spell_loss_percent']:.2f}% NPP decline and "
            f"{metadata['national_expanded_control_food_percent_difference']:.2f}% food-consumption difference."
        ),
        "*** p<0.01; ** p<0.05; * p<0.10. Evidence types identify priorities for targeted monitoring and contextual assessment.",
    ]
    for note in notes:
        ws.cell(note_row, 1, note)
        ws.cell(note_row, 1).font = Font(name="Times New Roman", size=8.0, italic=True, color="3F4B52")
        # Keep note rows unmerged for deterministic DOCX export.  With wrapping
        # disabled, Calc/Excel lets the text flow across the otherwise empty
        # cells, so the standalone workbook also remains legible.
        ws.cell(note_row, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        ws.row_dimensions[note_row].height = 17
        note_row += 1

    widths = [9, 17, 17, 20, 17, 22, 12, 13, 24]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 80
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.25
    ws.print_area = f"A1:I{note_row - 1}"
    wb.properties.title = "Regional Ecological-Economic Priority Evidence"
    wb.properties.subject = "Regional dry-spell sensitivity and expanded-control food-consumption relevance"
    wb.properties.creator = "Mike Li"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Regional Priority"]
    ws = wb["Regional Priority"]
    assert ws.max_column == 9
    assert not ws.merged_cells.ranges
    assert [ws.cell(row, 1).value for row in range(5, 11)] == [f"R{i}" for i in range(1, 7)]
    assert ws["I5"].value is not None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    table, metadata = build_rows()
    write_workbook(table, metadata)
    validate_workbook()
    print(table.to_string(index=False))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Saved evidence: {EVIDENCE.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
