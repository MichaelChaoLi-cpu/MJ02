#!/usr/bin/env python3
"""Consumption Harmonisation by Survey Wave.

Plan: Document the wave-specific construction of real annual per-capita total
and food consumption while keeping the appendix table compact and auditable.
Framework: AnaSOP Sections 4 and 5-7 Stage-2 measurement, exact survey timing,
instrument-regime, and analytical-sample rules. This is a measurement audit;
it does not estimate or alter the Stage-2 model.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "data/processed/cses_household_total_consumption_preprocessed.parquet"
RECALL_AUDIT = ROOT / (
    "data/exp/data-preprocessing/climate-npp-two-stage/consumption/"
    "nonfood_recall_annualization_audit.csv"
)
EVIDENCE = ROOT / (
    "data/exp/analysis/climate-npp/consumption-harmonisation-by-survey-wave"
)
OUTPUT = ROOT / "data/results/tables/Table_consumption_harmonisation_by_survey_wave.xlsx"

WAVES = ["2004", "2007", "2009", "2011-12", "2013", "2014", "2016", "2017", "2019", "2021"]

REGIME_GROUPS = [
    {
        "Survey waves": "2004",
        "Waves": ["2004"],
        "Instrument regime": "Diagnostic",
        "Non-food annualisation": "6-month ×2; annual ×1",
        "Housing treatment": "Separate ×12; actual rent only",
        "Education treatment": "Included in recall non-food",
        "Price conversion to 2021 riels": "Unavailable",
        "Main-outcome status": "Nominal diagnostic only",
    },
    {
        "Survey waves": "2007",
        "Waves": ["2007"],
        "Instrument regime": "Integrated housing",
        "Non-food annualisation": "1-month ×12; 6-month ×2; annual ×1",
        "Housing treatment": "Embedded in recall items 1–4",
        "Education treatment": "Included in recall non-food",
        "Price conversion to 2021 riels": "Food CPI¹; annual all-items CPI",
        "Main-outcome status": "Main eligible after gate",
    },
    {
        "Survey waves": "2009, 2011–12, 2013",
        "Waves": ["2009", "2011-12", "2013"],
        "Instrument regime": "2009–13 recall",
        "Non-food annualisation": "1-month ×12; 6-month ×2; annual ×1",
        "Housing treatment": "Separate ×12; actual/equivalent rent",
        "Education treatment": "Included in recall non-food",
        "Price conversion to 2021 riels": "Food CPI¹; annual all-items CPI",
        "Main-outcome status": "Main eligible after gate",
    },
    {
        "Survey waves": "2014, 2016, 2017",
        "Waves": ["2014", "2016", "2017"],
        "Instrument regime": "2014–17 expanded recall",
        "Non-food annualisation": "1-month ×12; 6-month ×2; annual ×1",
        "Housing treatment": "Separate ×12; actual/equivalent rent",
        "Education treatment": "Included in recall non-food",
        "Price conversion to 2021 riels": "Food CPI¹; annual all-items CPI",
        "Main-outcome status": "Main eligible after gate",
    },
    {
        "Survey waves": "2019, 2021",
        "Waves": ["2019", "2021"],
        "Instrument regime": "2019–21 itemised recall",
        "Non-food annualisation": "1-month ×12; 3-month ×4; 6-month ×2; annual ×1",
        "Housing treatment": "Separate ×12; maintenance not duplicated",
        "Education treatment": "Added from person education module",
        "Price conversion to 2021 riels": "Food CPI¹; all-items CPI; education CPI",
        "Main-outcome status": "Main eligible after gate",
    },
]


def validate_recall_factors() -> None:
    audit = pd.read_csv(RECALL_AUDIT)
    observed = {
        wave: tuple(sorted(audit.loc[audit["Survey Wave"].astype(str).eq(wave), "Annualization Factor"].unique()))
        for wave in WAVES
    }
    expected = {
        "2004": (1.0, 2.0),
        "2007": (1.0, 2.0, 12.0),
        "2009": (1.0, 2.0, 12.0),
        "2011-12": (1.0, 2.0, 12.0),
        "2013": (1.0, 2.0, 12.0),
        "2014": (1.0, 2.0, 12.0),
        "2016": (1.0, 2.0, 12.0),
        "2017": (1.0, 2.0, 12.0),
        "2019": (1.0, 2.0, 4.0, 12.0),
        "2021": (1.0, 2.0, 4.0, 12.0),
    }
    assert observed == expected, (observed, expected)


def build_table() -> pd.DataFrame:
    validate_recall_factors()
    frame = pd.read_parquet(
        INPUT,
        columns=[
            "Survey Wave",
            "Household ID",
            "Nominal Annual Food Consumption Riels",
            "Nominal Annual Recall Nonfood Expenditure Riels",
            "Nominal Annual Housing Services Riels",
            "Nominal Annual Supplemental Education Expenditure Riels",
            "Real 2021 Annual Total Consumption per Capita Riels",
            "Total Consumption Main Outcome Eligible",
        ],
    )
    frame["Survey Wave"] = frame["Survey Wave"].astype(str)
    support = (
        frame.groupby("Survey Wave", as_index=False)
        .agg(
            **{
                "Released households": ("Household ID", "size"),
                "Eligible households": ("Total Consumption Main Outcome Eligible", "sum"),
                "Food observed": ("Nominal Annual Food Consumption Riels", "count"),
                "Non-food observed": ("Nominal Annual Recall Nonfood Expenditure Riels", "count"),
                "Housing observed": ("Nominal Annual Housing Services Riels", "count"),
                "Education observed": ("Nominal Annual Supplemental Education Expenditure Riels", "count"),
                "Real total observed": ("Real 2021 Annual Total Consumption per Capita Riels", "count"),
            }
        )
        .set_index("Survey Wave")
        .loc[WAVES]
    )
    expected_eligible = [0, 3593, 11970, 3592, 3840, 12090, 3839, 3840, 9827, 9935]
    assert support["Eligible households"].astype(int).tolist() == expected_eligible

    rows = []
    for group in REGIME_GROUPS:
        waves = group["Waves"]
        released = int(support.loc[waves, "Released households"].sum())
        eligible = int(support.loc[waves, "Eligible households"].sum())
        rows.append(
            {
                "Survey waves": group["Survey waves"],
                "Instrument regime": group["Instrument regime"],
                "Non-food annualisation": group["Non-food annualisation"],
                "Housing treatment": group["Housing treatment"],
                "Education treatment": group["Education treatment"],
                "Price conversion to 2021 riels": group["Price conversion to 2021 riels"],
                "Main-outcome status": group["Main-outcome status"],
                "Released households": released,
                "Eligible households": eligible,
                "Retention rate": eligible / released,
            }
        )
    table = pd.DataFrame(rows)
    assert table["Released households"].sum() == 77_904
    assert table["Eligible households"].sum() == 62_526
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    table.to_csv(EVIDENCE / "consumption_harmonisation_by_instrument_regime.csv", index=False)
    support.reset_index().to_csv(EVIDENCE / "consumption_component_support_by_survey_wave.csv", index=False)
    return table


def write_workbook(table: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Consumption Harmonisation"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    blue = "5B9BD5"
    white = "FFFFFF"
    light = "F5F7F8"
    thin = Side(style="thin", color="C8D5DE")
    medium = Side(style="medium", color="7F9DB9")
    regime_fills = {
        "Diagnostic": "E7EAEC",
        "Integrated housing": "E8F0F6",
        "2009–13 recall": "E4F2F5",
        "2014–17 expanded recall": "E6F3EF",
        "2019–21 itemised recall": "FBF1D8",
    }

    ws.merge_cells("A1:J1")
    ws["A1"] = "Consumption Harmonisation by Survey Wave"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 29

    ws.merge_cells("A3:J3")
    ws["A3"] = (
        "Common rules: food prior-7-day value ×52 · divide by household size · "
        "log outcome in regressions · no winsorisation or routine imputation"
    )
    ws["A3"].fill = PatternFill("solid", fgColor=navy)
    ws["A3"].font = Font(name="Times New Roman", size=10.5, bold=True, color=white)
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A3"].border = Border(top=medium, bottom=medium)
    ws.row_dimensions[3].height = 25

    headers = list(table.columns)
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="Times New Roman", size=9.1, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[4].height = 34

    for row_offset, record in enumerate(table.itertuples(index=False, name=None), start=5):
        regime = str(record[1])
        base_fill = regime_fills[regime]
        for column, value in enumerate(record, start=1):
            cell = ws.cell(row_offset, column, value)
            cell.font = Font(
                name="Times New Roman",
                size=8.7,
                bold=column in {1, 7, 9},
                color="666666" if str(record[0]) == "2004" else "222222",
            )
            cell.alignment = Alignment(
                horizontal="right" if column in {8, 9, 10} else ("center" if column == 1 else "left"),
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin)
            if column in {1, 2}:
                cell.fill = PatternFill("solid", fgColor=base_fill)
            elif row_offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            if column in {8, 9}:
                cell.number_format = "#,##0"
            elif column == 10:
                cell.number_format = "0.0%"
        ws.row_dimensions[row_offset].height = 43

    total_row = 10
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
    total_cell = ws.cell(total_row, 1, "All released survey waves")
    total_cell.fill = PatternFill("solid", fgColor=navy)
    total_cell.font = Font(name="Times New Roman", size=9.2, bold=True, color=white)
    total_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    total_cell.border = Border(top=medium, bottom=medium)
    released_total = int(table["Released households"].sum())
    eligible_total = int(table["Eligible households"].sum())
    total_values = {
        8: released_total,
        9: eligible_total,
        10: eligible_total / released_total,
    }
    for column, value in total_values.items():
        cell = ws.cell(total_row, column, value)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=9.2, bold=True, color=white)
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = Border(top=medium, bottom=medium)
        cell.number_format = "0.0%" if column == 10 else "#,##0"
    ws.row_dimensions[total_row].height = 25

    note_row = 12
    notes = [
        "Notes: Food values are deflated with the interview-month food CPI when available; an annual food-CPI fallback is used when interview month is unavailable.",
        "Real total consumption equals food + recall non-food + housing services + the education supplement where required. It is divided by household size; the regression outcome is its natural logarithm.",
        "Main-outcome eligibility requires a main-linked household, all real components observed, and positive per-capita total consumption. CSES 2004 remains a nominal diagnostic because comparable CPI and imputed rent are unavailable.",
        "No monetary outcome is winsorised. Missing values are not imputed; only structural education skips are set to zero in 2019/2021 when no household member attends school.",
    ]
    for note in notes:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
        cell = ws.cell(note_row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.1, italic=True, color="3F4B52")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[note_row].height = 22
        note_row += 1

    widths = [20, 22, 34, 31, 28, 37, 25, 16, 16, 14]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "C5"
    ws.sheet_view.zoomScale = 70
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.18
    ws.page_margins.top = ws.page_margins.bottom = 0.22
    ws.print_area = "A1:J15"
    wb.properties.title = "Consumption Harmonisation by Survey Wave"
    wb.properties.subject = "Stage-2 CSES consumption measurement audit"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Consumption Harmonisation"]
    ws = wb["Consumption Harmonisation"]
    assert ws.max_column == 10
    assert not ws.merged_cells.ranges
    assert ws.max_row == 15
    assert [ws.cell(4, column).value for column in range(1, 11)] == [
        "Survey waves",
        "Instrument regime",
        "Non-food annualisation",
        "Housing treatment",
        "Education treatment",
        "Price conversion to 2021 riels",
        "Main-outcome status",
        "Released households",
        "Eligible households",
        "Retention rate",
    ]
    assert ws["H10"].value == 77_904
    assert ws["I10"].value == 62_526
    assert abs(float(ws["J10"].value) - 62_526 / 77_904) < 1e-12
    assert not any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in ws.iter_rows()
        for cell in row
    )
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    table = build_table()
    write_workbook(table)
    validate_workbook()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
