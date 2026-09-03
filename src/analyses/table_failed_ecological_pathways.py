#!/usr/bin/env python3
"""Build the standalone Appendix table of failed ecological pathways."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from table_failed_pathways_and_buffering_tests import COLUMNS, ecological_rows


ROOT = Path(__file__).resolve().parents[2]
ANALYSIS_OUTPUT = (
    ROOT
    / "data/exp/analysis/climate-welfare/failed-ecological-pathways/table_rows.csv"
)
OUTPUT = ROOT / "data/exp/legacy-results/tables/Table_failed_ecological_pathways.xlsx"


def main() -> None:
    table = pd.DataFrame(ecological_rows(), columns=COLUMNS)
    assert len(table) == 12
    ANALYSIS_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(ANALYSIS_OUTPUT, index=False)

    def one_row(family: str, outcome: str | None = None, exposure: str | None = None) -> pd.Series:
        selected = table.loc[table["Evidence family"].eq(family)]
        if outcome is not None:
            selected = selected.loc[selected["Outcome"].eq(outcome)]
        if exposure is not None:
            selected = selected.loc[selected["Exposure / modifier"].str.contains(exposure, regex=False)]
        if len(selected) != 1:
            raise ValueError(
                f"Expected one row for family={family!r}, outcome={outcome!r}, exposure={exposure!r}; found {len(selected)}"
            )
        return selected.iloc[0]

    wb = Workbook()
    ws = wb.active
    ws.title = "Ecological failures"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    header_fill = PatternFill("solid", fgColor="E5E5E5")
    ecology_fill = PatternFill("solid", fgColor="EAF3EE")
    prediction_fill = PatternFill("solid", fgColor="EAF1F8")

    def add(row: int, column: int, value, *, bold: bool = False, fill=None) -> None:
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(name="Times New Roman", size=9.2, bold=bold)
        cell.alignment = Alignment(
            horizontal="left" if column == 1 else "center",
            vertical="center",
            wrap_text=True,
        )
        if fill is not None:
            cell.fill = fill
        if isinstance(value, (float, np.floating)):
            cell.number_format = "0.000"

    ws.merge_cells("A1:I1")
    ws["A1"] = "Failed Ecological Pathways"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    specs = [
        ("All-land NPP", "Annual NPP pathway", "All-land NPP", None, ecology_fill),
        ("Cropland-weighted NPP", "Annual NPP pathway", "Cropland-weighted NPP", None, ecology_fill),
        ("Full-season EVI", "Full-season vegetation pathway", "Full-season EVI", None, ecology_fill),
        ("Full-season NDVI", "Full-season vegetation pathway", "Full-season NDVI", None, ecology_fill),
        ("NPP: monsoon only", "Annual NPP prediction", None, "Monsoon structure only", prediction_fill),
        ("NPP: joint model", "Annual NPP prediction", None, "Joint rainfall and monsoon structure", prediction_fill),
        ("EVI: monsoon only", "Full-season EVI prediction", None, "Monsoon structure only", prediction_fill),
        ("EVI: joint model", "Full-season EVI prediction", None, "Joint rainfall and monsoon structure", prediction_fill),
    ]
    add(3, 1, "Statistic", bold=True, fill=header_fill)
    ws.cell(3, 1).border = Border(top=medium, bottom=thin)
    for column, (label, _, _, _, _) in enumerate(specs, start=2):
        add(3, column, label, bold=True, fill=header_fill)
        ws.cell(3, column).border = Border(top=medium, bottom=thin)
    ws.row_dimensions[3].height = 36

    coefficient_records: dict[int, tuple[pd.Series, pd.Series]] = {}
    prediction_records: dict[int, pd.Series] = {}
    for column, (_, family, outcome, exposure, _) in enumerate(specs, start=2):
        if "prediction" in family.lower():
            prediction_records[column] = one_row(family, exposure=exposure)
        else:
            coefficient_records[column] = (
                one_row(family, outcome=outcome, exposure="Longest dry-spell"),
                one_row(family, outcome=outcome, exposure="Absolute heat"),
            )

    rows = [
        (4, "Specification"),
        (6, "Longest dry-spell coefficient"), (7, "95% interval"), (8, "Raw p-value"),
        (10, "Absolute-heat coefficient"), (11, "95% interval"), (12, "Raw p-value"),
        (14, "Held-out RMSE"), (15, "Improvement vs rainfall only"),
        (17, "Sample"), (18, "Fixed effects / inference"),
        (19, "Gate status"), (20, "Interpretation limit"),
    ]
    for row, label in rows:
        add(row, 1, label, bold=label in {"Gate status", "Interpretation limit"})
        ws.row_dimensions[row].height = 24 if row not in {18, 19, 20} else 32

    for column, (_, _, _, _, fill) in enumerate(specs, start=2):
        if column in coefficient_records:
            dry, heat = coefficient_records[column]
            values = {
                4: dry["Specification"],
                6: dry["Estimate / metric"], 7: dry["95% interval"], 8: dry["Raw p-value"],
                10: heat["Estimate / metric"], 11: heat["95% interval"], 12: heat["Raw p-value"],
                17: dry["Sample"], 18: dry["Fixed effects / inference"],
                19: dry["Gate status"], 20: dry["Interpretation limit"],
            }
        else:
            prediction = prediction_records[column]
            values = {
                4: prediction["Specification"], 14: prediction["Estimate / metric"],
                15: str(prediction["Adjusted / gate metric"]).replace("RMSE improvement=", ""),
                17: prediction["Sample"], 18: prediction["Fixed effects / inference"],
                19: prediction["Gate status"], 20: prediction["Interpretation limit"],
            }
        for row, _ in rows:
            add(row, column, values.get(row, "—"), fill=fill)
    for column in range(1, 10):
        ws.cell(20, column).border = Border(bottom=medium)

    notes = [
        "Notes: Coefficient significance does not override a failed held-out prediction or cross-sensor promotion gate.",
        "Annual NPP is measured in kg C m⁻²; full-season vegetation coefficients are outcome SD. Prediction rows use strict spatial-by-temporal cross-fitting.",
    ]
    for row, note in enumerate(notes, start=22):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = ws.cell(row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.5, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 18

    for column, width in enumerate([36] + [20] * 8, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.18
    ws.page_margins.right = 0.18
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.print_area = "A1:I23"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == ["Ecological failures"]
    assert table["Evidence family"].value_counts().sum() == 12
    assert abs(float(one_row("Annual NPP prediction", exposure="Joint rainfall")["Adjusted / gate metric"].split("=")[-1].replace("%", "")) + 1.53) < 0.01
    for row in check.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not any(error in cell.value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?"))
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
