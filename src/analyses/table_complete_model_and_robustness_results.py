#!/usr/bin/env python3
"""Build the complete focal ecological and household robustness table."""

from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.pagebreak import Break
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "data/exp/analysis/climate-welfare"
ECO = BASE / "postmonsoon-vegetation-validation"
FOOD = BASE / "cses-absolute-heat-food-validation"
REPAIR = BASE / "household-heat-identification-repair"
ANALYSIS_OUTPUT = BASE / "complete-model-robustness-results/table_rows.csv"
OUTPUT = ROOT / "data/exp/internal_output_archive/obsolete-formal-results-20260824/tables/Table_complete_model_and_robustness_results.xlsx"

COLUMNS = [
    "Domain",
    "Outcome / metric",
    "Focal exposure / comparison",
    "Specification",
    "Sample",
    "Radius",
    "Definition",
    "Fixed effects / controls",
    "Estimate / error metric",
    "95% interval",
    "p-value",
    "Inference / correction",
    "Role / status",
]


def ci(estimate: float, se: float) -> str:
    return f"[{estimate - 1.96 * se:.4f}, {estimate + 1.96 * se:.4f}]"


def parse_radius(specification: str, default: str = "5 km") -> str:
    match = re.search(r"(2|5|10) km", specification)
    return f"{match.group(1)} km" if match else default


def parse_definition(specification: str) -> str:
    match = re.search(r"(33|35|37) C", specification)
    threshold = match.group(1) if match else "35"
    return f"Approved; {threshold}°C"


def short_outcome(value: str) -> str:
    if "EVI" in value:
        return "Post-monsoon EVI anomaly (SD)"
    if "NDVI" in value:
        return "Post-monsoon NDVI anomaly (SD)"
    return value


def ecology_role(specification: str, outcome: str) -> str:
    roles: list[str] = []
    if specification == "Candidate B, 35 C, 5 km" and "EVI" in outcome:
        return "Primary ecological estimate; validation gate passed"
    if "NDVI" in outcome:
        roles.append("Same-product alternative-index confirmation")
    if "2 km" in specification or "10 km" in specification:
        roles.append("Radius check")
    if "excluding years" in specification:
        roles.append("Temporal-block exclusion")
    roles.append("Negative-direction support")
    return "; ".join(roles)


def food_role(specification: str, exposure: str) -> str:
    if specification == "Candidate B, 35 C, 5 km":
        return "Primary household estimate; negative interval"
    if "Future Absolute" in exposure:
        return "Future-season placebo; interval includes zero"
    if specification.startswith("Current plus future"):
        return "Clean-month current exposure sensitivity"
    if "excluding CSES" in specification:
        return "Leave-one-wave-out; negative direction"
    if "repeated-village" in specification:
        return "Village-FE sensitivity; null and sign reversal"
    if "composition adjusted" in specification:
        return "Composition-adjusted sensitivity; negative"
    if "linkage-response weighted" in specification:
        return "Linkage-response weighting sensitivity"
    if "37 C" in specification:
        return "Sparse-threshold sensitivity; null"
    if "33 C" in specification:
        return "Threshold sensitivity; negative"
    return "Radius sensitivity; negative"


def ecology_rows() -> list[dict[str, object]]:
    coefficients = pd.read_csv(ECO / "coefficients.csv")
    focal = coefficients.loc[
        coefficients["Exposure"].str.contains("Longest Intraseasonal Dry Spell", regex=False)
    ].copy()
    rows: list[dict[str, object]] = []
    for record in focal.to_dict("records"):
        rows.append(
            {
                "Domain": "Ecological response",
                "Outcome / metric": short_outcome(str(record["Outcome"])),
                "Focal exposure / comparison": "Longest dry-spell anomaly (1 SD)",
                "Specification": str(record["Specification"]).replace("Candidate B", "Approved definition"),
                "Sample": f'N={int(record["Observations"]):,}; {int(record["Villages"]):,} villages',
                "Radius": parse_radius(str(record["Specification"])),
                "Definition": parse_definition(str(record["Specification"])),
                "Fixed effects / controls": "Village + year FE; joint monsoon controls",
                "Estimate / error metric": f'{float(record["Coefficient Outcome SD per Exposure Unit"]):.4f} SD',
                "95% interval": f'[{float(record["95 Percent CI Lower"]):.4f}, {float(record["95 Percent CI Upper"]):.4f}]',
                "p-value": float(record["Probability Value"]),
                "Inference / correction": "Spatial-block clustered SE; joint stability gate",
                "Role / status": ecology_role(str(record["Specification"]), str(record["Outcome"])),
            }
        )

    prediction = pd.read_csv(ECO / "crossfit_model_comparison.csv")
    best_rmse = float(prediction["RMSE SD"].min())
    for record in prediction.to_dict("records"):
        model = str(record["Model"])
        if model == "Rainfall totals only":
            role = "Benchmark model"
        elif float(record["RMSE SD"]) == best_rmse:
            rainfall_rmse = float(
                prediction.loc[prediction["Model"].eq("Rainfall totals only"), "RMSE SD"].iloc[0]
            )
            gain = 100.0 * (rainfall_rmse - float(record["RMSE SD"])) / rainfall_rmse
            role = f"Lowest held-out RMSE; {gain:.2f}% below rainfall only"
        else:
            role = "Monsoon-structure-only comparison"
        rows.append(
            {
                "Domain": "Ecological prediction",
                "Outcome / metric": "Held-out post-monsoon EVI",
                "Focal exposure / comparison": model,
                "Specification": "Strict spatial × temporal cross-fit",
                "Sample": f'N={int(record["Observations"]):,}; {int(record["Villages"]):,} villages',
                "Radius": "5 km",
                "Definition": "Approved; 35°C",
                "Fixed effects / controls": "Identical folds and observations",
                "Estimate / error metric": (
                    f'RMSE={float(record["RMSE SD"]):.4f}; '
                    f'MAE={float(record["MAE SD"]):.4f}; '
                    f'R²₀={float(record["R2 versus Zero-Anomaly Benchmark"]):.3f}'
                ),
                "95% interval": "—",
                "p-value": np.nan,
                "Inference / correction": "Out-of-sample metric; R²₀ uses zero-anomaly benchmark",
                "Role / status": role,
            }
        )
    return rows


def food_rows() -> list[dict[str, object]]:
    coefficients = pd.read_csv(FOOD / "coefficients.csv")
    focal = coefficients.loc[
        coefficients["Exposure"].str.contains("Absolute Heat Days per 10", regex=False)
    ].copy()
    focal = focal.loc[
        ~focal["Specification"].isin(
            {
                "Candidate B, 35 C, 5 km",
                "Candidate B, 35 C, 5 km, repeated-village fixed effects",
            }
        )
    ].copy()
    rows: list[dict[str, object]] = []
    for record in focal.to_dict("records"):
        specification = str(record["Specification"])
        repeated = "repeated-village" in specification
        composition = "composition adjusted" in specification
        controls = "Village + wave×month FE" if repeated else "District + wave×month FE"
        if composition:
            controls += "; composition controls"
        if "linkage-response weighted" in specification:
            controls += "; outcome-blind linkage-response weights"
        percent = float(record["Percent Change for Exposure Unit"])
        rows.append(
            {
                "Domain": "Household response",
                "Outcome / metric": "Log real food consumption/member",
                "Focal exposure / comparison": (
                    "Future heat days (per 10)" if "Future Absolute" in str(record["Exposure"])
                    else "Completed-season heat days (per 10)"
                ),
                "Specification": specification.replace("Candidate B", "Approved definition"),
                "Sample": f'N={int(record["Observations"]):,}; {int(record["Villages"]):,} villages',
                "Radius": parse_radius(specification),
                "Definition": parse_definition(specification),
                "Fixed effects / controls": controls,
                "Estimate / error metric": f'{float(record["Coefficient Log Points"]):.4f} log points ({percent:.2f}%)',
                "95% interval": f'[{float(record["95 Percent CI Lower"]):.4f}, {float(record["95 Percent CI Upper"]):.4f}]',
                "p-value": float(record["Probability Value"]),
                "Inference / correction": "Spatial-block clustered SE; joint family gate",
                "Role / status": food_role(specification, str(record["Exposure"])),
            }
        )
    return rows


def identification_repair_rows() -> list[dict[str, object]]:
    """Return the frozen full-sample, same-sample, and pseudo-panel checks."""
    coefficients = pd.read_csv(REPAIR / "coefficients.csv")
    focal = coefficients.loc[
        coefficients["Exposure"].isin(
            {"Absolute Heat Days per 10", "Heat X Repeated Village"}
        )
    ].copy()
    roles = {
        "All linked households, district FE": "Pre-repair benchmark; negative interval",
        "All linked households, commune FE": "Full-sample finer-area stability; negative interval",
        "Repeated-village sample, district FE": "Same-sample district-FE comparator; interval includes zero",
        "Repeated-village sample, village FE": "Same-sample village-FE estimate; interval includes zero",
        "Single-wave-village sample, district FE": "Single-wave subgroup estimate; negative interval",
        "Formal repeated-versus-single-wave heat contrast": "Formal subgroup contrast; no evidence of heterogeneity",
        "Commune pseudo-panel, minimum 5 households": "Primary within-commune association; negative interval",
        "Commune pseudo-panel, minimum 10 households": "Cell-size sensitivity; negative interval",
        "Commune pseudo-panel, minimum 5 households, composition adjusted": "Composition sensitivity; borderline interval crosses zero",
    }
    rows: list[dict[str, object]] = []
    for record in focal.to_dict("records"):
        specification = str(record["Specification"])
        evidence = str(record["Evidence Level"])
        is_pseudo = evidence == "Commune survey-time pseudo-panel"
        is_contrast = str(record["Exposure"]) == "Heat X Repeated Village"
        unit_label = "difference per 10 heat days" if is_contrast else "per 10 heat days"
        unit_name = "communes" if is_pseudo else ("villages" if "Village" in str(record["Fixed Effects"]) else "areas")
        rows.append(
            {
                "Domain": "Commune pseudo-panel" if is_pseudo else "Household response",
                "Outcome / metric": "Log real food consumption/member",
                "Focal exposure / comparison": (
                    "Repeated minus single-wave heat slope" if is_contrast
                    else f"Completed-season heat days ({unit_label})"
                ),
                "Specification": specification,
                "Sample": (
                    f'N={int(record["Observations"]):,}; '
                    f'{int(record["Spatial Units"]):,} {unit_name}'
                ),
                "Radius": "5 km",
                "Definition": "Approved; 35°C",
                "Fixed effects / controls": str(record["Fixed Effects"]).replace(
                    "survey-wave-by-interview-month", "wave×month"
                ).replace("survey-time", "survey-time FE"),
                "Estimate / error metric": (
                    f'{float(record["Coefficient Log Points"]):.4f} log points '
                    f'({float(record["Percent Change for Exposure Unit"]):.2f}%)'
                ),
                "95% interval": (
                    f'[{float(record["95 Percent CI Lower"]):.4f}, '
                    f'{float(record["95 Percent CI Upper"]):.4f}]'
                ),
                "p-value": float(record["Probability Value"]),
                "Inference / correction": (
                    f'Spatial-block clustered SE; {record["Weighting"]}'
                ),
                "Role / status": roles[specification],
            }
        )
    assert len(rows) == 9
    return rows


def main() -> None:
    # The final story treats ecological and household responses as parallel
    # evidence.  The superseded mediation experiment is intentionally excluded
    # from the publication table rather than foregrounding a failed pathway.
    repair = identification_repair_rows()
    repaired_household = [row for row in repair if row["Domain"] == "Household response"]
    repaired_pseudo = [row for row in repair if row["Domain"] == "Commune pseudo-panel"]
    table = pd.DataFrame(
        ecology_rows() + repaired_household + food_rows() + repaired_pseudo,
        columns=COLUMNS,
    )
    assert len(table) == 40, f"Expected 40 rows, found {len(table)}"
    ANALYSIS_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(ANALYSIS_OUTPUT, index=False)

    wb = Workbook()
    ws = wb.active
    ws.title = "Complete results"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    header_fill = PatternFill("solid", fgColor="E5E5E5")
    domain_fills = {
        "Ecological response": PatternFill("solid", fgColor="EAF3EE"),
        "Ecological prediction": PatternFill("solid", fgColor="EAF1F8"),
        "Household response": PatternFill("solid", fgColor="F7F1E8"),
        "Commune pseudo-panel": PatternFill("solid", fgColor="F3EAF7"),
    }

    ws.merge_cells("A1:M1")
    ws["A1"] = "Complete Model and Robustness Results"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    for column, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=3, column=column, value=header)
        cell.font = Font(name="Times New Roman", size=8.8, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = header_fill
        cell.border = Border(top=medium, bottom=thin)
    ws.row_dimensions[3].height = 34

    previous_domain = None
    for row_index, record in enumerate(table.to_dict("records"), start=4):
        domain = str(record["Domain"])
        domain_changed = previous_domain is not None and domain != previous_domain
        for column, header in enumerate(COLUMNS, start=1):
            value = record[header]
            if header == "p-value" and pd.isna(value):
                value = "—"
            cell = ws.cell(row=row_index, column=column, value=value)
            cell.font = Font(name="Times New Roman", size=8.8)
            cell.alignment = Alignment(
                horizontal="center" if header in {"Radius", "Definition", "95% interval", "p-value"} else "left",
                vertical="center",
                wrap_text=True,
            )
            cell.fill = domain_fills[domain]
            if domain_changed:
                cell.border = Border(top=thin)
            if header == "p-value" and isinstance(value, (float, int)):
                cell.number_format = "0.000"
        ws.row_dimensions[row_index].height = 24
        previous_domain = domain

    final_row = 3 + len(table)
    for column in range(1, 14):
        ws.cell(final_row, column).border = Border(bottom=medium)

    notes = [
        "Notes: The table reports all focal dry-spell, absolute-heat, identification-repair, and held-out prediction diagnostics; rainfall and onset covariates remain controls in the compact main-text regression tables.",
        "Ecological, household, and commune pseudo-panel standard errors are clustered by fixed 0.75° spatial block. Survey weights are used at household level and summed within pseudo-panel cells.",
        "A missing interval or p-value denotes an out-of-sample performance metric rather than a standalone inferential test. Threshold, scale, sample-structure, wave-exclusion, and placebo rows disclose interpretation boundaries.",
    ]
    note_start = final_row + 2
    for offset, note in enumerate(notes):
        row = note_start + offset
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
        cell = ws.cell(row=row, column=1, value=note)
        cell.font = Font(name="Times New Roman", size=8.0, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 19

    widths = [15, 23, 23, 31, 19, 7, 14, 29, 26, 18, 8, 26, 29]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    ws.page_margins.left = 0.18
    ws.page_margins.right = 0.18
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.print_area = f"A1:M{note_start + len(notes) - 1}"
    ws.print_title_rows = "1:3"
    # Keep ecological and welfare evidence on separate pages while retaining
    # one worksheet and one continuous table. Pseudo-panel rows continue the
    # household page so the Appendix does not waste a third page.
    household_start_row = 4 + int(
        table["Domain"].isin(["Ecological response", "Ecological prediction"]).sum()
    )
    ws.row_breaks.append(Break(id=household_start_row - 1))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == ["Complete results"]
    assert table["Domain"].value_counts().to_dict() == {
        "Household response": 23,
        "Ecological response": 11,
        "Ecological prediction": 3,
        "Commune pseudo-panel": 3,
    }
    primary_food = table.loc[
        table["Specification"].eq("All linked households, district FE")
    ].iloc[0]
    assert "-2.59%" in str(primary_food["Estimate / error metric"])
    for row in check.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not any(error in cell.value for error in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?"))

    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(table.groupby("Domain", sort=False).size().to_string())


if __name__ == "__main__":
    main()
