#!/usr/bin/env python3
"""Analytical Samples, Variable Definitions, and Linkage.

Plan: Document climate, vegetation, village-buffer, household, and survey-wave
samples; actual interview timing; linkage loss; outcome support; and effective
spatial clustering in one main-text workbook.
Framework: AnaSOP workflow steps 1 and 3. The ecological sample follows Model A
at 5 km over 2001-2023; the household sample follows Model B with the last
fully completed May-October season and explicit public-village linkage.
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
CLIMATE = ROOT / "data/processed/cambodia_national_monsoon_timing_preprocessed.parquet"
VEGETATION = ROOT / "data/processed/cambodia_national_seasonal_vegetation_preprocessed.parquet"
VILLAGES = ROOT / "data/processed/cambodia_public_village_points_preprocessed.parquet"
ECOLOGY = ROOT / "data/processed/cambodia_public_village_monsoon_ecology_panel_preprocessed.parquet"
CSES = ROOT / "data/processed/cses_climate_ecology_welfare_frame_preprocessed.parquet"
MEDIATION = ROOT / "data/exp/analysis/climate-welfare/cses-absolute-heat-ndvi-mediation/results_summary.json"
OUTPUT = ROOT / "data/exp/legacy-results/tables/Table_analytical_samples_variable_definitions_and_linkage.xlsx"

ID = "National Village Point ID"
FOOD = "Real 2021 Food Consumption Value per Household Member Riels"
RAIN = "Last Complete Season May-October Precipitation Anomaly Z"
ONSET = "Last Complete Season Wet-Season Onset Anomaly Z Candidate B"
DRY = "Last Complete Season Longest Dry Spell Anomaly Z Candidate B"
HEAT = "Last Complete Season Absolute Heat Day Count 35 C Candidate B"

COLUMNS = [
    "Evidence family / analytical sample",
    "Period",
    "Unit",
    "Observations",
    "Spatial units",
    "Linked share",
    "Outcome support",
    "Effective clusters",
    "Key exclusion or limitation",
    "Analytical role",
]


def spatial_block(frame: pd.DataFrame) -> pd.Series:
    return (
        np.floor((frame["Point Longitude"] - 102.0) / 0.75).astype("Int64").astype(str)
        + "_"
        + np.floor((frame["Point Latitude"] - 10.0) / 0.75).astype("Int64").astype(str)
    )


def percent_text(numerator: int, denominator: int) -> str:
    return f"{numerator / denominator:.1%}" if denominator else "—"


def period_label(frame: pd.DataFrame, wave: int) -> str:
    if wave == 2019:
        clean = frame.loc[frame["Interview Timing Source Warning"].eq(0)]
        pairs = clean[["Interview Calendar Year", "Interview Month"]].drop_duplicates()
        start = pairs.sort_values(["Interview Calendar Year", "Interview Month"]).iloc[0]
        end = pairs.sort_values(["Interview Calendar Year", "Interview Month"]).iloc[-1]
        start_label = pd.Timestamp(int(start.iloc[0]), int(start.iloc[1]), 1).strftime("%b %Y")
        end_label = pd.Timestamp(int(end.iloc[0]), int(end.iloc[1]), 1).strftime("%b %Y")
        return f"{start_label}–{end_label}"
    return str(wave)


def build_table() -> tuple[pd.DataFrame, list[int]]:
    climate = pd.read_parquet(
        CLIMATE,
        columns=[
            "Climate Cell ID", "Year", "Wet-Season Onset DOY Candidate B",
            "Longest Intraseasonal Dry Spell Days Candidate B",
            "Post-Onset Absolute Heat Day Count 35 C Candidate B",
        ],
    )
    climate_complete = climate.iloc[:, 2:].notna().all(axis=1)

    vegetation = pd.read_parquet(
        VEGETATION,
        columns=[
            "Climate Cell ID", "Production Season Year",
            "November-February Mean EVI Anomaly Z",
            "November-February Mean NDVI Anomaly Z",
        ],
    )
    vegetation_complete = vegetation[
        ["November-February Mean EVI Anomaly Z", "November-February Mean NDVI Anomaly Z"]
    ].notna().all(axis=1)

    villages = pd.read_parquet(VILLAGES, columns=[ID])

    ecology_columns = [
        ID, "Buffer Radius km", "Year", "Point Longitude", "Point Latitude",
        "Village Buffer Mean November-February Mean EVI Anomaly Z",
        "Village Buffer Mean May October Precipitation Total mm Anomaly Z",
        "Village Buffer Mean Wet-Season Onset DOY Candidate B Anomaly Z",
        "Village Buffer Mean Longest Intraseasonal Dry Spell Days Candidate B Anomaly Z",
        "Village Buffer Mean Post-Onset Absolute Heat Day Count 35 C Candidate B",
    ]
    ecology = pd.read_parquet(ECOLOGY, columns=ecology_columns)
    ecological_sample = ecology.loc[
        ecology["Buffer Radius km"].eq(5) & ecology["Year"].between(2001, 2023)
    ].dropna(subset=ecology_columns[5:]).copy()
    ecological_blocks = spatial_block(ecological_sample).nunique()

    cses_columns = [
        "Survey Year", "Interview Calendar Year", "Interview Month",
        "Interview Timing Source Warning", "Household ID", "Village Code",
        "Climate Ecology Link Available", ID, "Point Longitude", "Point Latitude",
        FOOD, RAIN, ONSET, DRY, HEAT,
    ]
    cses = pd.read_parquet(CSES, columns=cses_columns)
    linked = cses.loc[cses["Climate Ecology Link Available"].eq(1)].copy()
    linked["Spatial Block"] = spatial_block(linked)
    linked["Primary Outcome Support"] = (
        linked[[FOOD, RAIN, ONSET, DRY, HEAT]].notna().all(axis=1) & linked[FOOD].gt(0)
    )
    repeated_ids = (
        linked.groupby("Village Code", observed=True)["Survey Year"].nunique().loc[lambda x: x > 1].index
    )
    repeated = linked.loc[linked["Village Code"].isin(repeated_ids)]

    mediation = json.loads(MEDIATION.read_text(encoding="utf-8"))["primary"]

    sample_rows = [
        {
            COLUMNS[0]: "National climate-definition frame",
            COLUMNS[1]: f"{int(climate['Year'].min())}–{int(climate['Year'].max())}",
            COLUMNS[2]: "Climate cell–year",
            COLUMNS[3]: len(climate),
            COLUMNS[4]: f"{climate['Climate Cell ID'].nunique():,} climate cells",
            COLUMNS[5]: None,
            COLUMNS[6]: f"Candidate B onset, dry spell, and heat complete: {climate_complete.mean():.1%}",
            COLUMNS[7]: "Not an estimation sample",
            COLUMNS[8]: "Outcome-blind 1991–2020 local reference",
            COLUMNS[9]: "Climate-definition and exposure support",
        },
        {
            COLUMNS[0]: "Season-aligned vegetation frame",
            COLUMNS[1]: f"{int(vegetation['Production Season Year'].min())}–{int(vegetation['Production Season Year'].max())}",
            COLUMNS[2]: "Climate cell–season",
            COLUMNS[3]: len(vegetation),
            COLUMNS[4]: f"{vegetation['Climate Cell ID'].nunique():,} climate cells",
            COLUMNS[5]: None,
            COLUMNS[6]: f"Post-monsoon EVI and NDVI complete: {vegetation_complete.mean():.1%}",
            COLUMNS[7]: "Not an estimation sample",
            COLUMNS[8]: "Production-season 2024 lacks complete November–February follow-up",
            COLUMNS[9]: "Ecological outcome construction",
        },
        {
            COLUMNS[0]: "National public-village frame",
            COLUMNS[1]: "Time-invariant",
            COLUMNS[2]: "Village point",
            COLUMNS[3]: len(villages),
            COLUMNS[4]: f"{villages[ID].nunique():,} villages",
            COLUMNS[5]: None,
            COLUMNS[6]: "Stable 2, 5, and 10 km buffer identifiers",
            COLUMNS[7]: "Not an estimation sample",
            COLUMNS[8]: "Public point represents village location, not polygon extent",
            COLUMNS[9]: "National ecological frame and risk geography",
        },
        {
            COLUMNS[0]: "Primary post-monsoon ecological sample",
            COLUMNS[1]: "2001–2023",
            COLUMNS[2]: "Village–season, 5 km",
            COLUMNS[3]: len(ecological_sample),
            COLUMNS[4]: f"{ecological_sample[ID].nunique():,} villages",
            COLUMNS[5]: None,
            COLUMNS[6]: "Post-monsoon EVI and joint climate family complete",
            COLUMNS[7]: f"{ecological_blocks} spatial blocks",
            COLUMNS[8]: "Complete cases; 2024 excluded from post-monsoon response",
            COLUMNS[9]: "Primary Model A estimation",
        },
        {
            COLUMNS[0]: "Strict spatial–temporal cross-fit sample",
            COLUMNS[1]: "2001–2023",
            COLUMNS[2]: "Held-out village–season",
            COLUMNS[3]: len(ecological_sample),
            COLUMNS[4]: f"{ecological_sample[ID].nunique():,} villages",
            COLUMNS[5]: None,
            COLUMNS[6]: "Same outcome and exposure support as Model A",
            COLUMNS[7]: f"25 held-out cells; {ecological_blocks} spatial blocks",
            COLUMNS[8]: "Training excludes evaluation spatial and temporal folds",
            COLUMNS[9]: "Incremental-prediction validation",
        },
        {
            COLUMNS[0]: "Harmonized CSES household frame",
            COLUMNS[1]: "2007–2021 waves",
            COLUMNS[2]: "Household",
            COLUMNS[3]: len(cses),
            COLUMNS[4]: f"{cses['Village Code'].nunique():,} coded villages",
            COLUMNS[5]: len(linked) / len(cses),
            COLUMNS[6]: f"Positive real food outcome: {(cses[FOOD] > 0).mean():.1%}",
            COLUMNS[7]: f"{linked['Spatial Block'].nunique()} linked spatial blocks",
            COLUMNS[8]: f"{len(cses) - len(linked):,} households lack an unambiguous public-point link",
            COLUMNS[9]: "Coverage and linkage denominator",
        },
        {
            COLUMNS[0]: "Linked primary food-consumption sample",
            COLUMNS[1]: "2007–2021 waves",
            COLUMNS[2]: "Linked household",
            COLUMNS[3]: len(linked),
            COLUMNS[4]: f"{linked['Village Code'].nunique():,} villages",
            COLUMNS[5]: 1.0,
            COLUMNS[6]: f"Food and Candidate B climate complete: {linked['Primary Outcome Support'].mean():.1%}",
            COLUMNS[7]: f"{linked['Spatial Block'].nunique()} spatial blocks",
            COLUMNS[8]: "Unlinked or ambiguous public-village matches excluded",
            COLUMNS[9]: "Primary Model B estimation",
        },
        {
            COLUMNS[0]: "Repeated-village confirmation sample",
            COLUMNS[1]: "2007–2021 waves",
            COLUMNS[2]: "Linked household",
            COLUMNS[3]: len(repeated),
            COLUMNS[4]: f"{repeated['Village Code'].nunique():,} villages",
            COLUMNS[5]: 1.0,
            COLUMNS[6]: f"Food and Candidate B climate complete: {repeated['Primary Outcome Support'].mean():.1%}",
            COLUMNS[7]: f"{repeated['Spatial Block'].nunique()} spatial blocks",
            COLUMNS[8]: "Village must appear in more than one survey wave",
            COLUMNS[9]: "Village-fixed-effect confirmation",
        },
        {
            COLUMNS[0]: "Strictly ordered NDVI mediation diagnostic",
            COLUMNS[1]: "2007–2021 waves",
            COLUMNS[2]: "Linked household",
            COLUMNS[3]: int(mediation["households"]),
            COLUMNS[4]: f"{int(mediation['villages']):,} villages",
            COLUMNS[5]: 1.0,
            COLUMNS[6]: f"Food, heat, and post-monsoon NDVI; {int(mediation['village_seasons']):,} village-seasons",
            COLUMNS[7]: f"{int(mediation['spatial_blocks'])} spatial blocks",
            COLUMNS[8]: "March–October interviews only; mediation gate failed",
            COLUMNS[9]: "Appendix mechanism diagnostic",
        },
    ]

    wave_rows = []
    for wave, frame in cses.groupby("Survey Year", observed=True, sort=True):
        linked_wave = frame.loc[frame["Climate Ecology Link Available"].eq(1)].copy()
        linked_wave["Spatial Block"] = spatial_block(linked_wave)
        complete = (
            linked_wave[[FOOD, RAIN, ONSET, DRY, HEAT]].notna().all(axis=1)
            & linked_wave[FOOD].gt(0)
        )
        warnings = int(frame["Interview Timing Source Warning"].eq(1).sum())
        limitation = f"{len(frame) - len(linked_wave):,} households unlinked"
        if warnings:
            limitation += f"; {warnings} interview-timing warning"
        wave_rows.append(
            {
                COLUMNS[0]: f"CSES {int(wave)}",
                COLUMNS[1]: period_label(frame, int(wave)),
                COLUMNS[2]: "Household",
                COLUMNS[3]: len(frame),
                COLUMNS[4]: f"{frame['Village Code'].nunique():,} coded villages",
                COLUMNS[5]: len(linked_wave) / len(frame),
                COLUMNS[6]: f"Food + climate among linked: {complete.mean():.1%}",
                COLUMNS[7]: f"{linked_wave['Spatial Block'].nunique()} linked spatial blocks",
                COLUMNS[8]: limitation,
                COLUMNS[9]: "Contributes to linked household analysis",
            }
        )

    table = pd.DataFrame(sample_rows + wave_rows, columns=COLUMNS)
    return table, [len(sample_rows)]


def write_workbook(table: pd.DataFrame, section_breaks: list[int]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Samples and linkage"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    medium_blue = "5B9BD5"
    pale_blue = "DDEBF7"
    pale_green = "E2F0D9"
    light_gray = "F3F3F3"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    medium = Side(style="medium", color="7F9DB9")
    vertical = Side(style="thin", color="D7E0E6")

    ws.merge_cells("A1:J1")
    ws["A1"] = "Analytical Samples, Variable Definitions, and Linkage"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    header_row = 3
    for column, label in enumerate(COLUMNS, start=1):
        cell = ws.cell(header_row, column, label)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            top=medium,
            bottom=medium,
            left=vertical if column in {5, 7, 9} else None,
        )
    ws.row_dimensions[header_row].height = 42

    row_cursor = 4
    section_rows: list[int] = []
    sample_break = section_breaks[0]
    for index, values in table.iterrows():
        if index in {0, sample_break}:
            section_label = (
                "A. Evidence frames and analytical samples"
                if index == 0 else "B. Survey-wave linkage and interview timing"
            )
            ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=10)
            cell = ws.cell(row_cursor, 1, section_label)
            cell.fill = PatternFill("solid", fgColor=medium_blue)
            cell.font = Font(name="Times New Roman", size=10, bold=True, color=white)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = Border(top=medium, bottom=thin)
            ws.row_dimensions[row_cursor].height = 21
            section_rows.append(row_cursor)
            row_cursor += 1

        for column, value in enumerate(values, start=1):
            if pd.isna(value):
                value = None
            cell = ws.cell(row_cursor, column, value)
            cell.font = Font(name="Times New Roman", size=9.2)
            cell.alignment = Alignment(
                horizontal="center" if column in {4, 6} else "left",
                vertical="center",
                wrap_text=True,
                indent=1 if column in {5, 7, 8, 9, 10} else 0,
            )
            cell.border = Border(
                bottom=thin,
                left=vertical if column in {5, 7, 9} else None,
            )
            if index % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=light_gray)
        ws.cell(row_cursor, 4).number_format = "#,##0"
        ws.cell(row_cursor, 6).number_format = "0.0%"
        if values[COLUMNS[0]] in {
            "Primary post-monsoon ecological sample",
            "Linked primary food-consumption sample",
        }:
            for cell in ws[row_cursor]:
                cell.fill = PatternFill("solid", fgColor=pale_green)
                cell.font = Font(name="Times New Roman", size=9.2, bold=True)
        elif values[COLUMNS[0]] == "Strict spatial–temporal cross-fit sample":
            for cell in ws[row_cursor]:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
        ws.row_dimensions[row_cursor].height = 39 if index < sample_break else 32
        row_cursor += 1

    notes_start = row_cursor + 1
    notes = [
        "Notes: Counts and linkage shares are unweighted. Linked share is the proportion of households with an unambiguous public-village climate–ecology link.",
        "Effective spatial clusters use fixed 0.75-degree blocks. The ecological comparison uses 25 strict spatial-by-temporal held-out cells.",
        "The 2019 survey ran from July 2019 through June 2020 after raw-date recovery; one household retains a source timing warning.",
    ]
    for offset, note in enumerate(notes):
        ws.merge_cells(start_row=notes_start + offset, start_column=1, end_row=notes_start + offset, end_column=10)
        cell = ws.cell(notes_start + offset, 1, note)
        cell.font = Font(name="Times New Roman", size=8.5, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[notes_start + offset].height = 19

    widths = [31, 18, 21, 14, 22, 14, 31, 24, 41, 32]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A{header_row}:J{row_cursor - 1}"
    ws.sheet_view.zoomScale = 60
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.18
    ws.page_margins.right = 0.18
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.print_area = f"A1:J{notes_start + len(notes) - 1}"

    wb.properties.title = "Analytical Samples, Variable Definitions, and Linkage"
    wb.properties.subject = "Climate-welfare analytical support and linkage audit"
    wb.properties.creator = "Mike Li"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate(table: pd.DataFrame) -> None:
    assert table.shape == (18, 10), table.shape
    assert int(table.loc[table[COLUMNS[0]].eq("National climate-definition frame"), COLUMNS[3]].iloc[0]) == 212_568
    assert int(table.loc[table[COLUMNS[0]].eq("Primary post-monsoon ecological sample"), COLUMNS[3]].iloc[0]) == 299_923
    assert int(table.loc[table[COLUMNS[0]].eq("Harmonized CSES household frame"), COLUMNS[3]].iloc[0]) == 62_920
    assert int(table.loc[table[COLUMNS[0]].eq("Linked primary food-consumption sample"), COLUMNS[3]].iloc[0]) == 46_445
    assert int(table.loc[table[COLUMNS[0]].eq("Repeated-village confirmation sample"), COLUMNS[3]].iloc[0]) == 21_724
    assert int(table.loc[table[COLUMNS[0]].eq("Strictly ordered NDVI mediation diagnostic"), COLUMNS[3]].iloc[0]) == 24_585
    assert table.loc[table[COLUMNS[0]].eq("CSES 2019"), COLUMNS[1]].iloc[0] == "Jul 2019–Jun 2020"

    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Samples and linkage"]
    ws = wb["Samples and linkage"]
    assert ws.max_column == 10
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    table, section_breaks = build_table()
    write_workbook(table, section_breaks)
    validate(table)
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
