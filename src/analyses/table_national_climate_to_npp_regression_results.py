#!/usr/bin/env python3
"""National Climate-to-NPP Regression Results.

Plan: Present the fixed Stage-1 national regression ladder and the alternative
heat-intensity model in compact Stargazer style.
Framework: AnaSOP Sections 5-7 village and year fixed effects, village-clustered
inference, a common complete sample, natural-unit reporting increments, and
separate heat-day and heat-degree-day specifications.
"""

from __future__ import annotations

from pathlib import Path

from linearmodels.iv import AbsorbingLS
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from workbook_layout import unmerge_display_spans


ROOT = Path(__file__).resolve().parents[2]
PANEL = ROOT / "data/processed/cses_public_village_cropland_npp_absolute_climate_panel_preprocessed.parquet"
OUTPUT = ROOT / "data/results/tables/Table_national_climate_to_npp_regression_results.xlsx"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/national-climate-to-npp-regression-results"

ID = "National Village Point ID"
YEAR = "Year"
OUTCOME = "Annual Strict-Cropland Mean NPP kg C per m2"
HEAT = "Village Buffer Mean Annual Heat Days at or Above 35 C"
HDD = "Village Buffer Mean Annual Heat Degree-Days Above 35 C"
RX5DAY = "Village Buffer Mean Annual Maximum Consecutive Five-Day Precipitation Rx5day mm"
DRY = "Village Buffer Mean Annual Maximum Consecutive Dry Days Below 1 mm"
RAIN = "Village Buffer Mean Annual Precipitation Total mm"

HEAT_10 = "Heat days ≥35 °C / 10"
HDD_10 = "Heat degree-days >35 °C / 10"
RX5DAY_10 = "Rx5day / 10 mm"
DRY_10 = "Maximum dry spell / 10 days"
RAIN_100 = "Annual precipitation / 100 mm"

MODELS = [
    ("(1)", "Heat only", [HEAT_10]),
    ("(2)", "+ rainfall", [HEAT_10, RAIN_100]),
    ("(3)", "Full primary", [HEAT_10, RX5DAY_10, DRY_10, RAIN_100]),
    ("(4)", "Full alternative", [HDD_10, RX5DAY_10, DRY_10, RAIN_100]),
]


def prepare_sample() -> pd.DataFrame:
    columns = [ID, YEAR, OUTCOME, HEAT, HDD, RX5DAY, DRY, RAIN]
    sample = pd.read_parquet(PANEL, columns=columns)
    sample = sample.loc[sample[YEAR].between(2001, 2021)].copy()
    for column in [OUTCOME, HEAT, HDD, RX5DAY, DRY, RAIN]:
        sample[column] = pd.to_numeric(sample[column], errors="coerce")
    sample = sample.dropna(subset=columns).reset_index(drop=True)
    sample[HEAT_10] = sample[HEAT] / 10.0
    sample[HDD_10] = sample[HDD] / 10.0
    sample[RX5DAY_10] = sample[RX5DAY] / 10.0
    sample[DRY_10] = sample[DRY] / 10.0
    sample[RAIN_100] = sample[RAIN] / 100.0
    return sample


def fit_model(sample: pd.DataFrame, regressors: list[str]) -> object:
    absorb = pd.DataFrame(
        {
            "Village fixed effect": sample[ID].astype("category"),
            "Calendar-year fixed effect": sample[YEAR].astype("category"),
        },
        index=sample.index,
    )
    clusters = pd.DataFrame(
        {"Village cluster": pd.Categorical(sample[ID]).codes},
        index=sample.index,
    )
    return AbsorbingLS(
        sample[OUTCOME].astype(float),
        sample[regressors].astype(float),
        absorb=absorb,
        drop_absorbed=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def formatted_coefficient(result: object, term: str) -> str:
    estimate = float(result.params[term])
    standard_error = float(result.std_errors[term])
    p_value = float(result.pvalues[term])
    return f"{estimate:.5f}{stars(p_value)}\n({standard_error:.5f})"


def estimate_models(sample: pd.DataFrame) -> tuple[dict[str, object], pd.DataFrame, pd.DataFrame]:
    results: dict[str, object] = {}
    coefficient_rows: list[dict[str, object]] = []
    support_rows: list[dict[str, object]] = []
    for number, label, regressors in MODELS:
        model = fit_model(sample, regressors)
        key = f"{number} {label}"
        results[key] = model
        intervals = model.conf_int(level=0.95)
        for term in regressors:
            coefficient_rows.append(
                {
                    "Model": key,
                    "Term": term,
                    "Estimate": float(model.params[term]),
                    "Clustered Standard Error": float(model.std_errors[term]),
                    "95 Percent CI Lower": float(intervals.loc[term, "lower"]),
                    "95 Percent CI Upper": float(intervals.loc[term, "upper"]),
                    "Probability Value": float(model.pvalues[term]),
                }
            )
        support_rows.append(
            {
                "Model": key,
                "Observations": int(model.nobs),
                "Villages": int(sample[ID].nunique()),
                "Years": int(sample[YEAR].nunique()),
                "Village Fixed Effects": "Yes",
                "Calendar-Year Fixed Effects": "Yes",
                "Village-Clustered Standard Errors": "Yes",
                "R Squared": float(model.rsquared),
                "Absorbed R Squared": float(model.absorbed_rsquared),
            }
        )
    return results, pd.DataFrame(coefficient_rows), pd.DataFrame(support_rows)


def write_workbook(results: dict[str, object], sample: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "National Stage 1"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    mid_blue = "5B9BD5"
    pale_green = "E2F0D9"
    pale_blue = "DDEBF7"
    light_gray = "F4F6F7"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7C9D6")
    medium = Side(style="medium", color="7F9DB9")
    vertical = Side(style="thin", color="D7E0E6")

    ws.merge_cells("A1:E1")
    ws["A1"] = "Table 2. National Climate-to-NPP Regression Results"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A3"] = "Dependent variable"
    ws["A3"].font = Font(name="Times New Roman", size=10, bold=True, color=white)
    ws["A3"].fill = PatternFill("solid", fgColor=navy)
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells("B3:E3")
    ws["B3"] = "Annual strict-cropland mean NPP (kg C m⁻² year⁻¹)"
    ws["B3"].font = Font(name="Times New Roman", size=10, bold=True, color=white)
    ws["B3"].fill = PatternFill("solid", fgColor=navy)
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.border = Border(top=medium, bottom=medium)
    ws.row_dimensions[3].height = 25

    ws["A4"] = "Variable"
    ws["A4"].fill = PatternFill("solid", fgColor=mid_blue)
    ws["A4"].font = Font(name="Times New Roman", size=10, bold=True, color=white)
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for column, (number, label, _) in enumerate(MODELS, start=2):
        cell = ws.cell(4, column, f"{number}\n{label}")
        cell.fill = PatternFill("solid", fgColor=mid_blue)
        cell.font = Font(name="Times New Roman", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=vertical)
    ws.row_dimensions[4].height = 36

    coefficient_terms = [HEAT_10, HDD_10, RX5DAY_10, DRY_10, RAIN_100]
    display_labels = {
        HEAT_10: "Heat days ≥35 °C (per 10 days)",
        HDD_10: "Heat degree-days >35 °C (per 10 degree-days)",
        RX5DAY_10: "Rx5day (per 10 mm)",
        DRY_10: "Maximum dry spell (per 10 days)",
        RAIN_100: "Annual precipitation (per 100 mm)",
    }

    row = 5
    model_keys = list(results)
    for index, term in enumerate(coefficient_terms):
        ws.cell(row, 1, display_labels[term])
        ws.cell(row, 1).font = Font(name="Times New Roman", size=10, bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for column, key in enumerate(model_keys, start=2):
            result = results[key]
            value = formatted_coefficient(result, term) if term in result.params.index else None
            cell = ws.cell(row, column, value)
            cell.font = Font(name="Times New Roman", size=10, bold=value is not None)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=vertical, bottom=thin)
            if column == 4:
                cell.fill = PatternFill("solid", fgColor=pale_green)
            elif column == 5:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
            elif index % 2:
                cell.fill = PatternFill("solid", fgColor=light_gray)
        ws.cell(row, 1).border = Border(bottom=thin)
        if index % 2:
            ws.cell(row, 1).fill = PatternFill("solid", fgColor=light_gray)
        ws.row_dimensions[row].height = 36
        row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row, 1, "Model structure and support")
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=mid_blue)
    ws.cell(row, 1).font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
    ws.cell(row, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 20
    row += 1

    statistic_rows = [
        ("Village fixed effects", lambda _: "Yes", None),
        ("Calendar-year fixed effects", lambda _: "Yes", None),
        ("Standard errors clustered by village", lambda _: "Yes", None),
        ("Observations", lambda model: int(model.nobs), "#,##0"),
        ("Villages", lambda _: int(sample[ID].nunique()), "#,##0"),
        ("Years", lambda _: int(sample[YEAR].nunique()), "0"),
        ("R²", lambda model: float(model.rsquared), "0.000"),
        ("Absorbed R²", lambda model: float(model.absorbed_rsquared), "0.000"),
    ]
    for index, (label, value_function, number_format) in enumerate(statistic_rows):
        label_cell = ws.cell(row, 1, label)
        label_cell.font = Font(name="Times New Roman", size=9.7, bold=index < 3)
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        label_cell.border = Border(bottom=thin)
        if index % 2:
            label_cell.fill = PatternFill("solid", fgColor=light_gray)
        for column, key in enumerate(model_keys, start=2):
            cell = ws.cell(row, column, value_function(results[key]))
            cell.font = Font(name="Times New Roman", size=9.7)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=vertical, bottom=thin)
            if number_format:
                cell.number_format = number_format
            if column == 4:
                cell.fill = PatternFill("solid", fgColor=pale_green)
            elif column == 5:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
            elif index % 2:
                cell.fill = PatternFill("solid", fgColor=light_gray)
        ws.row_dimensions[row].height = 24
        row += 1

    row += 1
    notes = [
        "Notes: Each coefficient cell reports estimate and village-clustered standard error in parentheses. All four models use the same 59,135 village-year observations, village fixed effects, and calendar-year fixed effects.",
        "Reporting increments are 10 heat days, 10 heat degree-days, 10 mm Rx5day, 10 dry-spell days, and 100 mm annual precipitation. Heat days and degree-days are never entered together.",
        "*** p<0.01; ** p<0.05; * p<0.10. Models are conditional associations, not causal estimates. Green marks the primary full model; blue marks the alternative heat-intensity model.",
    ]
    for note in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.6, italic=True, color="333333")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 22
        row += 1

    widths = [46, 23, 23, 23, 25]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.28
    ws.print_area = f"A1:E{row - 1}"
    wb.properties.title = "National Climate-to-NPP Regression Results"
    wb.properties.subject = "Stage-1 fixed-effects regression ladder"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    unmerge_display_spans(ws)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["National Stage 1"]
    ws = wb["National Stage 1"]
    assert ws.max_column == 5
    assert not ws.merged_cells.ranges
    assert ws["A1"].value == "Table 2. National Climate-to-NPP Regression Results"
    assert ws["B4"].value == "(1)\nHeat only"
    assert ws["E4"].value == "(4)\nFull alternative"
    for sheet_row in ws.iter_rows():
        for cell in sheet_row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    sample = prepare_sample()
    results, coefficients, support = estimate_models(sample)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    coefficients.to_csv(EVIDENCE / "national_climate_to_npp_regression_coefficients.csv", index=False)
    support.to_csv(EVIDENCE / "national_climate_to_npp_regression_support.csv", index=False)
    write_workbook(results, sample)
    validate()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(f"Common sample: {len(sample):,} village-years; {sample[ID].nunique():,} villages; {sample[YEAR].nunique()} years")
    print(coefficients.to_string(index=False))
    print(support.to_string(index=False))


if __name__ == "__main__":
    main()
