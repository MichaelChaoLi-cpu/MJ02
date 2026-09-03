#!/usr/bin/env python3
"""Climate-to-NPP robustness results.

Plan: summarize the prespecified Stage-1 robustness specifications and
leave-one-year coefficient ranges in one compact regression-style workbook.
Framework: AnaSOP Sections 5-8, Climate-to-NPP Robustness Results.
"""

from __future__ import annotations

import json
from math import erfc, sqrt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/exp/analysis/climate-npp/climate-to-npp-robustness-coefficients"
EVIDENCE = ROOT / "data/exp/analysis/climate-npp/climate-to-npp-robustness-results"
OUTPUT = ROOT / "data/results/tables/Table_climate_to_npp_robustness_results.xlsx"

COEFFICIENTS = SOURCE / "climate_to_npp_robustness_coefficients.csv"
LEAVE_ONE_YEAR = SOURCE / "leave_one_year_coefficient_diagnostics.csv"
SUMMARY = SOURCE / "climate_to_npp_robustness_summary.json"

EXPOSURES = {
    "Village Buffer Mean Annual Heat Days at or Above 35 C": "Heat days ≥35°C\n(per 10 days)",
    "Village Buffer Mean Annual Heat Degree-Days Above 35 C": "Degree-days >35°C\n(per 10 °C-days)",
    "Village Buffer Mean Annual Maximum Consecutive Five-Day Precipitation Rx5day mm": "Rx5day\n(per 10 mm)",
    "Village Buffer Mean Annual Maximum Consecutive Dry Days Below 1 mm": "Dry spell\n(per 10 days)",
    "Village Buffer Mean Annual Precipitation Total mm": "Annual rainfall\n(per 100 mm)",
}

SPECIFICATION_ORDER = [
    "Primary: 5 km strict cropland",
    "Annual rainfall omitted",
    "Residual annual rainfall",
    "Continuous NPP QC adjustment",
    "NPP QC lower 75%",
    "NPP QC lower 50%",
    "2 km strict cropland",
    "10 km strict cropland",
    "5 km inclusive agriculture",
    "Valid-pixel share ≥80%",
    "Valid-pixel share ≥95%",
    "Quadratic slope at median",
    "Alternative heat: degree-days",
    "Village + year clustered",
    "0.5° spatial-block clustered",
    "Conley spatial HAC: 25 km",
    "Conley spatial HAC: 50 km",
    "Conley spatial HAC: 100 km",
]

SPECIFICATION_LABELS = {
    "Primary: 5 km strict cropland": "Primary: 5 km strict cropland",
    "Annual rainfall omitted": "Annual rainfall omitted",
    "Residual annual rainfall": "Residual annual rainfall",
    "Continuous NPP QC adjustment": "+ continuous NPP QC",
    "NPP QC lower 75%": "NPP QC lower 75%",
    "NPP QC lower 50%": "NPP QC lower 50%",
    "2 km strict cropland": "2 km buffer",
    "10 km strict cropland": "10 km buffer",
    "5 km inclusive agriculture": "Inclusive agriculture",
    "Valid-pixel share ≥80%": "Valid-pixel share ≥80%",
    "Valid-pixel share ≥95%": "Valid-pixel share ≥95%",
    "Quadratic slope at median": "Quadratic slope at median",
    "Alternative heat: degree-days": "Alternative heat: degree-days",
    "Village + year clustered": "Two-way clustered SE",
    "0.5° spatial-block clustered": "Spatial-block clustered SE",
    "Conley spatial HAC: 25 km": "Conley spatial HAC: 25 km",
    "Conley spatial HAC: 50 km": "Conley spatial HAC: 50 km (focal)",
    "Conley spatial HAC: 100 km": "Conley spatial HAC: 100 km",
}

INFERENCE_LABELS = {
    "village clustered": "Village cluster",
    "village clustered; annual rainfall omitted": "Village cluster",
    "village clustered; residual rainfall": "Village cluster; residual rain",
    "village clustered; NPP QC covariate": "Village cluster; QC adjusted",
    "village clustered; NPP QC restriction": "Village cluster; QC restricted",
    "village clustered; delta method": "Village cluster; delta",
    "village_year": "Village + year",
    "spatial_block": "0.5° block",
    "conley_25km": "Bartlett 25 km",
    "conley_50km": "Bartlett 50 km",
    "conley_100km": "Bartlett 100 km",
}


def normal_p_value(estimate: float, standard_error: float) -> float:
    """Two-sided normal-reference p-value used consistently for display stars."""
    return erfc(abs(estimate / standard_error) / sqrt(2.0))


def significance_stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def build_tables() -> tuple[pd.DataFrame, pd.DataFrame]:
    coefficients = pd.read_csv(COEFFICIENTS).copy()
    leave_one_year = pd.read_csv(LEAVE_ONE_YEAR).copy()

    assert len(coefficients) == 71
    assert set(coefficients["Specification"]) == set(SPECIFICATION_ORDER)
    assert set(coefficients["Exposure"]) == set(EXPOSURES)
    assert leave_one_year["Omitted Year"].nunique() == 21

    # Convert kg C m^-2 to g C m^-2 for readable natural-unit coefficients.
    for column in ["Estimate", "Standard Error", "95 Percent CI Lower", "95 Percent CI Upper"]:
        coefficients[f"{column} (g)"] = coefficients[column] * 1000.0
    leave_one_year["Estimate (g)"] = leave_one_year["Estimate"] * 1000.0

    rows: list[dict[str, object]] = []
    for specification in SPECIFICATION_ORDER:
        subset = coefficients.loc[coefficients["Specification"].eq(specification)].set_index("Exposure")
        row: dict[str, object] = {"Specification": SPECIFICATION_LABELS[specification]}
        for exposure, label in EXPOSURES.items():
            if exposure not in subset.index:
                row[label] = ""
                continue
            result = subset.loc[exposure]
            p_value = normal_p_value(float(result["Estimate"]), float(result["Standard Error"]))
            stars = significance_stars(p_value)
            row[label] = (
                f"{float(result['Estimate (g)']):.3f}{stars}\n"
                f"({float(result['Standard Error (g)']):.3f})"
            )
        first = subset.iloc[0]
        row["Observations"] = int(first["Observations"])
        row["Villages"] = int(first["Villages"])
        row["Inference"] = INFERENCE_LABELS[str(first["Covariance"])]
        rows.append(row)
    table = pd.DataFrame(rows)

    primary = coefficients.loc[
        coefficients["Specification"].eq("Primary: 5 km strict cropland")
    ].set_index("Exposure")
    loo_rows: list[dict[str, object]] = []
    for exposure, label in EXPOSURES.items():
        if "Heat Degree-Days" in exposure:
            continue
        subset = leave_one_year.loc[leave_one_year["Exposure"].eq(exposure)]
        estimate_min = float(subset["Estimate (g)"].min())
        estimate_max = float(subset["Estimate (g)"].max())
        loo_rows.append(
            {
                "Exposure": label.replace("\n", " "),
                "Primary estimate": float(primary.loc[exposure, "Estimate (g)"]),
                "Leave-one-year minimum": estimate_min,
                "Leave-one-year maximum": estimate_max,
                "Sign stable": "Yes" if estimate_min * estimate_max > 0 else "No",
            }
        )
    loo_summary = pd.DataFrame(loo_rows)

    EVIDENCE.mkdir(parents=True, exist_ok=True)
    coefficients.to_csv(EVIDENCE / "robustness_coefficient_evidence.csv", index=False)
    table.to_csv(EVIDENCE / "robustness_regression_table.csv", index=False)
    loo_summary.to_csv(EVIDENCE / "leave_one_year_range_summary.csv", index=False)
    return table, loo_summary


def write_workbook(table: pd.DataFrame, loo_summary: pd.DataFrame) -> None:
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    thresholds = summary["npp_qc_thresholds"]
    coefficients = pd.read_csv(COEFFICIENTS)
    qc50_years = int(
        coefficients.loc[coefficients["Specification"].eq("NPP QC lower 50%"), "Years"].iloc[0]
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Climate-NPP Robustness"
    ws.sheet_view.showGridLines = False

    navy = "1F4E78"
    blue = "3F7CAC"
    teal = "3A9D8F"
    white = "FFFFFF"
    light = "F4F6F7"
    pale_blue = "E8F0F7"
    thin = Side(style="thin", color="C8D5DE")
    medium = Side(style="medium", color="7F9DB9")

    ws.merge_cells("A1:I1")
    ws["A1"] = "Climate-to-NPP Robustness Results"
    ws["A1"].font = Font(name="Times New Roman", size=15, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 29

    ws.merge_cells("A3:A4")
    ws["A3"] = "Specification"
    ws.merge_cells("B3:F3")
    ws["B3"] = "Annual cropland NPP change (g C m⁻² year⁻¹)"
    for column, label in [(7, "Observations"), (8, "Villages"), (9, "Inference")]:
        ws.merge_cells(start_row=3, start_column=column, end_row=4, end_column=column)
        ws.cell(3, column, label)
    for column in [1, 2, 7, 8, 9]:
        cell = ws.cell(3, column)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=medium)
    ws.row_dimensions[3].height = 23

    exposure_labels = list(EXPOSURES.values())
    for column, header in enumerate(exposure_labels, start=2):
        cell = ws.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor=blue if column in {2, 3, 5} else teal)
        cell.font = Font(name="Times New Roman", size=8.7, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[4].height = 34

    data_columns = ["Specification", *exposure_labels, "Observations", "Villages", "Inference"]
    for row_number, record in enumerate(table[data_columns].itertuples(index=False, name=None), start=5):
        for column, value in enumerate(record, start=1):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(
                name="Times New Roman",
                size=8.8 if column not in {2, 3, 4, 5, 6} else 8.4,
                bold=(row_number == 5 and column == 1),
            )
            cell.alignment = Alignment(
                horizontal="left" if column in {1, 9} else "center",
                vertical="center",
                wrap_text=True,
                indent=1 if column in {1, 9} else 0,
            )
            cell.border = Border(bottom=thin)
            if row_number == 5:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
            elif row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            if column in {7, 8}:
                cell.number_format = "#,##0"
        ws.row_dimensions[row_number].height = 35

    section_row = 5 + len(table) + 1
    ws.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=9)
    ws.cell(section_row, 1, "Leave-one-year stability")
    ws.cell(section_row, 1).fill = PatternFill("solid", fgColor=navy)
    ws.cell(section_row, 1).font = Font(name="Times New Roman", size=9.5, bold=True, color=white)
    ws.cell(section_row, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[section_row].height = 23

    loo_headers = [
        "Exposure",
        "Primary estimate",
        "Minimum",
        "Maximum",
        "Sign stable",
    ]
    for column, header in enumerate(loo_headers, start=1):
        end_column = 5 if column == 1 else column + 4
        start_column = 1 if column == 1 else end_column
        if column == 1:
            ws.merge_cells(start_row=section_row + 1, start_column=1, end_row=section_row + 1, end_column=5)
        cell = ws.cell(section_row + 1, start_column, header)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="Times New Roman", size=8.8, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=medium)
    ws.row_dimensions[section_row + 1].height = 26

    for row_number, record in enumerate(
        loo_summary.itertuples(index=False, name=None),
        start=section_row + 2,
    ):
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=5)
        values = [record[0], record[1], record[2], record[3], record[4]]
        target_columns = [1, 6, 7, 8, 9]
        for column, value in zip(target_columns, values):
            cell = ws.cell(row_number, column, value)
            cell.font = Font(name="Times New Roman", size=8.8)
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "center",
                vertical="center",
                indent=1 if column == 1 else 0,
            )
            cell.border = Border(bottom=thin)
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            if column in {6, 7, 8}:
                cell.number_format = "0.000"
        ws.row_dimensions[row_number].height = 23

    note_row = section_row + 2 + len(loo_summary) + 1
    notes = [
        "Notes: All models include village and year fixed effects. Coefficients are in natural units, not standardized; standard errors are in parentheses.",
        "*** p<0.01, ** p<0.05, * p<0.10 using two-sided normal-reference tests. The quadratic row reports the slope at the sample median with delta-method uncertainty.",
        "Conley rows use same-year Bartlett spatial HAC with 25, 50, and 100 km cutoffs; 50 km is focal. The degree-day coefficient uses its own natural 10 °C-day increment.",
        "Rainfall sensitivities omit annual rainfall or replace it with the village/year-residual component after Rx5day adjustment.",
        (
            f"NPP QC checks add the filled-days percentage or retain observations at or below the pooled P75 ({thresholds['NPP QC lower 75%']:.2f}%) "
            f"and P50 ({thresholds['NPP QC lower 50%']:.2f}%) cutoffs; the P50 sample spans {qc50_years} years."
        ),
        "The leave-one-year panel reports the smallest and largest coefficient after omitting each year from 2001–2021.",
    ]
    for note in notes:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=9)
        cell = ws.cell(note_row, 1, note)
        cell.font = Font(name="Times New Roman", size=8.0, italic=True, color="3F4B52")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[note_row].height = 22
        note_row += 1

    widths = [31, 16, 16, 16, 16, 16, 13, 11, 21]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 80
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.25
    ws.print_area = f"A1:I{note_row - 1}"
    wb.properties.title = "Climate-to-NPP Robustness Results"
    wb.properties.subject = "Stage-1 fixed-effects robustness and leave-one-year stability"
    wb.properties.creator = "Mike Li"
    # Preserve cell content while removing presentation-only merges for DOCX.
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    # Values from an unmerged display span live in their upper-left cells.
    # Left alignment prevents centred text from being clipped at the print
    # boundary, while unwrapped notes flow across the empty cells to their right.
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws["B3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    ws.cell(section_row, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for row_number in range(note_row - len(notes), note_row):
        ws.cell(row_number, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        ws.row_dimensions[row_number].height = 17
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


def validate_workbook() -> None:
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == ["Climate-NPP Robustness"]
    ws = wb["Climate-NPP Robustness"]
    assert ws.max_row == 36
    assert ws.max_column == 9
    assert not ws.merged_cells.ranges
    assert ws["A5"].value == "Primary: 5 km strict cropland"
    assert ws["A19"].value == "Spatial-block clustered SE"
    assert ws["A21"].value == "Conley spatial HAC: 50 km (focal)"
    assert ws["A26"].value.startswith("Heat days")
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"))


def main() -> None:
    table, loo_summary = build_tables()
    write_workbook(table, loo_summary)
    validate_workbook()
    print(f"Saved: {OUTPUT.relative_to(ROOT)}")
    print(table.to_string(index=False))
    print("\nLeave-one-year summary")
    print(loo_summary.to_string(index=False))


if __name__ == "__main__":
    main()
