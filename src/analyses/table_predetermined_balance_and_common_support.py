#!/usr/bin/env python3
"""Predetermined Balance and Common Support.

Plan: Compare four outcome-blind comparison-group constructions on a common
2001--2007 baseline and expose the balance--sample-size--geography frontier.
Framework: AnaSOP Section 5 balance/overlap gate and Section 7 Step 2.  No
post-conflict NPP observations are read or estimated.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy.spatial.distance import cdist
from sklearn.linear_model import LogisticRegression
from sklearn.preprocessing import StandardScaler


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from figure_conflict_exposure_geography_and_evidence import (  # noqa: E402
    BALANCE_FEATURES,
    REPORT_FEATURES,
    build_baseline,
    capped_calibration_weights,
    weighted_variance,
)


OUT = ROOT / "data/exp/legacy-results/tables/Table_predetermined_balance_and_common_support.xlsx"
AUDIT_DIR = ROOT / "data/exp/experiment-design/cambodia-thailand-village-common-support/method-comparison"
METHOD_OUT = AUDIT_DIR / "support_method_summary.csv"
VARIABLE_OUT = AUDIT_DIR / "support_method_variable_diagnostics.csv"
WEIGHTS_OUT = AUDIT_DIR / "support_method_village_weights.csv"
METADATA_OUT = AUDIT_DIR / "support_method_comparison_metadata.json"

BORDER_POOL_MAX_KM = 200.0
GEOGRAPHIC_CALIPER_KM = 150.0
PROPENSITY_CALIPER_SD = 0.20
MATCHES_PER_TREATED = 5

METHOD_ORDER = [
    "Unweighted border pool",
    "Pooled capped calibration",
    "Overlap weighting",
    "Geographic 1:5 matching",
    "Sector-stratified calibration",
]

METHOD_SHORT = {
    "Unweighted border pool": "Raw pool",
    "Pooled capped calibration": "Pooled calibration",
    "Overlap weighting": "Overlap weighting",
    "Geographic 1:5 matching": "Geographic matching",
    "Sector-stratified calibration": "Sector calibration",
}


def normalize(weights: np.ndarray) -> np.ndarray:
    weights = np.asarray(weights, dtype=float)
    total = weights.sum()
    if total <= 0:
        raise ValueError("Weights must have positive mass")
    return weights / total


def effective_sample_size(weights: np.ndarray) -> float:
    weights = normalize(weights)
    return float(1.0 / np.square(weights).sum())


def weighted_ks(
    treated: np.ndarray,
    treated_weights: np.ndarray,
    controls: np.ndarray,
    control_weights: np.ndarray,
) -> float:
    treated_weights = normalize(treated_weights)
    control_weights = normalize(control_weights)
    support = np.sort(np.unique(np.concatenate([treated, controls])))

    t_order = np.argsort(treated)
    c_order = np.argsort(controls)
    t_values = treated[t_order]
    c_values = controls[c_order]
    t_cumulative = np.cumsum(treated_weights[t_order])
    c_cumulative = np.cumsum(control_weights[c_order])

    t_position = np.searchsorted(t_values, support, side="right") - 1
    c_position = np.searchsorted(c_values, support, side="right") - 1
    t_cdf = np.where(t_position >= 0, t_cumulative[np.maximum(t_position, 0)], 0.0)
    c_cdf = np.where(c_position >= 0, c_cumulative[np.maximum(c_position, 0)], 0.0)
    return float(np.max(np.abs(t_cdf - c_cdf)))


def haversine_matrix(left: pd.DataFrame, right: pd.DataFrame) -> np.ndarray:
    lon1 = np.radians(left["Point Longitude"].to_numpy(float))[:, None]
    lat1 = np.radians(left["Point Latitude"].to_numpy(float))[:, None]
    lon2 = np.radians(right["Point Longitude"].to_numpy(float))[None, :]
    lat2 = np.radians(right["Point Latitude"].to_numpy(float))[None, :]
    term = (
        np.sin((lat2 - lat1) / 2) ** 2
        + np.cos(lat1) * np.cos(lat2) * np.sin((lon2 - lon1) / 2) ** 2
    )
    return 6371.0 * 2 * np.arcsin(np.sqrt(np.clip(term, 0, 1)))


def build_sample() -> tuple[pd.DataFrame, pd.DataFrame]:
    baseline = build_baseline()
    treated = baseline[
        baseline["Candidate Affected District"].eq(1)
        & baseline["Eligible pre-conflict data"]
    ].copy().reset_index(drop=True)
    controls = baseline[
        baseline["Candidate Affected District"].eq(0)
        & baseline["Eligible pre-conflict data"]
        & baseline["Border distance"].le(BORDER_POOL_MAX_KM)
    ].copy().reset_index(drop=True)
    if len(treated) != 88 or len(controls) != 6283:
        raise RuntimeError(f"Unexpected support sample: {len(treated)} treated, {len(controls)} controls")
    return treated, controls


def pooled_calibration(treated: pd.DataFrame, controls: pd.DataFrame) -> dict:
    control_weights, details = capped_calibration_weights(treated, controls)
    return {
        "method": "Pooled capped calibration",
        "estimand": "ATT-like pooled candidate districts",
        "treated_weights": np.ones(len(treated)),
        "control_weights": control_weights,
        "sector_control_weights": {},
        "settings": details,
    }


def overlap_weighting(treated: pd.DataFrame, controls: pd.DataFrame) -> dict:
    sample = pd.concat([treated, controls], ignore_index=True)
    scaler = StandardScaler()
    matrix = scaler.fit_transform(sample[REPORT_FEATURES])
    treatment = np.r_[np.ones(len(treated)), np.zeros(len(controls))]
    model = LogisticRegression(C=1e6, max_iter=20_000, solver="lbfgs")
    model.fit(matrix, treatment)
    propensity = model.predict_proba(matrix)[:, 1]
    return {
        "method": "Overlap weighting",
        "estimand": "ATO overlap population",
        "treated_weights": 1 - propensity[: len(treated)],
        "control_weights": propensity[len(treated) :],
        "sector_control_weights": {},
        "settings": {
            "logistic_features": REPORT_FEATURES,
            "treated_propensity_min": float(propensity[: len(treated)].min()),
            "treated_propensity_max": float(propensity[: len(treated)].max()),
            "control_propensity_max": float(propensity[len(treated) :].max()),
        },
    }


def geographic_matching(treated: pd.DataFrame, controls: pd.DataFrame) -> dict:
    sample = pd.concat([treated, controls], ignore_index=True)
    scaler = StandardScaler()
    matrix = scaler.fit_transform(sample[BALANCE_FEATURES])
    treatment = np.r_[np.ones(len(treated)), np.zeros(len(controls))]
    propensity_model = LogisticRegression(C=0.1, max_iter=10_000, class_weight="balanced")
    propensity_model.fit(matrix, treatment)
    propensity = propensity_model.predict_proba(matrix)[:, 1]
    logit = np.log(propensity / (1 - propensity))
    caliper = PROPENSITY_CALIPER_SD * np.std(logit)

    treated_matrix = matrix[: len(treated)]
    control_matrix = matrix[len(treated) :]
    covariate_distance = cdist(treated_matrix, control_matrix)
    geographic_distance = haversine_matrix(treated, controls)

    control_counts = np.zeros(len(controls))
    sector_counts = {
        sector: np.zeros(len(controls))
        for sector in treated["Candidate Conflict Sector"].dropna().unique()
    }
    retained = np.zeros(len(treated))
    selected_distances: list[float] = []
    for index in range(len(treated)):
        eligible = (
            (np.abs(logit[len(treated) :] - logit[index]) <= caliper)
            & (geographic_distance[index] <= GEOGRAPHIC_CALIPER_KM)
        )
        if not eligible.any():
            continue
        eligible_indices = np.flatnonzero(eligible)
        selected = eligible_indices[
            np.argsort(covariate_distance[index, eligible])[:MATCHES_PER_TREATED]
        ]
        retained[index] = 1
        control_counts[selected] += 1
        sector = treated.loc[index, "Candidate Conflict Sector"]
        sector_counts[sector][selected] += 1
        selected_distances.extend(geographic_distance[index, selected].tolist())
    return {
        "method": "Geographic 1:5 matching",
        "estimand": "ATT in propensity/geographic support",
        "treated_weights": retained,
        "control_weights": control_counts,
        "sector_control_weights": sector_counts,
        "settings": {
            "propensity_caliper_sd": PROPENSITY_CALIPER_SD,
            "realized_logit_caliper": float(caliper),
            "geographic_caliper_km": GEOGRAPHIC_CALIPER_KM,
            "matches_per_treated": MATCHES_PER_TREATED,
            "maximum_selected_pair_distance_km": float(max(selected_distances)),
        },
    }


def sector_calibration(treated: pd.DataFrame, controls: pd.DataFrame) -> dict:
    combined = np.zeros(len(controls))
    sector_weights: dict[str, np.ndarray] = {}
    details: dict[str, dict] = {}
    for sector, sector_treated in treated.groupby("Candidate Conflict Sector", observed=True):
        weights, fit = capped_calibration_weights(sector_treated, controls)
        sector_weights[str(sector)] = weights
        combined += len(sector_treated) / len(treated) * weights
        details[str(sector)] = fit
    return {
        "method": "Sector-stratified calibration",
        "estimand": "ATT-like, sector shares preserved",
        "treated_weights": np.ones(len(treated)),
        "control_weights": combined,
        "sector_control_weights": sector_weights,
        "settings": details,
    }


def raw_pool(treated: pd.DataFrame, controls: pd.DataFrame) -> dict:
    return {
        "method": "Unweighted border pool",
        "estimand": "Unadjusted reference",
        "treated_weights": np.ones(len(treated)),
        "control_weights": np.ones(len(controls)),
        "sector_control_weights": {},
        "settings": {},
    }


def variable_diagnostics(
    method: dict,
    treated: pd.DataFrame,
    controls: pd.DataFrame,
    raw_treated: pd.DataFrame,
    raw_controls: pd.DataFrame,
) -> pd.DataFrame:
    treated_weights = normalize(method["treated_weights"])
    control_weights = normalize(method["control_weights"])
    rows = []
    for variable in REPORT_FEATURES:
        treated_values = treated[variable].to_numpy(float)
        control_values = controls[variable].to_numpy(float)
        reference_sd = np.sqrt(
            (
                np.var(raw_treated[variable].to_numpy(float), ddof=1)
                + np.var(raw_controls[variable].to_numpy(float), ddof=1)
            )
            / 2
        )
        treated_mean = np.average(treated_values, weights=treated_weights)
        control_mean = np.average(control_values, weights=control_weights)
        smd = (treated_mean - control_mean) / reference_sd
        variance_ratio = weighted_variance(treated_values, treated_weights) / weighted_variance(
            control_values, control_weights
        )
        rows.append(
            {
                "Method": method["method"],
                "Variable": variable,
                "Core calibration variable": variable in BALANCE_FEATURES,
                "Treated weighted mean": treated_mean,
                "Control weighted mean": control_mean,
                "Standardized mean difference": smd,
                "Absolute standardized mean difference": abs(smd),
                "Treated/control variance ratio": variance_ratio,
                "Symmetric variance ratio": max(variance_ratio, 1 / variance_ratio),
                "Weighted KS distance": weighted_ks(
                    treated_values,
                    treated_weights,
                    control_values,
                    control_weights,
                ),
            }
        )
    return pd.DataFrame(rows)


def sector_balance(
    method: dict,
    treated: pd.DataFrame,
    controls: pd.DataFrame,
    raw_treated: pd.DataFrame,
    raw_controls: pd.DataFrame,
) -> tuple[float, float]:
    worst_smd = 0.0
    worst_ks = 0.0
    treated_method_weights = np.asarray(method["treated_weights"], float)
    for sector in treated["Candidate Conflict Sector"].dropna().unique():
        sector_mask = treated["Candidate Conflict Sector"].eq(sector).to_numpy()
        sector_treated = treated.loc[sector_mask]
        sector_treated_weights = treated_method_weights[sector_mask]
        if sector_treated_weights.sum() <= 0:
            continue
        if sector in method["sector_control_weights"]:
            sector_control_weights = method["sector_control_weights"][sector]
        else:
            sector_control_weights = method["control_weights"]
        sector_control_weights = normalize(sector_control_weights)
        sector_treated_weights = normalize(sector_treated_weights)
        for variable in BALANCE_FEATURES:
            reference_sd = np.sqrt(
                (
                    np.var(raw_treated[variable].to_numpy(float), ddof=1)
                    + np.var(raw_controls[variable].to_numpy(float), ddof=1)
                )
                / 2
            )
            smd = (
                np.average(sector_treated[variable], weights=sector_treated_weights)
                - np.average(controls[variable], weights=sector_control_weights)
            ) / reference_sd
            ks = weighted_ks(
                sector_treated[variable].to_numpy(float),
                sector_treated_weights,
                controls[variable].to_numpy(float),
                sector_control_weights,
            )
            worst_smd = max(worst_smd, abs(smd))
            worst_ks = max(worst_ks, ks)
    return worst_smd, worst_ks


def nearest_treated_distance(treated: pd.DataFrame, controls: pd.DataFrame) -> np.ndarray:
    return haversine_matrix(controls, treated).min(axis=1)


def method_summary(
    method: dict,
    diagnostics: pd.DataFrame,
    treated: pd.DataFrame,
    controls: pd.DataFrame,
    raw_treated: pd.DataFrame,
    raw_controls: pd.DataFrame,
    control_nearest_treated_km: np.ndarray,
) -> dict:
    treated_weights = normalize(method["treated_weights"])
    control_weights = normalize(method["control_weights"])
    sorted_index = np.argsort(control_weights)[::-1]
    cumulative = np.cumsum(control_weights[sorted_index])
    n95 = int(np.searchsorted(cumulative, 0.95) + 1)
    selected95 = sorted_index[:n95]
    worst_sector_smd, worst_sector_ks = sector_balance(
        method, treated, controls, raw_treated, raw_controls
    )
    return {
        "Method": method["method"],
        "Estimand": method["estimand"],
        "Treated retained": int(np.count_nonzero(np.asarray(method["treated_weights"]) > 0)),
        "Treated ESS": effective_sample_size(method["treated_weights"]),
        "Controls with positive weight": int(np.count_nonzero(np.asarray(method["control_weights"]) > 0)),
        "Controls carrying 95% weight": n95,
        "Control ESS": effective_sample_size(method["control_weights"]),
        "Maximum control weight": float(control_weights.max()),
        "Maximum absolute SMD, core": float(
            diagnostics.loc[
                diagnostics["Core calibration variable"], "Absolute standardized mean difference"
            ].max()
        ),
        "Maximum absolute SMD, all": float(
            diagnostics["Absolute standardized mean difference"].max()
        ),
        "Maximum weighted KS": float(diagnostics["Weighted KS distance"].max()),
        "Worst symmetric variance ratio": float(diagnostics["Symmetric variance ratio"].max()),
        "Worst sector-specific core SMD": worst_sector_smd,
        "Worst sector-specific core KS": worst_sector_ks,
        "Maximum nearest-treated distance among 95% controls km": float(
            control_nearest_treated_km[selected95].max()
        ),
    }


def weights_long(methods: list[dict], treated: pd.DataFrame, controls: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for method in methods:
        treated_weights = normalize(method["treated_weights"])
        control_weights = normalize(method["control_weights"])
        treated_part = treated[[
            "National Village Point ID",
            "Public Village Name",
            "Candidate Conflict Sector",
            "Point Longitude",
            "Point Latitude",
        ]].copy()
        treated_part["Method"] = method["method"]
        treated_part["Analysis role"] = "Candidate treated"
        treated_part["Analysis weight"] = treated_weights
        control_part = controls[[
            "National Village Point ID",
            "Public Village Name",
            "Candidate Conflict Sector",
            "Point Longitude",
            "Point Latitude",
        ]].copy()
        control_part["Method"] = method["method"]
        control_part["Analysis role"] = "Candidate control"
        control_part["Analysis weight"] = control_weights
        rows.extend([treated_part, control_part])
    return pd.concat(rows, ignore_index=True)


def write_workbook(summary: pd.DataFrame, variable: pd.DataFrame, metadata: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance and support"
    ws.sheet_view.showGridLines = False

    dark = "1F4E78"
    mid = "D9EAF7"
    light = "F3F6F8"
    orange = "F4B183"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="B7C9D6")

    ws.merge_cells("A1:L1")
    ws["A1"] = "Predetermined Balance and Common Support"
    ws["A1"].font = Font(name="Arial", size=15, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=dark)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:L2")
    ws["A2"] = (
        "Outcome-blind comparison using 2001–2007 pre-conflict data; 5 km village buffers; "
        "88 provisional treated villages and 6,283 border-context controls."
    )
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="404040")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30

    ws.merge_cells("A4:L4")
    ws["A4"] = "A. Method-level balance, support, and spatial concentration"
    ws["A4"].font = Font(name="Arial", size=10, bold=True, color="000000")
    ws["A4"].fill = PatternFill("solid", fgColor=orange)

    summary_columns = [
        ("Method", "Method"),
        ("Estimand", "Estimand"),
        ("Treated retained", "Treated\nretained"),
        ("Treated ESS", "Treated\nESS"),
        ("Controls carrying 95% weight", "Controls for\n95% weight"),
        ("Control ESS", "Control\nESS"),
        ("Maximum control weight", "Maximum\ncontrol weight"),
        ("Maximum absolute SMD, core", "Maximum |SMD|\ncore variables"),
        ("Maximum absolute SMD, all", "Maximum |SMD|\nall variables"),
        ("Maximum weighted KS", "Maximum\nweighted KS"),
        ("Worst sector-specific core SMD", "Worst sector\ncore |SMD|"),
        ("Maximum nearest-treated distance among 95% controls km", "95% support max\ndistance (km)"),
    ]
    for column, (_, label) in enumerate(summary_columns, start=1):
        cell = ws.cell(5, column, label)
        cell.font = Font(name="Arial", size=8, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)
    ws.row_dimensions[5].height = 36

    summary_indexed = summary.set_index("Method").loc[METHOD_ORDER].reset_index()
    for row_offset, record in enumerate(summary_indexed.to_dict("records"), start=6):
        for column, (key, _) in enumerate(summary_columns, start=1):
            cell = ws.cell(row_offset, column, record[key])
            cell.font = Font(name="Arial", size=8)
            cell.alignment = Alignment(
                horizontal="left" if column <= 2 else "center",
                vertical="center",
                wrap_text=column <= 2,
            )
            if row_offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=light)
            cell.border = Border(bottom=Border(bottom=thin_gray).bottom)
        ws.row_dimensions[row_offset].height = 30

    for row in range(6, 11):
        ws.cell(row, 3).number_format = "0"
        ws.cell(row, 4).number_format = "0.0"
        ws.cell(row, 5).number_format = "0"
        ws.cell(row, 6).number_format = "0.0"
        ws.cell(row, 7).number_format = "0.0%"
        for column in range(8, 12):
            ws.cell(row, column).number_format = "0.000"
        ws.cell(row, 12).number_format = "0.0"

    ws.merge_cells("A13:L13")
    ws["A13"] = "B. Variable-level standardized differences and distributional overlap"
    ws["A13"].font = Font(name="Arial", size=10, bold=True)
    ws["A13"].fill = PatternFill("solid", fgColor=orange)

    detail_headers = [
        "Variable",
        "Core?",
        "Raw |SMD|",
        "Pooled calibration |SMD|",
        "Overlap |SMD|",
        "Geographic match |SMD|",
        "Sector calibration |SMD|",
        "Raw KS",
        "Pooled calibration KS",
        "Overlap KS",
        "Geographic match KS",
        "Sector calibration KS",
    ]
    for column, label in enumerate(detail_headers, start=1):
        cell = ws.cell(14, column, label)
        cell.font = Font(name="Arial", size=8, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)
    ws.row_dimensions[14].height = 40

    smd_pivot = variable.pivot(
        index=["Variable", "Core calibration variable"],
        columns="Method",
        values="Absolute standardized mean difference",
    )
    ks_pivot = variable.pivot(
        index=["Variable", "Core calibration variable"],
        columns="Method",
        values="Weighted KS distance",
    )
    for row_offset, feature in enumerate(REPORT_FEATURES, start=15):
        core = feature in BALANCE_FEATURES
        ws.cell(row_offset, 1, feature)
        ws.cell(row_offset, 2, "Yes" if core else "No")
        for method_index, method in enumerate(METHOD_ORDER, start=3):
            ws.cell(row_offset, method_index, float(smd_pivot.loc[(feature, core), method]))
        for method_index, method in enumerate(METHOD_ORDER, start=8):
            ws.cell(row_offset, method_index, float(ks_pivot.loc[(feature, core), method]))
        for column in range(1, 13):
            cell = ws.cell(row_offset, column)
            cell.font = Font(name="Arial", size=8)
            cell.alignment = Alignment(horizontal="left" if column == 1 else "center", vertical="center")
            cell.border = Border(bottom=thin_gray)
            if row_offset % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=light)
        for column in range(3, 13):
            ws.cell(row_offset, column).number_format = "0.000"

    note_row = 30
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=12)
    ws.cell(note_row, 1, "C. Interpretation and audit notes")
    ws.cell(note_row, 1).font = Font(name="Arial", size=10, bold=True)
    ws.cell(note_row, 1).fill = PatternFill("solid", fgColor=orange)
    notes = [
        "No method is declared to pass automatically. Human review must consider mean balance, full-distribution overlap, effective sample size, sector balance, and geography jointly.",
        "Pooled capped calibration uses a 2% maximum control-weight cap. Sector calibration applies the same cap separately to Preah Vihear and Ta Moan–Ta Krabey, preserving their 25/88 and 63/88 shares.",
        f"Geographic matching uses 1:{MATCHES_PER_TREATED} replacement matching, a {GEOGRAPHIC_CALIPER_KM:.0f} km geographic caliper, and a propensity-logit caliper equal to {PROPENSITY_CALIPER_SD:.2f} pooled SD; these settings remain diagnostic rather than frozen.",
        "Overlap weighting targets the overlap population and therefore changes the estimand and downweights candidate treated villages with weak support.",
        "Treatment provenance remains unresolved: the 88 villages use a provisional affected-district definition, so no method in this table authorizes a causal post-conflict claim.",
    ]
    for index, note in enumerate(notes, start=note_row + 1):
        ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=12)
        cell = ws.cell(index, 1, f"• {note}")
        cell.font = Font(name="Arial", size=8)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = PatternFill("solid", fgColor="FFF2CC" if index == note_row + 1 else white)
        ws.row_dimensions[index].height = 30 if index != note_row + 3 else 38

    widths = [30, 29, 11, 11, 13, 11, 13, 14, 14, 12, 13, 16]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    ws.freeze_panes = "C6"
    ws.auto_filter.ref = "A5:L10"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:L{note_row + len(notes)}"
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

    ws["H5"].comment = Comment("SMD uses the unweighted pooled treated-control SD as the common denominator.", "Mike Li")
    ws["J5"].comment = Comment("Weighted two-sample Kolmogorov-Smirnov distance; larger values indicate weaker distributional overlap.", "Mike Li")
    ws["F5"].comment = Comment("Effective sample size equals 1 / sum(normalized weight squared).", "Mike Li")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    # Compact structural verification after write.
    check = load_workbook(OUT, data_only=False)
    if check.sheetnames != ["Balance and support"]:
        raise RuntimeError(f"Workbook must contain one sheet; found {check.sheetnames}")
    sheet = check["Balance and support"]
    if sheet.max_row < 35 or sheet.max_column != 12:
        raise RuntimeError(f"Unexpected workbook extent: {sheet.max_row} rows x {sheet.max_column} columns")
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(error in cell.value for error in ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"]):
                raise RuntimeError(f"Formula error token in {cell.coordinate}: {cell.value}")


def main() -> None:
    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    treated, controls = build_sample()
    methods = [
        raw_pool(treated, controls),
        pooled_calibration(treated, controls),
        overlap_weighting(treated, controls),
        geographic_matching(treated, controls),
        sector_calibration(treated, controls),
    ]
    control_distance = nearest_treated_distance(treated, controls)

    variable_frames = []
    summary_rows = []
    for method in methods:
        diagnostics = variable_diagnostics(method, treated, controls, treated, controls)
        variable_frames.append(diagnostics)
        summary_rows.append(
            method_summary(
                method,
                diagnostics,
                treated,
                controls,
                treated,
                controls,
                control_distance,
            )
        )
    variable = pd.concat(variable_frames, ignore_index=True)
    summary = pd.DataFrame(summary_rows)
    weights = weights_long(methods, treated, controls)

    metadata = {
        "outcome_blind_pre_period": [2001, 2007],
        "primary_buffer_radius_km": 5,
        "candidate_treated_villages": len(treated),
        "candidate_control_villages": len(controls),
        "candidate_control_border_distance_max_km": BORDER_POOL_MAX_KM,
        "methods": {method["method"]: method["settings"] for method in methods},
        "automatic_pass_rule": None,
        "human_review_required": True,
        "treatment_status": "provisional affected-district definition; not frozen",
    }

    summary.to_csv(METHOD_OUT, index=False)
    variable.to_csv(VARIABLE_OUT, index=False)
    weights.to_csv(WEIGHTS_OUT, index=False)
    METADATA_OUT.write_text(json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    write_workbook(summary, variable, metadata)

    print(summary.to_string(index=False))
    print(f"Saved: {OUT.relative_to(ROOT)}")
    print(f"Audit: {AUDIT_DIR.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
