#!/usr/bin/env python3
"""One-sheet workbook styling shared by the LongNTL planned tables."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def write_one_sheet(
    table: pd.DataFrame,
    *,
    output: Path,
    sheet: str,
    title: str,
    note: str,
    widths: list[int],
    numeric_columns: set[str] | None = None,
) -> None:
    if len(widths) != len(table.columns):
        raise ValueError("Column-width count does not match the table")
    numeric_columns = numeric_columns or set()
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name=sheet, index=False, startrow=3)
        ws = writer.book[sheet]
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(table.columns))
        ws.cell(1, 1, title)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(table.columns))
        ws.cell(2, 1, note)
        navy = "1F4E78"
        stripe = "EAF1F7"
        thin = Side(style="thin", color="B8C2CC")
        ws.cell(1, 1).font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        ws.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
        ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 24
        ws.cell(2, 1).font = Font(name="Arial", size=9, italic=True, color="444444")
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 34
        for cell in ws[4]:
            cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        ws.row_dimensions[4].height = 38
        for row_index in range(5, 5 + len(table)):
            fill = PatternFill("solid", fgColor="FFFFFF" if row_index % 2 else stripe)
            for column_index, column in enumerate(table.columns, start=1):
                cell = ws.cell(row_index, column_index)
                cell.font = Font(name="Arial", size=8.5)
                cell.fill = fill
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="D8DEE4"))
                if column in numeric_columns:
                    cell.number_format = "0.000"
        for column_index, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(4, column_index).column_letter].width = width
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{ws.cell(4 + len(table), len(table.columns)).coordinate}"
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:4"
